"""Excel I/O helpers for bulk import (distributors) and report export.

Built on :mod:`openpyxl` (already a project dependency). Synchronous parse
and serialize functions are wrapped in ``asyncio.to_thread`` so handlers
remain non-blocking.
"""
from __future__ import annotations

import asyncio
import io
from collections.abc import Iterable, Sequence
from typing import Any

from openpyxl import Workbook, load_workbook


def _read_rows_blocking(
    file_bytes: bytes,
    *,
    sheet_name: str | None,
    header: bool,
) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    rows_iter = sheet.iter_rows(values_only=True)

    headers: list[str]
    if header:
        first = next(rows_iter, None)
        if first is None:
            return [], []
        headers = [str(cell) if cell is not None else "" for cell in first]
    else:
        peek = next(rows_iter, None)
        if peek is None:
            return [], []
        headers = [f"col_{i}" for i in range(len(peek))]
        rows_iter = iter([peek, *rows_iter])

    parsed: list[dict[str, Any]] = []
    for raw_row in rows_iter:
        if all(cell is None for cell in raw_row):
            continue
        parsed.append(
            {
                headers[i]: raw_row[i]
                for i in range(min(len(headers), len(raw_row)))
            }
        )
    return headers, parsed


async def read_workbook(
    file_bytes: bytes,
    *,
    sheet_name: str | None = None,
    header: bool = True,
) -> tuple[list[str], list[dict[str, Any]]]:
    """Parse an .xlsx into ``(headers, rows)``. Empty rows are skipped."""
    return await asyncio.to_thread(
        _read_rows_blocking, file_bytes, sheet_name=sheet_name, header=header
    )


def _write_rows_blocking(
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    *,
    sheet_name: str,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def write_workbook(
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    *,
    sheet_name: str = "Sheet1",
) -> bytes:
    """Serialize ``rows`` (each a sequence aligned with ``headers``) into .xlsx bytes."""
    return await asyncio.to_thread(
        _write_rows_blocking, headers, rows, sheet_name=sheet_name
    )
