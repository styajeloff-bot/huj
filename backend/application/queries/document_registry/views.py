"""Shared authorized registry views for HTTP, monetization and notifications."""
import base64
import hashlib
import json
from datetime import date, datetime
from io import BytesIO
from typing import Any
from uuid import UUID
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession

from application.errors import ServiceError
from domain.document_registry import STATUS_LABELS, TYPES
from infrastructure.repositories import document_registry_repository as repo


def today_msk() -> date:
    return datetime.now(ZoneInfo("Europe/Moscow")).date()


async def resolve_actor(session: AsyncSession, user: dict[str, Any], notification_company_id: UUID | None = None) -> dict[str, Any]:
    company_id = notification_company_id or user.get("company_id")
    if notification_company_id and user["role"] == "carcraft_employee":
        raise ServiceError("Администратор не выбирает компанию уведомления", 403)
    actor = await repo.resolve_actor(session, UUID(str(user["id"])), user["role"], UUID(str(company_id)) if company_id else None)
    if actor is None:
        raise ServiceError("Нет доступа к выбранной компании", 403)
    return actor


async def get_document(session: AsyncSession, document_id: UUID, actor: dict[str, Any]) -> dict[str, Any]:
    document = await repo.get_document(session, document_id, actor, today=today_msk())
    if document is None:
        raise ServiceError("Документ не найден", 404)
    return document


async def history(session: AsyncSession, document_id: UUID, actor: dict[str, Any]) -> dict[str, Any]:
    await get_document(session, document_id, actor)
    versions = await repo.version_projection(session, [document_id])
    return {"items": versions.get(document_id, [])}


def _public_group(group: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in group.items() if key not in ("documents", "created_at")}


async def list_groups(session: AsyncSession, actor: dict[str, Any], filters: dict[str, Any], page: int, page_size: int) -> dict[str, Any]:
    groups, total = await repo.list_groups(session, actor, filters, today_msk(), offset=(page - 1) * page_size, limit=page_size)
    return {"items": [_public_group(group) for group in groups], "pagination": {"page": page, "page_size": page_size, "total": total}}


async def group_documents(session: AsyncSession, group_id: UUID, actor: dict[str, Any], document_type: str | None, page: int, page_size: int) -> dict[str, Any]:
    documents = await repo.get_group_documents(session, group_id, actor, today_msk())
    if not documents:
        raise ServiceError("Группа не найдена", 404)
    if document_type:
        documents = [document for document in documents if document["document_type"] == document_type]
    return {"items": documents[(page - 1) * page_size:page * page_size], "pagination": {"page": page, "page_size": page_size, "total": len(documents)}}


def company_labels(value: dict[str, Any]) -> list[str]:
    labels = ["Платформа МЛ"] if value["platform_ml"] else []
    roles = {"leasing_company": "ЛК", "distributor": "Дистрибьютор", "dealer": "Дилер"}
    for role, label in roles.items():
        labels.extend(dict.fromkeys(f"{label} {company['name']}" for company in value["companies"] if company["role"] == role))
    return labels


def table_row(group: dict[str, Any]) -> dict[str, Any]:
    document = group["display_document"]
    linked = []
    for child in group["documents"]:
        if child["id"] == document["id"]:
            continue
        companies = company_labels(group["participants"]) or company_labels(child["related_companies"]) or ["Без компании"]
        linked.append({"id": child["id"], "document_type": child["document_type"], "contract_number": child["contract_number"], "label": f"{TYPES[child['document_type']]} + {child['contract_number']} + {', '.join(companies)}", "url": f"/workspace/document-registry?document={child['id']}"})
    return {**_public_group(group), **document, "related_documents": linked}


def _cursor_parts(value: Any) -> tuple[str, str]:
    if not isinstance(value, list) or len(value) != 2 or not all(isinstance(part, str) for part in value):
        raise ValueError("Invalid cursor shape")
    return value[0], value[1]


def _decode_cursor(cursor: str | None) -> tuple[datetime, UUID] | None:
    if not cursor:
        return None
    try:
        value = _cursor_parts(json.loads(base64.urlsafe_b64decode(cursor.encode())))
        timestamp = datetime.fromisoformat(value[0])
        identifier = UUID(value[1])
    except (ValueError, TypeError, IndexError, UnicodeError) as exc:
        raise ServiceError("Некорректный указатель страницы", 422) from exc
    if timestamp.tzinfo is None:
        raise ServiceError("Некорректный указатель страницы", 422)
    return timestamp, identifier


async def table(session: AsyncSession, actor: dict[str, Any], filters: dict[str, Any], cursor: str | None, limit: int) -> dict[str, Any]:
    groups, _ = await repo.list_groups(session, actor, filters, today_msk(), limit=limit + 1, cursor=_decode_cursor(cursor))
    has_more = len(groups) > limit
    groups = groups[:limit]
    next_cursor = None
    if has_more and groups:
        last = groups[-1]
        next_cursor = base64.urlsafe_b64encode(json.dumps([last["created_at"].isoformat(), str(last["group_id"])]).encode()).decode()
    return {"items": [table_row(group) for group in groups], "pagination": {"has_more": has_more, "next_cursor": next_cursor, "limit": limit}}


def _xlsx(rows: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Справочник документов"
    sheet.append(["Тип документа", "Номер", "Название", "Участники", "Связанные компании", "Связанные документы", "Марка", "Модель", "Действует с", "Действует по", "Статус", "Кол-во файлов"])
    for row in rows:
        version = row["current_version"]
        values = [TYPES[row["document_type"]], row["contract_number"], row["name"], "\n".join(company_labels(row["participants"])), "\n".join(company_labels(row["related_companies"])), " + ".join(document["label"] for document in row["related_documents"]), row["participants"]["mark_name"], row["participants"]["model_name"], version["valid_from"].strftime("%d.%m.%Y"), version["valid_to"].strftime("%d.%m.%Y") if version["valid_to"] else "—", STATUS_LABELS[row["status"]], str(len(version["files"]))]
        sheet.append([str(value) if value else "—" for value in values])
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for column in ("A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L"):
        sheet.column_dimensions[column].width = 36 if column in ("C", "D", "E", "F") else 20
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


async def export(session: AsyncSession, actor: dict[str, Any], filters: dict[str, Any]) -> bytes:
    from asyncio import to_thread
    groups, _ = await repo.list_groups(session, actor, filters, today_msk(), limit=None)
    return await to_thread(_xlsx, [table_row(group) for group in groups])


async def candidates(session: AsyncSession, actor: dict[str, Any], context: dict[str, Any]) -> dict[str, Any]:
    try:
        groups = await repo.candidate_groups(session, actor, context, today_msk())
        return {"items": [document for group in groups for document in group["documents"] if document["status"] in ("active", "expiring", "pending")], "groups": [_public_group(group) for group in groups]}
    except ValueError as exc:
        raise ServiceError(str(exc), 409) from exc


async def company_lookup(session: AsyncSession, role: str, search: str) -> dict[str, Any]:
    return {"items": await repo.lookup_companies(session, role, search)}


async def catalog_lookup(session: AsyncSession, fields: str, mark_id: UUID | None) -> dict[str, Any]:
    if fields == "models" and mark_id is None:
        raise ServiceError("Выберите марку", 422)
    return {"items": await repo.lookup_catalog(session, fields, mark_id)}


async def check_number(session: AsyncSession, number: str) -> dict[str, bool]:
    return {"available": await repo.number_available(session, number)}


async def deletion_usages(session: AsyncSession, document_id: UUID) -> dict[str, Any]:
    usages = await repo.deletion_usages(session, document_id)
    if usages is None:
        raise ServiceError("Документ не найден", 404)
    fingerprint = hashlib.sha256(json.dumps(usages, sort_keys=True, default=str).encode()).hexdigest()
    return {"document_ids": usages["document_ids"], "programs": usages["programs"], "fingerprint": fingerprint}
