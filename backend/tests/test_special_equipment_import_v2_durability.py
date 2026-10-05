from __future__ import annotations

import base64
import hashlib
from contextlib import AbstractAsyncContextManager
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, ClassVar, cast
from unittest.mock import AsyncMock
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.dialects import postgresql
from sqlalchemy.exc import IntegrityError
from sqlalchemy.sql.dml import Insert

from application import special_equipment_import as import_service
from application import special_equipment_import_v2 as import_v2
from application.errors import ServiceError
from application.tasks import special_equipment_import as import_tasks
from domain.special_equipment_import import (
    ImportAggregateSemanticConflictError,
    ImportContractError,
    ImportErrorPolicy,
    ImportMode,
    ImportStatus,
)
from infrastructure.models.special_equipment import (
    SpecialEquipmentAttribute,
    SpecialEquipmentAttributeGroup,
    SpecialEquipmentCategory,
    SpecialEquipmentCategoryAttribute,
    SpecialEquipmentColor,
    SpecialEquipmentModificationAttributeValue,
    SpecialEquipmentProduct,
    SpecialEquipmentProductAttachment,
    SpecialEquipmentProductImage,
    SpecialEquipmentTrim,
    SpecialEquipmentTrimAttribute,
    SpecialEquipmentTrimAttributeValue,
)
from infrastructure.repositories import (
    special_equipment_import_repository as import_repository,
)
from infrastructure.services.special_equipment_import_images import TransferredImage
from infrastructure.services.special_equipment_xlsx import (
    JsonlRows,
    _normalize_cell,
    build_template_v7,
    create_row_stores,
    parse_xlsx,
)
from tests.special_equipment_factories import special_equipment_directory


def test_v2_apply_drops_fields_absent_from_target_table() -> None:
    attribute_id = uuid4()

    values = import_repository._v2_coerce_values(
        SpecialEquipmentAttribute.__table__,
        {
            "id": attribute_id,
            "name": "Грузоподъёмность",
            "slug": "gruzopodemnost",
        },
    )

    assert values == {
        "id": attribute_id,
        "name": "Грузоподъёмность",
    }


@pytest.mark.asyncio
async def test_v4_apply_persists_trim_color_values_and_product_assignments(
    db_session: Any,
) -> None:
    mark, model, modification = special_equipment_directory(
        mark_name="Import v4 mark",
        model_name="Import v4 model",
        modification_name="Import v4 modification",
    )
    group = SpecialEquipmentAttributeGroup(
        code=f"group-{uuid4().hex}",
        name=f"Import v4 group {uuid4().hex}",
        slug=f"group-{uuid4().hex}",
        sort_order=0,
    )
    attribute = SpecialEquipmentAttribute(
        code=f"power-{uuid4().hex}",
        name="Import v4 power",
        attribute_group_id=group.id,
        data_type="number",
        filter_kind="range",
    )
    db_session.add_all([mark, model, modification, group, attribute])
    await db_session.flush()

    trim_id = uuid4()
    body_color_id = uuid4()
    interior_color_id = uuid4()
    product_id = uuid4()
    trim_key = f"{modification.code}:premium"
    plan = {
        "trims": [
            {
                "id": trim_id,
                "code": "premium",
                "operation": "ADD",
                "values": {
                    "name": "Premium",
                    "slug": "premium",
                    "modification_id": modification.id,
                    "sort_order": 0,
                    "is_active": True,
                },
                "_aggregate_kind": "trim",
                "_aggregate_code": trim_key,
                "_sheet_code": "Комплектации",
                "_row_number": 2,
            }
        ],
        "colors": [
            {
                "id": body_color_id,
                "code": "black",
                "operation": "ADD",
                "values": {
                    "name": "Чёрный",
                    "applicability": "body",
                    "is_active": True,
                },
                "_aggregate_kind": "color",
                "_aggregate_code": "black",
                "_sheet_code": "Цвета",
                "_row_number": 2,
            },
            {
                "id": interior_color_id,
                "code": "gray",
                "operation": "ADD",
                "values": {
                    "name": "Серый",
                    "applicability": "interior",
                    "is_active": True,
                },
                "_aggregate_kind": "color",
                "_aggregate_code": "gray",
                "_sheet_code": "Цвета",
                "_row_number": 3,
            },
        ],
        "trim_attributes": [
            {
                "operation": "ADD",
                "values": {
                    "trim_id": trim_id,
                    "attribute_id": attribute.id,
                    "group_id": group.id,
                    "is_required": True,
                    "is_filterable": True,
                    "sort_order": 0,
                },
                "_aggregate_kind": "trim",
                "_aggregate_code": trim_key,
                "_sheet_code": "Характеристики комплектаций",
                "_row_number": 2,
            }
        ],
        "trim_attribute_values": [
            {
                "operation": "SET",
                "values": {
                    "trim_id": trim_id,
                    "attribute_id": attribute.id,
                    "value_number": Decimal("350"),
                    "value_text": None,
                    "value_boolean": None,
                    "option_id": None,
                },
                "_aggregate_kind": "trim",
                "_aggregate_code": trim_key,
                "_sheet_code": "Значения комплектаций",
                "_row_number": 2,
            }
        ],
        "products": [
            {
                "id": product_id,
                "code": f"truck-{uuid4().hex}",
                "operation": "ADD",
                "values": {
                    "slug": f"truck-{uuid4().hex}",
                    "modification_id": modification.id,
                    "trim_id": trim_id,
                    "body_color_id": body_color_id,
                    "interior_color_id": interior_color_id,
                    "condition": "new",
                    "no_vin": True,
                    "price": Decimal("1000000"),
                    "currency_code": "RUB",
                    "publication_status": "draft",
                    "sale_status": "available",
                },
                "_aggregate_kind": "product",
                "_aggregate_code": "truck-v4",
                "_sheet_code": "Объявления",
                "_row_number": 2,
            }
        ],
    }

    result = await import_repository.apply_normalized_plan(
        db_session,
        job_id=uuid4(),
        mode=ImportMode.APPEND.value,
        plan=plan,
    )

    product = await db_session.get(SpecialEquipmentProduct, product_id)
    value = await db_session.get(
        SpecialEquipmentTrimAttributeValue,
        (trim_id, attribute.id),
    )
    assert result["rejected_aggregates"] == []
    assert product is not None
    assert (
        product.trim_id,
        product.body_color_id,
        product.interior_color_id,
    ) == (trim_id, body_color_id, interior_color_id)
    assert value is not None
    assert value.value_number == Decimal("350")


@pytest.mark.asyncio
async def test_v4_full_snapshot_archives_missing_trims_colors_and_values(
    db_session: Any,
) -> None:
    mark, model, modification = special_equipment_directory(
        mark_name="Snapshot v4 mark",
        model_name="Snapshot v4 model",
        modification_name="Snapshot v4 modification",
    )
    group = SpecialEquipmentAttributeGroup(
        code=f"snapshot-group-{uuid4().hex}",
        name=f"Snapshot v4 group {uuid4().hex}",
        slug=f"snapshot-group-{uuid4().hex}",
        sort_order=0,
    )
    attribute = SpecialEquipmentAttribute(
        code=f"snapshot-power-{uuid4().hex}",
        name="Snapshot v4 power",
        attribute_group_id=group.id,
        data_type="number",
        filter_kind="range",
    )
    keep_trim = SpecialEquipmentTrim(
        modification_id=modification.id,
        code="keep",
        name="Keep",
        slug="keep",
        sort_order=0,
    )
    missing_trim = SpecialEquipmentTrim(
        modification_id=modification.id,
        code="missing",
        name="Missing",
        slug="missing",
        sort_order=1,
    )
    keep_color = SpecialEquipmentColor(
        code=f"keep-{uuid4().hex}",
        name=f"Keep {uuid4().hex}",
        applicability="both",
    )
    missing_color = SpecialEquipmentColor(
        code=f"missing-{uuid4().hex}",
        name=f"Missing {uuid4().hex}",
        applicability="both",
    )
    db_session.add_all(
        [
            mark,
            model,
            modification,
            group,
            attribute,
            keep_trim,
            missing_trim,
            keep_color,
            missing_color,
        ]
    )
    await db_session.flush()
    missing_link = SpecialEquipmentTrimAttribute(
        trim_id=missing_trim.id,
        attribute_id=attribute.id,
        group_id=group.id,
        is_required=True,
        is_filterable=True,
        sort_order=0,
    )
    missing_value = SpecialEquipmentTrimAttributeValue(
        trim_id=missing_trim.id,
        attribute_id=attribute.id,
        value_number=Decimal("1"),
    )
    db_session.add_all([missing_link, missing_value])
    await db_session.flush()
    counts = {"created": 0, "updated": 0, "archived": 0, "removed": 0}

    await import_repository._v2_delete_missing_snapshot_rows(
        db_session,
        plan={},
        counts=counts,
        include_v4_catalog=False,
    )
    await db_session.refresh(missing_trim)
    await db_session.refresh(missing_color)
    assert missing_trim.is_active is True
    assert missing_color.is_active is True
    assert (
        await db_session.scalar(
            select(SpecialEquipmentTrimAttributeValue).where(
                SpecialEquipmentTrimAttributeValue.trim_id == missing_trim.id,
                SpecialEquipmentTrimAttributeValue.attribute_id == attribute.id,
            )
        )
        is not None
    )

    await import_repository._v2_delete_missing_snapshot_rows(
        db_session,
        plan={
            "trims": [
                {"id": keep_trim.id, "operation": "SET"},
            ],
            "colors": [
                {"id": keep_color.id, "operation": "SET"},
            ],
        },
        counts=counts,
    )

    await db_session.refresh(keep_trim)
    await db_session.refresh(missing_trim)
    await db_session.refresh(keep_color)
    await db_session.refresh(missing_color)
    assert keep_trim.is_active is True
    assert missing_trim.is_active is False
    assert keep_color.is_active is True
    assert missing_color.is_active is False
    assert (
        await db_session.scalar(
            select(SpecialEquipmentTrimAttribute).where(
                SpecialEquipmentTrimAttribute.trim_id == missing_trim.id,
                SpecialEquipmentTrimAttribute.attribute_id == attribute.id,
            )
        )
        is None
    )
    assert (
        await db_session.scalar(
            select(SpecialEquipmentTrimAttributeValue).where(
                SpecialEquipmentTrimAttributeValue.trim_id == missing_trim.id,
                SpecialEquipmentTrimAttributeValue.attribute_id == attribute.id,
            )
        )
        is None
    )


@pytest.mark.asyncio
async def test_v2_apply_full_snapshot_deletes_conflicting_modification_value_before_applying_trim_value(
    db_session: Any,
) -> None:
    mark, model, modification = special_equipment_directory(
        mark_name=f"Mark {uuid4().hex[:6]}",
        model_name=f"Model {uuid4().hex[:6]}",
        modification_name=f"Modification {uuid4().hex[:6]}",
    )
    group = SpecialEquipmentAttributeGroup(
        code=f"group-{uuid4().hex[:6]}",
        name="Group",
        slug=f"group-{uuid4().hex[:6]}",
        sort_order=0,
    )
    attribute = SpecialEquipmentAttribute(
        code=f"power-{uuid4().hex[:6]}",
        name="Power",
        attribute_group_id=group.id,
        data_type="number",
        filter_kind="range",
    )
    trim = SpecialEquipmentTrim(
        modification_id=modification.id,
        code=f"trim-{uuid4().hex[:6]}",
        name="Trim",
        slug=f"trim-{uuid4().hex[:6]}",
        sort_order=0,
    )
    db_session.add_all([mark, model, modification, group, attribute, trim])
    await db_session.flush()

    # S0 has modification attribute value
    mod_value = SpecialEquipmentModificationAttributeValue(
        modification_id=modification.id,
        attribute_id=attribute.id,
        value_number=Decimal("100"),
    )
    db_session.add(mod_value)
    await db_session.flush()

    # FULL_SNAPSHOT plan does NOT have modification attribute values,
    # but introduces trim attribute value for the same attribute.
    trim_key = f"{modification.code}:{trim.code}"
    plan: dict[str, list[dict[str, Any]]] = {
        "marks": [
            {
                "id": mark.id,
                "code": mark.code,
                "operation": "SET",
                "values": {"name": mark.name, "slug": mark.slug, "is_active": True},
                "_aggregate_kind": "mark",
                "_aggregate_code": mark.code,
            }
        ],
        "models": [
            {
                "id": model.id,
                "code": model.code,
                "operation": "SET",
                "values": {"name": model.name, "slug": model.slug, "mark_id": mark.id, "is_active": True},
                "_aggregate_kind": "model",
                "_aggregate_code": model.code,
            }
        ],
        "modifications": [
            {
                "id": modification.id,
                "code": modification.code,
                "operation": "SET",
                "values": {"name": modification.name, "slug": modification.slug, "model_id": model.id, "is_active": True},
                "_aggregate_kind": "modification",
                "_aggregate_code": modification.code,
            }
        ],
        "trims": [
            {
                "id": trim.id,
                "code": trim.code,
                "operation": "SET",
                "values": {"name": trim.name, "slug": trim.slug, "modification_id": modification.id, "sort_order": 0, "is_active": True},
                "_aggregate_kind": "trim",
                "_aggregate_code": trim_key,
            }
        ],
        "attribute_groups": [
            {
                "id": group.id,
                "code": group.code,
                "operation": "SET",
                "values": {"name": group.name, "slug": group.slug, "sort_order": 0, "is_active": True},
                "_aggregate_kind": "attribute_group",
                "_aggregate_code": group.code,
            }
        ],
        "attributes": [
            {
                "id": attribute.id,
                "code": attribute.code,
                "operation": "SET",
                "values": {"name": attribute.name, "attribute_group_id": group.id, "data_type": "number", "filter_kind": "range", "is_active": True},
                "_aggregate_kind": "attribute",
                "_aggregate_code": attribute.code,
            }
        ],
        "trim_attributes": [
            {
                "operation": "SET",
                "values": {
                    "trim_id": trim.id,
                    "attribute_id": attribute.id,
                    "group_id": group.id,
                    "is_required": False,
                    "is_filterable": False,
                    "sort_order": 0,
                },
                "_aggregate_kind": "trim",
                "_aggregate_code": trim_key,
            }
        ],
        "trim_attribute_values": [
            {
                "operation": "SET",
                "values": {
                    "trim_id": trim.id,
                    "attribute_id": attribute.id,
                    "value_number": Decimal("200"),
                    "value_text": None,
                    "value_boolean": None,
                    "option_id": None,
                },
                "_aggregate_kind": "trim",
                "_aggregate_code": trim_key,
            }
        ],
    }

    result = await import_repository.apply_normalized_plan(
        db_session,
        job_id=uuid4(),
        mode=ImportMode.FULL_SNAPSHOT.value,
        plan=plan,
    )

    assert result["rejected_aggregates"] == []
    mod_id = modification.id
    attr_id = attribute.id
    trim_id = trim.id
    db_session.expire_all()
    # Old modification value in S0 must be deleted
    assert (
        await db_session.get(
            SpecialEquipmentModificationAttributeValue,
            (mod_id, attr_id),
        )
        is None
    )
    # New trim value must be persisted
    trim_val = await db_session.get(
        SpecialEquipmentTrimAttributeValue,
        (trim_id, attr_id),
    )
    assert trim_val is not None
    assert trim_val.value_number == Decimal("200")


@pytest.mark.asyncio
async def test_v2_apply_full_snapshot_explicit_attribute_deactivation_not_blocked_by_s0_category_attributes(
    db_session: Any,
) -> None:
    group = SpecialEquipmentAttributeGroup(
        code=f"group-{uuid4().hex[:6]}",
        name="Group",
        slug=f"group-{uuid4().hex[:6]}",
        sort_order=0,
    )
    attribute = SpecialEquipmentAttribute(
        code=f"attr-{uuid4().hex[:6]}",
        name="Attribute",
        attribute_group_id=group.id,
        data_type="number",
        filter_kind="range",
        is_active=True,
    )
    category = SpecialEquipmentCategory(
        code=f"cat-{uuid4().hex[:6]}",
        name="Category",
        slug=f"cat-{uuid4().hex[:6]}",
        is_active=True,
        sort_order=0,
        usage_metric="mileage_km",
    )
    db_session.add_all([group, attribute, category])
    await db_session.flush()

    # S0 has category_attribute link
    cat_attr = SpecialEquipmentCategoryAttribute(
        category_id=category.id,
        attribute_id=attribute.id,
        group_id=group.id,
        is_required=False,
        is_filterable=False,
        is_visible=True,
        sort_order=0,
    )
    db_session.add(cat_attr)
    await db_session.flush()

    # FULL_SNAPSHOT plan explicitly deactivates attribute and does not include category_attributes
    plan: dict[str, list[dict[str, Any]]] = {
        "attribute_groups": [
            {
                "id": group.id,
                "code": group.code,
                "operation": "SET",
                "values": {"name": group.name, "slug": group.slug, "sort_order": 0, "is_active": True},
                "_aggregate_kind": "attribute_group",
                "_aggregate_code": group.code,
            }
        ],
        "categories": [
            {
                "id": category.id,
                "code": category.code,
                "operation": "SET",
                "values": {"name": category.name, "slug": category.slug, "is_active": True, "sort_order": 0, "usage_metric": "mileage_km"},
                "_aggregate_kind": "category",
                "_aggregate_code": category.code,
            }
        ],
        "attributes": [
            {
                "id": attribute.id,
                "code": attribute.code,
                "operation": "SET",
                "values": {"name": attribute.name, "attribute_group_id": group.id, "data_type": "number", "filter_kind": "range", "is_active": False},
                "_aggregate_kind": "attribute",
                "_aggregate_code": attribute.code,
            }
        ],
    }

    result = await import_repository.apply_normalized_plan(
        db_session,
        job_id=uuid4(),
        mode=ImportMode.FULL_SNAPSHOT.value,
        plan=plan,
    )

    assert result["rejected_aggregates"] == []
    cat_id = category.id
    attr_id = attribute.id
    db_session.expire_all()
    await db_session.refresh(attribute)
    assert attribute.is_active is False
    assert (
        await db_session.get(
            SpecialEquipmentCategoryAttribute,
            (cat_id, attr_id),
        )
        is None
    )


@pytest.mark.asyncio
async def test_v3_image_transfer_replaces_private_urls_with_staged_keys(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    product_id = uuid4()
    plan: dict[str, JsonlRows] = create_row_stores(
        tmp_path / "image-plan",
        import_tasks.DATA_SHEET_HEADERS,
    )
    plan["products"].append(
        {
            "id": product_id,
            "code": "truck-1",
            "operation": "SET",
            "values": {
                "_image_action": "replace",
                "_image_source_url": "https://media.example.test/secret.png?token=x",
            },
            "_sheet_code": "Объявления",
            "_row_number": 2,
            "_aggregate_kind": "product",
            "_aggregate_code": "truck-1",
        }
    )

    async def transfer(**kwargs: Any) -> TransferredImage:
        assert kwargs["source_url"].startswith("https://media.example.test/")
        return TransferredImage(
            storage_key="special-equipment/staged/job/products/image.webp",
            content_sha256="a" * 64,
            size_bytes=123,
        )

    monkeypatch.setattr(import_tasks, "transfer_temporary_image", transfer)
    issues = import_v2.V2IssueCollector()
    summary = await import_tasks._transfer_plan_images(
        job_id=uuid4(),
        mode=ImportMode.PATCH,
        plan=plan,
        issues=issues,
        storage=object(),  # type: ignore[arg-type]
    )

    values = next(iter(plan["products"]))["values"]
    assert values == {
        "_primary_image_action": "replace",
        "_primary_image_storage_key": (
            "special-equipment/staged/job/products/image.webp"
        ),
    }
    assert summary == {
        "requested": 1,
        "transferred": 1,
        "reused": 0,
        "optionalFailures": 0,
        "blockingFailures": 0,
    }
    assert issues.error_count == 0


@pytest.mark.asyncio
async def test_v3_apply_replaces_then_clears_only_primary_product_image(
    db_session: Any,
) -> None:
    mark, model, modification = special_equipment_directory(
        mark_name="Import image mark",
        model_name="Import image model",
        modification_name="Import image modification",
    )
    db_session.add_all([mark, model, modification])
    await db_session.flush()
    product = SpecialEquipmentProduct(
        code=f"image-import-{uuid4()}",
        modification_id=modification.id,
        slug=f"image-import-{uuid4()}",
        condition="new",
        no_vin=True,
        price=100,
        currency_code="RUB",
        publication_status="draft",
        sale_status="available",
    )
    db_session.add(product)
    await db_session.flush()
    old_primary = SpecialEquipmentProductImage(
        product_id=product.id,
        storage_key=f"old-primary-{uuid4()}.webp",
        sort_order=0,
        is_primary=True,
    )
    additional = SpecialEquipmentProductImage(
        product_id=product.id,
        storage_key=f"additional-{uuid4()}.webp",
        sort_order=1,
        is_primary=False,
    )
    db_session.add_all([old_primary, additional])
    await db_session.flush()

    staged_key = f"special-equipment/staged/{uuid4()}/primary.webp"
    await import_repository._v2_apply_product_primary_image(
        db_session,
        product_id=product.id,
        values={
            "_primary_image_action": "replace",
            "_primary_image_storage_key": staged_key,
        },
    )
    await db_session.flush()

    replaced = list(
        (
            await db_session.execute(
                select(SpecialEquipmentProductImage)
                .where(SpecialEquipmentProductImage.product_id == product.id)
                .order_by(SpecialEquipmentProductImage.sort_order)
            )
        ).scalars()
    )
    assert [(item.storage_key, item.is_primary) for item in replaced] == [
        (staged_key, True),
        (additional.storage_key, False),
    ]

    await import_repository._v2_apply_product_primary_image(
        db_session,
        product_id=product.id,
        values={"_primary_image_action": "clear"},
    )
    await db_session.flush()
    cleared = list(
        (
            await db_session.execute(
                select(SpecialEquipmentProductImage).where(
                    SpecialEquipmentProductImage.product_id == product.id
                )
            )
        ).scalars()
    )
    assert [
        (item.storage_key, item.is_primary, item.sort_order) for item in cleared
    ] == [(additional.storage_key, False, 0)]


@pytest.mark.asyncio
async def test_relation_replacement_reinserts_mixed_existing_and_new_rows() -> None:
    owner_id = uuid4()
    existing_target_id = uuid4()
    new_target_id = uuid4()
    created_by = uuid4()
    created_at = datetime.now(UTC)

    class MappingResult:
        def mappings(self) -> MappingResult:
            return self

        def all(self) -> list[dict[str, Any]]:
            return [
                {
                    "product_id": owner_id,
                    "attachment_product_id": existing_target_id,
                    "position": 0,
                    "created_by": created_by,
                    "created_at": created_at,
                }
            ]

    class CompileSession:
        def __init__(self) -> None:
            self.statements: list[Any] = []

        async def execute(self, statement: Any) -> MappingResult:
            statement.compile(dialect=postgresql.dialect())
            self.statements.append(statement)
            return MappingResult()

    session = CompileSession()
    counts = {"created": 0, "updated": 0, "removed": 0}
    await import_repository._v2_replace_product_relation_set(
        session,  # type: ignore[arg-type]
        family="product_attachments",
        rows=[
            {
                "operation": "SET",
                "values": {
                    "product_id": owner_id,
                    "attachment_product_id": existing_target_id,
                    "position": 2,
                },
            },
            {
                "operation": "ADD",
                "values": {
                    "product_id": owner_id,
                    "attachment_product_id": new_target_id,
                    "position": 3,
                },
            },
        ],
        counts=counts,
        snapshot=False,
    )

    assert counts == {"created": 1, "updated": 1, "removed": 0}
    inserts = [
        statement
        for statement in session.statements
        if isinstance(statement, Insert)
        and statement.table is SpecialEquipmentProductAttachment.__table__
    ]
    assert len(inserts) == 2


def _job(**overrides: Any) -> dict[str, Any]:
    now = datetime.now(UTC)
    values: dict[str, Any] = {
        "id": uuid4(),
        "requested_by": uuid4(),
        "request_hash": "request-hash",
        "original_filename": "catalog.xlsx",
        "mode": "PATCH",
        "error_policy": "BEST_EFFORT",
        "template_version": 2,
        "status": "awaiting_upload",
        "phase": "awaiting_upload",
        "actual_size_bytes": 0,
        "expected_size_bytes": 1024,
        "rows_done": 0,
        "rows_total": None,
        "entities_done": 0,
        "entities_total": None,
        "images_done": 0,
        "images_total": 0,
        "issues_total": 0,
        "summary": {},
        "preview_hash": None,
        "catalog_revision": None,
        "applied_revision": None,
        "error_code": None,
        "error_detail": None,
        "created_at": now,
        "updated_at": now,
        "source_object_key": "imports/source.xlsx",
        "multipart_upload_id": "upload-1",
        "upload_expires_at": now.replace(year=now.year + 1),
        "cancellation_requested": False,
    }
    values.update(overrides)
    return values


class _Storage:
    def __init__(self) -> None:
        self.initiated = 0
        self.aborted: list[tuple[str, str]] = []

    async def initiate_multipart(self, _key: str, _content_type: str) -> str:
        self.initiated += 1
        return f"upload-{self.initiated}"

    async def abort_multipart(self, *, key: str, upload_id: str) -> None:
        self.aborted.append((key, upload_id))


class _TaskSessionContext:
    async def __aenter__(self) -> object:
        return object()

    async def __aexit__(
        self,
        exc_type: type[BaseException] | None,
        exc: BaseException | None,
        traceback: Any,
    ) -> bool:
        return False


class _ArtifactStorage:
    async def download_to_path(self, *_args: Any, **_kwargs: Any) -> str:
        return "a" * 64


class _WorkerSession:
    def __init__(self) -> None:
        self.commits = 0

    async def commit(self) -> None:
        self.commits += 1


class _WorkerSessionContext:
    sessions: ClassVar[list[_WorkerSession]] = []

    async def __aenter__(self) -> _WorkerSession:
        session = _WorkerSession()
        self.sessions.append(session)
        return session

    async def __aexit__(
        self,
        exc_type: type[BaseException] | None,
        exc: BaseException | None,
        traceback: Any,
    ) -> bool:
        return False


def _trim_conflict_plan() -> tuple[dict[str, list[dict[str, Any]]], UUID, UUID, UUID]:
    modification_id = uuid4()
    trim_id = uuid4()
    attribute_id = uuid4()
    return (
        {
            "modification_attribute_values": [
                {
                    "operation": "SET",
                    "values": {
                        "modification_id": modification_id,
                        "attribute_id": attribute_id,
                    },
                    "_sheet_code": "Характеристики модификаций",
                    "_row_number": 9,
                    "_aggregate_kind": "modification",
                    "_aggregate_code": "modification-race",
                }
            ]
        },
        modification_id,
        trim_id,
        attribute_id,
    )


def _configure_worker_trim_race(
    monkeypatch: pytest.MonkeyPatch,
    *,
    job: dict[str, Any],
    plan: dict[str, list[dict[str, Any]]],
    modification_id: UUID,
    trim_id: UUID,
    attribute_id: UUID,
    expected_modification_value_markers: frozenset[tuple[UUID, UUID]] = frozenset(),
    expected_trim_value_markers: (
        frozenset[tuple[UUID, UUID, UUID]]
    ) = frozenset(),
) -> None:
    async def get_job(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        return job

    async def current_revision(*_args: Any, **_kwargs: Any) -> int:
        return int(job["catalog_revision"]) + 1

    async def matches(*_args: Any, **_kwargs: Any) -> bool:
        return True

    async def value_state(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        return {
            "modification_by_trim": {trim_id: modification_id},
            "trim_values": {(trim_id, modification_id, attribute_id)},
            "modification_values": set(),
        }

    _WorkerSessionContext.sessions = []
    monkeypatch.setattr(import_tasks, "AsyncSessionLocal", _WorkerSessionContext)
    monkeypatch.setattr(import_tasks.repo, "get_job", get_job)
    monkeypatch.setattr(import_tasks.repo, "get_catalog_revision", current_revision)
    monkeypatch.setattr(import_tasks.repo, "lock_expected_product_versions", matches)
    monkeypatch.setattr(import_tasks.repo, "lock_active_seller_companies", matches)
    monkeypatch.setattr(import_tasks.repo, "lock_import_archive_targets", matches)
    monkeypatch.setattr(
        import_tasks.repo,
        "lock_and_read_trim_modification_value_state",
        value_state,
    )
    monkeypatch.setattr(
        import_tasks,
        "get_special_equipment_import_storage",
        _ArtifactStorage,
    )
    monkeypatch.setattr(
        import_tasks,
        "read_rows_archive",
        lambda *_args, **_kwargs: ({}, plan),
    )
    monkeypatch.setattr(
        import_tasks,
        "_validate_normalized_artifact",
        lambda **_kwargs: (
            {},
            set(),
            set(expected_modification_value_markers),
            set(expected_trim_value_markers),
        ),
    )
    monkeypatch.setattr(
        import_tasks,
        "_artifact_matches_locked_job",
        lambda *_args, **_kwargs: True,
    )


@pytest.mark.asyncio
async def test_worker_atomic_trim_race_finishes_with_semantic_code(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan, modification_id, trim_id, attribute_id = _trim_conflict_plan()
    job = _job(
        status=ImportStatus.APPLYING.value,
        normalized_artifact_key="imports/normalized-race.zip",
        mode=ImportMode.FULL_SNAPSHOT.value,
        error_policy=ImportErrorPolicy.ATOMIC.value,
        template_version=4,
        catalog_revision=7,
        source_code="catalog_v2",
    )
    marked: dict[str, Any] = {}

    async def mark_failed(job_id: UUID, code: str, detail: str) -> None:
        marked.update(job_id=job_id, code=code, detail=detail)

    _configure_worker_trim_race(
        monkeypatch,
        job=job,
        plan=plan,
        modification_id=modification_id,
        trim_id=trim_id,
        attribute_id=attribute_id,
    )
    monkeypatch.setattr(import_tasks, "_mark_failed", mark_failed)

    await import_tasks.apply_special_equipment_import.original_func(str(job["id"]))

    assert marked == {
        "job_id": job["id"],
        "code": "ATTRIBUTE_ALREADY_ASSIGNED_IN_TRIM",
        "detail": "Характеристика уже заполнена в комплектации модификации",
    }


@pytest.mark.asyncio
async def test_worker_best_effort_trim_race_rejects_only_conflicting_row(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan, modification_id, trim_id, attribute_id = _trim_conflict_plan()
    job = _job(
        status=ImportStatus.APPLYING.value,
        normalized_artifact_key="imports/normalized-race.zip",
        mode=ImportMode.PATCH.value,
        error_policy=ImportErrorPolicy.BEST_EFFORT.value,
        template_version=4,
        catalog_revision=7,
        source_code="catalog_v2",
    )
    issue = {
        "sheet_code": "Характеристики модификаций",
        "row_number": 9,
        "column_name": None,
        "severity": "error",
        "code": "ATTRIBUTE_ALREADY_ASSIGNED_IN_TRIM",
        "message": "Характеристика уже заполнена в комплектации модификации",
        "raw_value_preview": None,
        "entity_type": "modification",
        "external_key": "modification-race",
    }
    updates: list[dict[str, Any]] = []
    appended: list[dict[str, Any]] = []

    async def apply_plan(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        return {
            "counts": {"created": 0, "updated": 0, "archived": 0, "removed": 0},
            "rejected_aggregates": [dict(issue)],
        }

    async def no_issues(*_args: Any, **_kwargs: Any) -> list[dict[str, Any]]:
        return []

    async def append_issues(*_args: Any, issues: list[dict[str, Any]], **_kwargs: Any) -> None:
        appended.extend(issues)

    async def no_op(*_args: Any, **_kwargs: Any) -> None:
        return None

    async def increment(*_args: Any, **_kwargs: Any) -> int:
        return 9

    async def update(*_args: Any, **values: Any) -> dict[str, Any]:
        updates.append(values)
        return {**job, **values}

    _configure_worker_trim_race(
        monkeypatch,
        job=job,
        plan=plan,
        modification_id=modification_id,
        trim_id=trim_id,
        attribute_id=attribute_id,
    )
    monkeypatch.setattr(import_tasks.repo, "apply_normalized_plan", apply_plan)
    monkeypatch.setattr(import_tasks.repo, "list_issues", no_issues)
    monkeypatch.setattr(import_tasks.repo, "append_issues", append_issues)
    monkeypatch.setattr(
        import_tasks.repo,
        "mark_job_staged_media_referenced",
        no_op,
    )
    monkeypatch.setattr(import_tasks.repo, "increment_catalog_revision", increment)
    monkeypatch.setattr(import_tasks.repo, "update_job", update)

    await import_tasks.apply_special_equipment_import.original_func(str(job["id"]))

    assert [row["code"] for row in appended] == [
        "ATTRIBUTE_ALREADY_ASSIGNED_IN_TRIM"
    ]
    assert updates[-1]["status"] == ImportStatus.COMPLETED_WITH_WARNINGS.value
    assert updates[-1]["summary"]["rejectedAggregates"] == 1


@pytest.mark.asyncio
async def test_worker_unrelated_revision_change_remains_preview_stale(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan, modification_id, trim_id, attribute_id = _trim_conflict_plan()
    job = _job(
        status=ImportStatus.APPLYING.value,
        normalized_artifact_key="imports/normalized-stale.zip",
        mode=ImportMode.PATCH.value,
        error_policy=ImportErrorPolicy.BEST_EFFORT.value,
        template_version=4,
        catalog_revision=7,
        source_code="catalog_v2",
    )
    updates: list[dict[str, Any]] = []

    async def no_conflict_state(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        return {
            "modification_by_trim": {trim_id: modification_id},
            "trim_values": set(),
            "modification_values": set(),
        }

    async def update(*_args: Any, **values: Any) -> dict[str, Any]:
        updates.append(values)
        return {**job, **values}

    async def unexpected_apply(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        raise AssertionError("stale plan must not be applied")

    _configure_worker_trim_race(
        monkeypatch,
        job=job,
        plan=plan,
        modification_id=modification_id,
        trim_id=trim_id,
        attribute_id=attribute_id,
    )
    monkeypatch.setattr(
        import_tasks.repo,
        "lock_and_read_trim_modification_value_state",
        no_conflict_state,
    )
    monkeypatch.setattr(import_tasks.repo, "update_job", update)
    monkeypatch.setattr(
        import_tasks.repo,
        "apply_normalized_plan",
        unexpected_apply,
    )

    await import_tasks.apply_special_equipment_import.original_func(str(job["id"]))

    assert updates[-1]["status"] == ImportStatus.PREVIEW_STALE.value
    assert updates[-1]["error_code"] == "PREVIEW_STALE"


@pytest.mark.asyncio
async def test_worker_does_not_misclassify_preexisting_v4_snapshot_value(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan, modification_id, trim_id, attribute_id = _trim_conflict_plan()
    job = _job(
        status=ImportStatus.APPLYING.value,
        normalized_artifact_key="imports/normalized-v4-snapshot.zip",
        mode=ImportMode.FULL_SNAPSHOT.value,
        error_policy=ImportErrorPolicy.ATOMIC.value,
        template_version=4,
        catalog_revision=7,
        source_code="catalog_v2",
    )
    updates: list[dict[str, Any]] = []

    async def update(*_args: Any, **values: Any) -> dict[str, Any]:
        updates.append(values)
        return {**job, **values}

    async def unexpected_apply(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        raise AssertionError("unrelated revision must keep the preview stale")

    _configure_worker_trim_race(
        monkeypatch,
        job=job,
        plan=plan,
        modification_id=modification_id,
        trim_id=trim_id,
        attribute_id=attribute_id,
        expected_trim_value_markers=frozenset(
            {(trim_id, modification_id, attribute_id)}
        ),
    )
    monkeypatch.setattr(import_tasks.repo, "update_job", update)
    monkeypatch.setattr(
        import_tasks.repo,
        "apply_normalized_plan",
        unexpected_apply,
    )

    await import_tasks.apply_special_equipment_import.original_func(str(job["id"]))

    assert updates[-1]["status"] == ImportStatus.PREVIEW_STALE.value
    assert updates[-1]["error_code"] == "PREVIEW_STALE"


@pytest.mark.asyncio
async def test_create_retry_reuses_job_and_does_not_start_second_multipart(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    stored: dict[str, Any] | None = None

    async def acquire(*_args: Any, **_kwargs: Any) -> None:
        return None

    async def get_existing(*_args: Any, **_kwargs: Any) -> dict[str, Any] | None:
        return stored

    async def create_job(
        *_args: Any,
        values: dict[str, Any],
        **_kwargs: Any,
    ) -> dict[str, Any]:
        nonlocal stored
        stored = _job(**values)
        return stored

    monkeypatch.setattr(
        import_service.repo,
        "acquire_create_idempotency_lock",
        acquire,
    )
    monkeypatch.setattr(
        import_service.repo,
        "get_job_by_idempotency",
        get_existing,
    )
    monkeypatch.setattr(import_service.repo, "create_job", create_job)
    storage = _Storage()
    command = import_service.CreateImportCommand(
        requested_by=uuid4(),
        idempotency_key="same-request-key",
        filename="catalog.xlsx",
        size=1024,
        mode=ImportMode.PATCH,
        template_version=8,
    )

    first, first_replay = await import_service.create_import(
        command,
        object(),  # type: ignore[arg-type]
        storage,  # type: ignore[arg-type]
    )
    second, second_replay = await import_service.create_import(
        command,
        object(),  # type: ignore[arg-type]
        storage,  # type: ignore[arg-type]
    )

    assert first["id"] == second["id"]
    assert (first_replay, second_replay) == (False, True)
    assert storage.initiated == 1


@pytest.mark.asyncio
async def test_create_maps_deleted_target_race_and_aborts_multipart(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def no_op(*_args: Any, **_kwargs: Any) -> None:
        return None

    async def no_existing(*_args: Any, **_kwargs: Any) -> None:
        return None

    async def missing_target(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        raise import_repository.ImportTargetWarehouseUnavailableError

    monkeypatch.setattr(import_service.repo, "warehouse_is_active", AsyncMock(return_value=True))
    monkeypatch.setattr(import_service.repo, "acquire_create_idempotency_lock", no_op)
    monkeypatch.setattr(import_service.repo, "get_job_by_idempotency", no_existing)
    monkeypatch.setattr(import_service.repo, "create_job", missing_target)
    storage = _Storage()
    command = import_service.CreateImportCommand(
        requested_by=uuid4(),
        idempotency_key="target-deleted-race",
        filename="catalog.xlsx",
        size=1024,
        mode=ImportMode.PATCH,
        template_version=8,
        target_warehouse_id=uuid4(),
    )

    with pytest.raises(ServiceError) as raised:
        await import_service.create_import(
            command,
            object(),  # type: ignore[arg-type]
            storage,  # type: ignore[arg-type]
        )

    assert (raised.value.status_code, str(raised.value)) == (
        422,
        "WAREHOUSE_NOT_FOUND",
    )
    assert storage.aborted and storage.aborted[0][1] == "upload-1"


@pytest.mark.asyncio
async def test_v3_normalized_artifact_apply_becomes_preview_stale(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    job = _job(
        status="applying",
        normalized_artifact_key="imports/normalized-v3.zip",
    )
    marked: dict[str, Any] = {}

    async def get_job(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        return job

    async def mark_preview_stale(job_id: Any, detail: str) -> None:
        marked.update(job_id=job_id, detail=detail)

    monkeypatch.setattr(import_tasks, "AsyncSessionLocal", _TaskSessionContext)
    monkeypatch.setattr(import_tasks.repo, "get_job", get_job)
    monkeypatch.setattr(
        import_tasks,
        "get_special_equipment_import_storage",
        _ArtifactStorage,
    )
    monkeypatch.setattr(
        import_tasks,
        "read_rows_archive",
        lambda *_args, **_kwargs: ({"schema_version": 3}, {}),
    )
    monkeypatch.setattr(import_tasks, "_mark_preview_stale", mark_preview_stale)

    await import_tasks.apply_special_equipment_import.original_func(str(job["id"]))

    assert marked == {
        "job_id": job["id"],
        "detail": import_tasks._ARTIFACT_SCHEMA_STALE_DETAIL,
    }


@pytest.mark.asyncio
async def test_repeated_multipart_part_is_idempotent_without_storage_write(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    raw_digest = hashlib.sha256(b"same-part").digest()
    digest = raw_digest.hex()
    content_digest = f"sha-256=:{base64.b64encode(raw_digest).decode()}:"
    job = _job()
    existing = {
        "part_number": 1,
        "byte_start": 0,
        "byte_end": 1023,
        "size_bytes": 1024,
        "sha256": digest,
    }
    retries = 0

    async def get_job(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        return job

    async def no_op(*_args: Any, **_kwargs: Any) -> None:
        return None

    async def get_part(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        return existing

    async def mark_retry(*_args: Any, **_kwargs: Any) -> None:
        nonlocal retries
        retries += 1

    monkeypatch.setattr(import_service.repo, "get_job", get_job)
    monkeypatch.setattr(import_service.repo, "acquire_part_lock", no_op)
    monkeypatch.setattr(import_service.repo, "get_part", get_part)
    monkeypatch.setattr(import_service.repo, "mark_part_retry", mark_retry)
    command = import_service.PutImportPartCommand(
        import_id=job["id"],
        requested_by=job["requested_by"],
        part_number=1,
        byte_start=0,
        byte_end=1023,
        content_length=1024,
        content_digest=content_digest,
        chunks=_empty_chunks(),
    )

    result = await import_service.put_import_part(
        command,
        object(),  # type: ignore[arg-type]
        object(),  # type: ignore[arg-type]
    )

    assert result["idempotent"] is True
    assert retries == 1


async def _empty_chunks() -> Any:
    if False:
        yield b""


@pytest.mark.asyncio
async def test_cancel_aborts_incomplete_multipart_and_keeps_catalog_untouched(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    job = _job()

    async def get_job(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        return job

    async def update_job(
        *_args: Any,
        **values: Any,
    ) -> dict[str, Any]:
        return {**job, **values, "updated_at": datetime.now(UTC)}

    monkeypatch.setattr(import_service.repo, "get_job", get_job)
    monkeypatch.setattr(import_service.repo, "update_job", update_job)
    storage = _Storage()

    result = await import_service.cancel_import(
        object(),  # type: ignore[arg-type]
        storage,  # type: ignore[arg-type]
        import_id=job["id"],
        requested_by=job["requested_by"],
    )

    assert result["status"] == "cancelled"
    assert storage.aborted == [(job["source_object_key"], job["multipart_upload_id"])]


@pytest.mark.asyncio
async def test_best_effort_preview_exposes_errors_as_nonblocking_warnings(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    job = _job(
        status="preview_ready",
        preview_hash="a" * 64,
        catalog_revision=7,
        summary={"errors": 3, "warnings": 2},
    )

    async def get_job(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        return job

    monkeypatch.setattr(import_service.repo, "get_job", get_job)
    preview = await import_service.get_preview(
        object(),  # type: ignore[arg-type]
        import_id=job["id"],
        requested_by=job["requested_by"],
    )

    assert preview["blockingIssues"] == 0
    assert preview["warnings"] == 5
    assert preview["canApply"] is True


@pytest.mark.asyncio
async def test_apply_requires_matching_preview_hash_and_destructive_confirmation(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    job = _job(
        status="preview_ready",
        preview_hash="a" * 64,
        summary={"requiresDestructiveConfirmation": True},
    )

    async def get_job(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        return job

    async def update_job(
        *_args: Any,
        **values: Any,
    ) -> dict[str, Any]:
        return {**job, **values, "updated_at": datetime.now(UTC)}

    monkeypatch.setattr(import_service.repo, "get_job", get_job)
    monkeypatch.setattr(import_service.repo, "update_job", update_job)
    with pytest.raises(ServiceError, match="PREVIEW_STALE"):
        await import_service.request_import_application(
            import_service.ApplyImportCommand(
                import_id=job["id"],
                requested_by=job["requested_by"],
                if_match="b" * 64,
                confirm_destructive_changes=True,
            ),
            object(),  # type: ignore[arg-type]
        )
    with pytest.raises(ServiceError, match="DESTRUCTIVE_CONFIRMATION_REQUIRED"):
        await import_service.request_import_application(
            import_service.ApplyImportCommand(
                import_id=job["id"],
                requested_by=job["requested_by"],
                if_match=job["preview_hash"],
                confirm_destructive_changes=False,
            ),
            object(),  # type: ignore[arg-type]
        )

    result = await import_service.request_import_application(
        import_service.ApplyImportCommand(
            import_id=job["id"],
            requested_by=job["requested_by"],
            if_match=job["preview_hash"],
            confirm_destructive_changes=True,
        ),
        object(),  # type: ignore[arg-type]
    )
    assert result["status"] == "applying"


class _NestedTransaction(AbstractAsyncContextManager[None]):
    async def __aenter__(self) -> None:
        return None

    async def __aexit__(
        self,
        exc_type: type[BaseException] | None,
        exc: BaseException | None,
        traceback: Any,
    ) -> bool:
        return False


class _ApplySession:
    nested_transactions = 0
    flushed = False

    def begin_nested(self) -> _NestedTransaction:
        self.nested_transactions += 1
        return _NestedTransaction()

    async def flush(self) -> None:
        self.flushed = True


def _aggregate_plan() -> dict[str, list[dict[str, Any]]]:
    return {
        "marks": [
            {
                "id": uuid4(),
                "code": "bad",
                "operation": "ADD",
                "_aggregate_kind": "mark",
                "_aggregate_code": "bad",
                "_sheet_code": "Марки",
                "_row_number": 2,
            },
            {
                "id": uuid4(),
                "code": "good",
                "operation": "ADD",
                "_aggregate_kind": "mark",
                "_aggregate_code": "good",
                "_sheet_code": "Марки",
                "_row_number": 3,
            },
        ]
    }


@pytest.mark.asyncio
@pytest.mark.parametrize("mode", ["APPEND", "PATCH"])
async def test_partial_modes_rollback_only_conflicting_aggregate(
    mode: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def apply_aggregate(
        _session: Any,
        aggregate: dict[str, list[dict[str, Any]]],
        *,
        counts: dict[str, int],
    ) -> None:
        row = aggregate["marks"][0]
        counts["created"] += 1
        if row["code"] == "bad":
            raise IntegrityError("insert", {}, Exception("conflict"))

    monkeypatch.setattr(
        import_repository,
        "_v2_apply_aggregate",
        apply_aggregate,
    )
    session = _ApplySession()
    result = await import_repository.apply_normalized_plan(
        session,  # type: ignore[arg-type]
        job_id=uuid4(),
        mode=mode,
        plan=_aggregate_plan(),
    )

    assert result["counts"]["created"] == 1
    assert len(result["rejected_aggregates"]) == 1
    assert result["rejected_aggregates"][0]["external_key"] == "bad"
    assert session.nested_transactions == 2
    assert session.flushed is True


@pytest.mark.asyncio
async def test_full_replacement_propagates_conflict_before_cleanup(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    relations_cleaned = False
    retire_called = False

    async def apply_aggregate(*_args: Any, **_kwargs: Any) -> None:
        raise IntegrityError("insert", {}, Exception("conflict"))

    async def delete_relations(*_args: Any, **_kwargs: Any) -> set[UUID]:
        nonlocal relations_cleaned
        relations_cleaned = True
        return set()

    async def retire_entities(*_args: Any, **_kwargs: Any) -> None:
        nonlocal retire_called
        retire_called = True

    monkeypatch.setattr(
        import_repository,
        "_v2_apply_aggregate",
        apply_aggregate,
    )
    monkeypatch.setattr(
        import_repository,
        "_v2_delete_missing_snapshot_relations",
        delete_relations,
    )
    monkeypatch.setattr(
        import_repository,
        "_v2_retire_missing_snapshot_entities",
        retire_entities,
    )
    with pytest.raises(ImportAggregateSemanticConflictError):
        await import_repository.apply_normalized_plan(
            _ApplySession(),  # type: ignore[arg-type]
            job_id=uuid4(),
            mode="FULL_SNAPSHOT",
            plan=_aggregate_plan(),
        )

    assert relations_cleaned is True
    assert retire_called is False


@pytest.mark.asyncio
async def test_full_replacement_runs_global_missing_row_cleanup(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def apply_aggregate(
        _session: Any,
        _aggregate: Any,
        *,
        counts: dict[str, int],
        snapshot: bool = False,
    ) -> None:
        assert snapshot is True
        counts["created"] += 1

    async def delete_relations(
        _session: Any,
        *,
        plan: Any,
        counts: dict[str, int],
        include_v4_catalog: bool,
    ) -> set[UUID]:
        assert plan == _aggregate_plan_value
        assert include_v4_catalog is True
        counts["removed"] += 4
        return set()

    async def retire_entities(
        _session: Any,
        *,
        plan: Any,
        counts: dict[str, int],
        include_v4_catalog: bool,
        target_warehouse_id: Any = None,
    ) -> None:
        assert plan == _aggregate_plan_value
        assert include_v4_catalog is True

    _aggregate_plan_value = _aggregate_plan()
    monkeypatch.setattr(
        import_repository,
        "_v2_apply_aggregate",
        apply_aggregate,
    )
    monkeypatch.setattr(
        import_repository,
        "_v2_delete_missing_snapshot_relations",
        delete_relations,
    )
    monkeypatch.setattr(
        import_repository,
        "_v2_retire_missing_snapshot_entities",
        retire_entities,
    )
    result = await import_repository.apply_normalized_plan(
        _ApplySession(),  # type: ignore[arg-type]
        job_id=uuid4(),
        mode="FULL_SNAPSHOT",
        plan=_aggregate_plan_value,
    )

    assert result["counts"] == {
        "created": 2,
        "updated": 0,
        "archived": 0,
        "removed": 4,
    }
    assert result["rejected_aggregates"] == []


@pytest.mark.parametrize(
    ("retry_count", "expected_status"),
    [(3, "failed_retryable"), (4, "failed")],
)
@pytest.mark.asyncio
async def test_task_retry_state_is_bounded_and_releases_lease(
    retry_count: int,
    expected_status: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    updates: dict[str, Any] = {}

    class Session:
        async def commit(self) -> None:
            return None

    class SessionContext(AbstractAsyncContextManager[Session]):
        async def __aenter__(self) -> Session:
            return Session()

        async def __aexit__(
            self,
            exc_type: type[BaseException] | None,
            exc: BaseException | None,
            traceback: Any,
        ) -> bool:
            return False

    async def get_job(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
        return {"retry_count": retry_count}

    async def update_job(*_args: Any, **values: Any) -> dict[str, Any]:
        updates.update(values)
        return values

    monkeypatch.setattr(import_tasks, "AsyncSessionLocal", SessionContext)
    monkeypatch.setattr(import_tasks.repo, "get_job", get_job)
    monkeypatch.setattr(import_tasks.repo, "update_job", update_job)

    await import_tasks._mark_retryable(uuid4(), "TEMPORARY", "temporary")

    assert updates["status"] == expected_status
    assert updates["lease_owner"] is None
    if expected_status == "failed":
        assert updates["lease_until"] is None
        assert updates["error_code"] == "RETRIES_EXHAUSTED"
    else:
        assert updates["lease_until"] is not None


def test_parser_rejects_formula_and_oversized_text(tmp_path: Path) -> None:
    formula_path = tmp_path / "formula.xlsx"
    formula_path.write_bytes(build_template_v7(mode=ImportMode.PATCH))
    workbook = load_workbook(formula_path)
    workbook["Марки"].append(("Добавить", "kamaz", "=1+1", "Да"))
    workbook.save(formula_path)
    workbook.close()
    with pytest.raises(ImportContractError, match="XLSX_FORMULA_FORBIDDEN"):
        parse_xlsx(formula_path, expected_mode=ImportMode.PATCH)

    with pytest.raises(ImportContractError, match="CELL_TEXT_TOO_LONG"):
        _normalize_cell("я" * 100_001)


def test_full_replacement_requires_all_business_sheets(tmp_path: Path) -> None:
    source = tmp_path / "incomplete-full.xlsx"
    source.write_bytes(build_template_v7(mode=ImportMode.FULL_SNAPSHOT))
    workbook = load_workbook(source)
    del workbook["Категории объявлений"]
    workbook.save(source)
    workbook.close()

    with pytest.raises(
        ImportContractError,
        match="FULL_SNAPSHOT_SHEETS_MISSING",
    ):
        parse_xlsx(source, expected_mode=ImportMode.FULL_SNAPSHOT)


def test_v2_coerce_values_handles_defaults_and_nullability() -> None:
    # 1. NOT NULL column with default/server_default: None should be omitted
    res_def = import_repository._v2_coerce_values(
        SpecialEquipmentCategory.__table__,
        {
            "id": uuid4(),
            "code": "test_cat",
            "name": "Тест",
            "is_attachment_category": None,
        },
    )
    assert "is_attachment_category" not in res_def

    # 2. Nullable column: None should be kept
    res_nullable = import_repository._v2_coerce_values(
        SpecialEquipmentCategory.__table__,
        {
            "id": uuid4(),
            "code": "test_cat",
            "name": "Тест",
            "image_key": None,
        },
    )
    assert "image_key" in res_nullable
    assert res_nullable["image_key"] is None

    # 3. NOT NULL column without default: None should be kept
    res_not_null = import_repository._v2_coerce_values(
        SpecialEquipmentCategory.__table__,
        {
            "id": uuid4(),
            "code": "test_cat",
            "name": "Тест",
            "usage_metric": None,
        },
    )
    assert "usage_metric" in res_not_null
    assert res_not_null["usage_metric"] is None


@pytest.mark.asyncio
async def test_main_scenario_categories_blank_attachment_regression(
    db_session: Any,
) -> None:
    category_ids = [uuid4() for _ in range(11)]
    categories = [
        {
            "id": category_ids[i],
            "code": f"faw-cat-{i}",
            "operation": "ADD",
            "values": {
                "name": f"FAW Категория {i}",
                "slug": f"faw-cat-{i}-{uuid4().hex[:6]}",
                "usage_metric": "mileage_km",
                "is_active": True,
                "sort_order": i,
                # is_attachment_category omitted/None
            },
            "_aggregate_kind": "category",
            "_aggregate_code": f"faw-cat-{i}",
            "_sheet_code": "Категории",
            "_row_number": i + 2,
        }
        for i in range(11)
    ]
    # Relation between cat 0 (parent) and cat 1 (child)
    category_relations = [
        {
            "operation": "ADD",
            "values": {
                "parent_id": category_ids[0],
                "child_id": category_ids[1],
                "sort_order": 0,
            },
            "_aggregate_kind": "category",
            "_aggregate_code": "faw-cat-1",
            "_sheet_code": "Связи категорий",
            "_row_number": 2,
        }
    ]

    plan = {
        "categories": categories,
        "category_relations": category_relations,
    }

    # 1. APPEND mode
    result_append = await import_repository.apply_normalized_plan(
        db_session,
        job_id=uuid4(),
        mode=ImportMode.APPEND.value,
        plan=plan,
    )
    assert result_append["rejected_aggregates"] == []
    assert result_append["counts"]["created"] >= 11

    # Verify all 11 categories in DB have is_attachment_category == False
    for cid in category_ids:
        cat = await db_session.get(SpecialEquipmentCategory, cid)
        assert cat is not None
        assert cat.is_attachment_category is False

    # 2. PATCH mode (with blanks)
    patch_categories = [
        {
            **cat,
            "operation": "SET",
            "values": {
                **cast("dict[str, Any]", cat["values"]),
                "name": f"FAW Категория {i} (patch)",
            },
        }
        for i, cat in enumerate(categories)
    ]
    patch_category_relations = [
        {**rel, "operation": "SET"} for rel in category_relations
    ]
    patch_plan = {
        "categories": patch_categories,
        "category_relations": patch_category_relations,
    }
    result_patch = await import_repository.apply_normalized_plan(
        db_session,
        job_id=uuid4(),
        mode=ImportMode.PATCH.value,
        plan=patch_plan,
    )
    assert result_patch["rejected_aggregates"] == []
    for cid in category_ids:
        cat = await db_session.get(SpecialEquipmentCategory, cid)
        assert cat is not None
        assert cat.is_attachment_category is False

    # 3. FULL_SNAPSHOT mode
    result_snapshot = await import_repository.apply_normalized_plan(
        db_session,
        job_id=uuid4(),
        mode=ImportMode.FULL_SNAPSHOT.value,
        plan=patch_plan,
    )
    assert result_snapshot["rejected_aggregates"] == []
    for cid in category_ids:
        cat = await db_session.get(SpecialEquipmentCategory, cid)
        assert cat is not None
        assert cat.is_attachment_category is False


@pytest.mark.asyncio
async def test_patch_and_snapshot_attachment_category_and_active_retention(
    db_session: Any,
) -> None:
    cat_id = uuid4()
    cat_code = f"cat-retention-{uuid4().hex[:6]}"
    category = SpecialEquipmentCategory(
        id=cat_id,
        code=cat_code,
        name=f"Retention Category {uuid4().hex[:6]}",
        slug=f"retention-{uuid4().hex[:6]}",
        usage_metric="engine_hours",
        is_attachment_category=True,
        is_active=True,
        sort_order=10,
    )
    db_session.add(category)
    await db_session.flush()

    # 1. PATCH with blank is_attachment_category: keeps True
    patch_plan = {
        "categories": [
            {
                "id": cat_id,
                "code": cat_code,
                "operation": "SET",
                "values": {
                    "name": "Updated name",
                    "slug": category.slug,
                    "usage_metric": "engine_hours",
                    "is_attachment_category": True,  # retained from current
                    "is_active": True,
                },
                "_aggregate_kind": "category",
                "_aggregate_code": cat_code,
                "_sheet_code": "Категории",
                "_row_number": 2,
            }
        ]
    }
    res_patch = await import_repository.apply_normalized_plan(
        db_session,
        job_id=uuid4(),
        mode=ImportMode.PATCH.value,
        plan=patch_plan,
    )
    assert res_patch["rejected_aggregates"] == []
    cat_db = await db_session.get(SpecialEquipmentCategory, cat_id)
    assert cat_db is not None
    assert cat_db.is_attachment_category is True

    # 2. PATCH with clear_fields: becomes False
    clear_plan = {
        "categories": [
            {
                "id": cat_id,
                "code": cat_code,
                "operation": "SET",
                "values": {
                    "name": "Updated name",
                    "slug": category.slug,
                    "usage_metric": "engine_hours",
                    "is_attachment_category": False,  # cleared to default False
                    "is_active": True,
                },
                "_aggregate_kind": "category",
                "_aggregate_code": cat_code,
                "_sheet_code": "Категории",
                "_row_number": 2,
            }
        ]
    }
    res_clear = await import_repository.apply_normalized_plan(
        db_session,
        job_id=uuid4(),
        mode=ImportMode.PATCH.value,
        plan=clear_plan,
    )
    assert res_clear["rejected_aggregates"] == []
    await db_session.refresh(cat_db)
    assert cat_db.is_attachment_category is False

    # 3. Reset to True and test FULL_SNAPSHOT with blank fields retaining current
    cat_db.is_attachment_category = True
    cat_db.is_active = True
    await db_session.flush()

    snapshot_plan = {
        "categories": [
            {
                "id": cat_id,
                "code": cat_code,
                "operation": "SET",
                "values": {
                    "name": "Snapshot Category",
                    "slug": category.slug,
                    "usage_metric": "engine_hours",
                    "is_attachment_category": True,  # planner resolves current
                    "is_active": True,  # planner resolves current
                },
                "_aggregate_kind": "category",
                "_aggregate_code": cat_code,
                "_sheet_code": "Категории",
                "_row_number": 2,
            }
        ]
    }
    res_snap = await import_repository.apply_normalized_plan(
        db_session,
        job_id=uuid4(),
        mode=ImportMode.FULL_SNAPSHOT.value,
        plan=snapshot_plan,
    )
    assert res_snap["rejected_aggregates"] == []
    await db_session.refresh(cat_db)
    assert cat_db.is_attachment_category is True
    assert cat_db.is_active is True


@pytest.mark.asyncio
async def test_integrity_error_translation_and_cascade_rejection(
    db_session: Any,
) -> None:
    # 1. Prepare existing category in DB
    existing_slug = f"dup-slug-{uuid4().hex[:6]}"
    existing_cat = SpecialEquipmentCategory(
        id=uuid4(),
        code=f"cat-exist-{uuid4().hex[:6]}",
        name="Существующая категория",
        slug=existing_slug,
        usage_metric="mileage_km",
        is_attachment_category=False,
        is_active=True,
        sort_order=0,
    )
    db_session.add(existing_cat)
    await db_session.flush()

    # 2. In PATCH mode: try to insert a category with same slug (unique violation 23505)
    bad_cat_id = uuid4()
    bad_cat_code = f"cat-dup-{uuid4().hex[:6]}"
    child_cat_id = uuid4()
    child_cat_code = f"cat-child-{uuid4().hex[:6]}"

    patch_plan = {
        "categories": [
            {
                "id": bad_cat_id,
                "code": bad_cat_code,
                "operation": "ADD",
                "values": {
                    "name": "Дубликат",
                    "slug": existing_slug,  # will trigger uq_special_equipment_categories_slug
                    "usage_metric": "mileage_km",
                    "is_active": True,
                    "sort_order": 1,
                },
                "_aggregate_kind": "category",
                "_aggregate_code": bad_cat_code,
                "_sheet_code": "Категории",
                "_row_number": 3,
            },
            {
                "id": child_cat_id,
                "code": child_cat_code,
                "operation": "ADD",
                "values": {
                    "name": "Дочерняя категория",
                    "slug": f"child-{uuid4().hex[:6]}",
                    "usage_metric": "mileage_km",
                    "is_active": True,
                    "sort_order": 2,
                },
                "_aggregate_kind": "category",
                "_aggregate_code": child_cat_code,
                "_sheet_code": "Категории",
                "_row_number": 4,
            },
        ],
        "category_relations": [
            {
                "operation": "ADD",
                "values": {
                    "parent_id": bad_cat_id,  # references bad_cat_id which failed!
                    "child_id": child_cat_id,
                    "sort_order": 0,
                },
                "_aggregate_kind": "category",
                "_aggregate_code": child_cat_code,
                "_sheet_code": "Связи категорий",
                "_row_number": 5,
            }
        ],
    }

    result = await import_repository.apply_normalized_plan(
        db_session,
        job_id=uuid4(),
        mode=ImportMode.PATCH.value,
        plan=patch_plan,
    )
    assert len(result["rejected_aggregates"]) == 2
    root_rejection = next(
        r for r in result["rejected_aggregates"] if r["external_key"] == bad_cat_code
    )
    assert root_rejection["code"] == "UNIQUE_CONFLICT"
    assert (
        root_rejection["message"]
        == "Категория с таким названием уже существует с другим кодом"
    )

    child_rejection = next(
        r for r in result["rejected_aggregates"] if r["external_key"] == child_cat_code
    )
    assert child_rejection["code"] == "DEPENDENCY_NOT_APPLIED"
    assert (
        f"Не применено, так как категория `{bad_cat_code}` не загружена"
        in child_rejection["message"]
    )

    # 3. In FULL_SNAPSHOT mode: same conflict raises ImportAggregateSemanticConflictError
    with pytest.raises(ImportAggregateSemanticConflictError) as exc_info:
        await import_repository.apply_normalized_plan(
            db_session,
            job_id=uuid4(),
            mode=ImportMode.FULL_SNAPSHOT.value,
            plan=patch_plan,
        )
    assert exc_info.value.code == "UNIQUE_CONFLICT"
    assert "Лист «Категории»" in str(exc_info.value)
    assert bad_cat_code in str(exc_info.value)

