from __future__ import annotations

import argparse
import csv
import hashlib
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

DEFAULT_STOCK_PATHS = [
    Path("/Users/a_belianskii/Documents/multileasing/мерс + байк + амбертрак.xlsx"),
    Path("/Users/a_belianskii/Documents/multileasing/Сток в наличии 29.01 ver 3.xlsx"),
]
DEFAULT_VEHICLES_CSV = Path(__file__).resolve().parent / "generated_data" / "vehicles.csv"
DEFAULT_EXPORT = Path("~/projects/carcraft-leadgenerator/leasing_export").expanduser()
DEFAULT_VEHICLES_MASTER = DEFAULT_EXPORT / "import" / "full" / "vehicles_master.csv"
DEFAULT_OUTPUT = DEFAULT_EXPORT / "import" / "full" / "catalog_real_stock"
DEFAULT_CHUNK_SIZE = 20_000
VIN_ALPHABET = "ABCDEFGHJKLMNPRSTUVWXYZ0123456789"
REQUIRED_OUTPUT_COLUMNS = ["ID", "Статус", "created_at", "dealer_id", "warehouse_id"]


def _read_stock_rows(paths: list[Path]) -> tuple[list[str], list[dict[str, Any]]]:
    headers: list[str] | None = None
    rows: list[dict[str, Any]] = []
    for path in paths:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            sheet_headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
            if headers is None:
                headers = [str(h) if h is not None else "" for h in sheet_headers]
            for values in ws.iter_rows(min_row=2, values_only=True):
                row = {
                    str(header) if header is not None else "": value
                    for header, value in zip(sheet_headers, values, strict=False)
                }
                if not row.get("VIN"):
                    continue
                rows.append(row)
        finally:
            wb.close()
    if headers is None or not rows:
        raise RuntimeError("No stock rows with VIN found")
    for col in REQUIRED_OUTPUT_COLUMNS:
        if col not in headers:
            headers.append(col)
    return headers, rows


def _read_source_vehicles(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def _read_master(path: Path) -> dict[str, dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return {row["id"]: row for row in csv.DictReader(f, delimiter=";")}


def _generated_vin(vehicle_id: str, used: set[str]) -> str:
    salt = 0
    while True:
        digest = hashlib.blake2b(f"{vehicle_id}:{salt}".encode(), digest_size=16).digest()
        n = int.from_bytes(digest, "big")
        vin = "".join(VIN_ALPHABET[(n >> (i * 5)) % len(VIN_ALPHABET)] for i in range(17))
        if vin not in used:
            used.add(vin)
            return vin
        salt += 1


def _coerce_price(value: str) -> int | str:
    try:
        return int(float(value))
    except ValueError:
        return value


def _build_output_row(
    *,
    template: dict[str, Any],
    vehicle_ref: dict[str, str],
    master: dict[str, str],
    used_vins: set[str],
    first_use: bool,
) -> dict[str, Any]:
    row = dict(template)
    vehicle_id = vehicle_ref["vehicle_id"]
    row["ID"] = vehicle_id
    row["VIN"] = str(template.get("VIN")).strip() if first_use else _generated_vin(vehicle_id, used_vins)
    if first_use:
        used_vins.add(row["VIN"])
    row["Стоимость "] = _coerce_price(vehicle_ref.get("unit_price", ""))
    row["Специальная стоимость "] = _coerce_price(vehicle_ref.get("unit_price", ""))
    row["Статус"] = master.get("status") or "available"
    row["created_at"] = master.get("created_at") or ""
    row["dealer_id"] = master.get("dealer_id") or ""
    row["warehouse_id"] = master.get("warehouse_id") or ""
    return row


def _write_chunk(path: Path, headers: list[str], rows: list[dict[str, Any]]) -> None:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Лист1")
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])
    wb.save(path)


def generate_excel_chunks(
    *,
    stock_paths: list[Path],
    vehicles_csv: Path,
    vehicles_master_csv: Path,
    output_dir: Path,
    chunk_size: int,
) -> list[Path]:
    headers, stock_rows = _read_stock_rows(stock_paths)
    source_vehicles = _read_source_vehicles(vehicles_csv)
    master_by_id = _read_master(vehicles_master_csv)
    output_dir.mkdir(parents=True, exist_ok=True)

    used_vins: set[str] = set()
    chunk_rows: list[dict[str, Any]] = []
    generated: list[Path] = []
    stock_len = len(stock_rows)

    for index, vehicle_ref in enumerate(source_vehicles):
        vehicle_id = vehicle_ref["vehicle_id"]
        master = master_by_id.get(vehicle_id)
        if master is None:
            continue
        template_index = index % stock_len
        output_row = _build_output_row(
            template=stock_rows[template_index],
            vehicle_ref=vehicle_ref,
            master=master,
            used_vins=used_vins,
            first_use=index < stock_len,
        )
        chunk_rows.append(output_row)
        if len(chunk_rows) >= chunk_size:
            path = output_dir / f"part-{len(generated) + 1:04d}.xlsx"
            _write_chunk(path, headers, chunk_rows)
            generated.append(path)
            chunk_rows = []

    if chunk_rows:
        path = output_dir / f"part-{len(generated) + 1:04d}.xlsx"
        _write_chunk(path, headers, chunk_rows)
        generated.append(path)
    return generated


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--vehicles-csv", type=Path, default=DEFAULT_VEHICLES_CSV)
    parser.add_argument("--vehicles-master-csv", type=Path, default=DEFAULT_VEHICLES_MASTER)
    parser.add_argument("--chunk-size", type=int, default=DEFAULT_CHUNK_SIZE)
    args = parser.parse_args()
    generated = generate_excel_chunks(
        stock_paths=DEFAULT_STOCK_PATHS,
        vehicles_csv=args.vehicles_csv,
        vehicles_master_csv=args.vehicles_master_csv,
        output_dir=args.output_dir,
        chunk_size=args.chunk_size,
    )
    for path in generated:
        print(path)  # noqa: T201


if __name__ == "__main__":
    main()
