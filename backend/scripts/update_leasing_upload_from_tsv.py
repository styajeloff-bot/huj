from __future__ import annotations

import argparse
import csv
import logging
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

CLIENT_FIELD_MAP = {
    "company_name": "Название компании",
    "director": "ФИО генерального директора",
    "company_phone": "Телефон",
    "company_email": "Email",
}

LEASING_COMPANY_FIELD_MAP = {
    "company_name": "Название компании",
    "company_phone": "Телефон",
    "company_email": "Email",
}


def normalize_inn(value: Any) -> str:
    if value is None:
        return ""

    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]

    return text


def clean_value(value: Any) -> str:
    if value is None:
        return ""

    return str(value).strip()


def read_first_by_inn(path: Path) -> dict[str, dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file, delimiter="\t")
        if not reader.fieldnames or "ИНН" not in reader.fieldnames:
            raise ValueError(f"{path}: не найден столбец ИНН")

        rows_by_inn: dict[str, dict[str, str]] = {}
        for row in reader:
            inn = normalize_inn(row.get("ИНН"))
            if inn and inn not in rows_by_inn:
                rows_by_inn[inn] = {key: clean_value(value) for key, value in row.items()}

    return rows_by_inn


def header_indexes(ws: Worksheet) -> dict[str, int]:
    return {
        clean_value(cell.value): cell.column
        for cell in ws[1]
        if clean_value(cell.value)
    }


def update_sheet(ws: Worksheet, rows_by_inn: dict[str, dict[str, str]], field_map: dict[str, str]) -> int:
    headers = header_indexes(ws)
    if "company_inn" not in headers:
        raise ValueError(f"лист {ws.title}: не найден столбец company_inn")

    missing_headers = [header for header in field_map if header not in headers]
    if missing_headers:
        raise ValueError(f"лист {ws.title}: не найдены столбцы {', '.join(missing_headers)}")

    changed_cells = 0
    inn_column = headers["company_inn"]

    for row_index in range(2, ws.max_row + 1):
        inn = normalize_inn(ws.cell(row=row_index, column=inn_column).value)
        source_row = rows_by_inn.get(inn)
        if not source_row:
            continue

        for excel_header, tsv_header in field_map.items():
            value = clean_value(source_row.get(tsv_header))
            if not value:
                continue

            cell = ws.cell(row=row_index, column=headers[excel_header])
            if cell.value != value:
                cell.value = value
                changed_cells += 1

    return changed_cells


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Обновляет листы clients и leasing_companies в Excel данными из TSV по company_inn/ИНН.",
    )
    parser.add_argument("--source-file", type=Path, default=Path("leasing_upload_v4.xlsx"))
    parser.add_argument("--clients-file", type=Path, default=Path("clients.tsv"))
    parser.add_argument("--leasing-companies-file", type=Path, default=Path("leasing_companies.tsv"))
    parser.add_argument("--output-file", type=Path, default=Path("output.xlsx"))
    return parser.parse_args()


def main() -> None:
    args = parse_args()

    clients_by_inn = read_first_by_inn(args.clients_file)
    leasing_companies_by_inn = read_first_by_inn(args.leasing_companies_file)

    workbook = load_workbook(args.source_file)
    for sheet_name in ("clients", "leasing_companies"):
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"{args.source_file}: не найден лист {sheet_name}")

    clients_changed = update_sheet(workbook["clients"], clients_by_inn, CLIENT_FIELD_MAP)
    leasing_companies_changed = update_sheet(
        workbook["leasing_companies"],
        leasing_companies_by_inn,
        LEASING_COMPANY_FIELD_MAP,
    )

    workbook.save(args.output_file)

    logger = logging.getLogger(__name__)
    logger.info("Создан файл: %s", args.output_file)
    logger.info("clients: изменено ячеек %s", clients_changed)
    logger.info("leasing_companies: изменено ячеек %s", leasing_companies_changed)


if __name__ == "__main__":
    main()
