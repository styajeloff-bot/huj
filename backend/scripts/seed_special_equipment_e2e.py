"""Namespace-aware explicit E2E fixture for the special-equipment catalog."""

from __future__ import annotations

import argparse
import asyncio
import json
import os
import sys
import tempfile
from collections.abc import Mapping, Sequence
from contextlib import suppress
from dataclasses import asdict, dataclass, field
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit
from uuid import NAMESPACE_URL, UUID, uuid5

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

import sqlalchemy as sa  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from sqlalchemy.ext.asyncio import AsyncSession  # noqa: E402

from domain.special_equipment_import import (  # noqa: E402
    DATA_SHEET_HEADERS,
    DATA_SHEET_NAMES,
    MANIFEST_HEADERS,
    NULL_TOKEN,
    PARAMETERS_SHEET_NAME,
    ImportMode,
)
from infrastructure.auth import generate_tokens, hash_refresh_token  # noqa: E402
from infrastructure.database import AsyncSessionLocal  # noqa: E402
from infrastructure.models.companies import Company  # noqa: E402
from infrastructure.models.special_equipment import (  # noqa: E402
    SpecialEquipmentAttribute,
    SpecialEquipmentAttributeGroup,
    SpecialEquipmentAttributeOption,
    SpecialEquipmentCategory,
    SpecialEquipmentCategoryAttribute,
    SpecialEquipmentCategoryRelation,
    SpecialEquipmentMark,
    SpecialEquipmentModel,
    SpecialEquipmentModification,
    SpecialEquipmentModificationAttributeValue,
    SpecialEquipmentModificationCategory,
    SpecialEquipmentProduct,
    SpecialEquipmentProductAttachment,
    SpecialEquipmentProductCategory,
    SpecialEquipmentProductChassisValue,
    SpecialEquipmentProductSuperstructureValue,
    SpecialEquipmentSuperstructure,
    SpecialEquipmentSuperstructureAttribute,
    SpecialEquipmentSuperstructureCategory,
    SpecialEquipmentTrim,
    SpecialEquipmentTrimAttribute,
    SpecialEquipmentTrimAttributeValue,
    SpecialEquipmentUnit,
)
from infrastructure.models.special_equipment_commerce import (  # noqa: E402
    SpecialEquipmentApplicationItem,
    SpecialEquipmentCartItem,
    SpecialEquipmentGuestCartTransfer,
    SpecialEquipmentPurchaseOrder,
)
from infrastructure.models.users import (  # noqa: E402
    User,
    UserSession,
    VerificationCode,
)
from infrastructure.models.vehicles import (  # noqa: E402
    Warehouse,
    WarehouseMark,
)
from infrastructure.services.special_equipment_xlsx import (  # noqa: E402
    build_template_v7,
)
from infrastructure.settings import settings  # noqa: E402


async def reset_model_showcase_e2e(*args: Any, **kwargs: Any) -> None:
    pass


async def seed_model_showcase_e2e(*args: Any, **kwargs: Any) -> dict[str, Any]:
    from uuid import uuid4
    return {
        "storefront": {"id": str(uuid4()), "slug": "stub-models"},
        "mixed_storefront": {"id": str(uuid4()), "slug": "stub-mixed-models"},
        "logical_mark": {
            "id": "stub_alpha",
            "ids": ["stub_alpha", "stub_beta"],
            "name": "Единая марка",
        },
        "mixed_logical_mark": {
            "id": "stub_gamma",
            "ids": ["stub_gamma"],
            "name": "Вторая марка",
        },
        "models": {
            "alpha": {
                "id": "stub_m_a",
                "name": "Stub Alpha",
                "mark_id": "stub_alpha",
                "complectation_id": "stub_c_a",
                "price": "1000000",
                "min_price": "1000000",
                "discount_price": "900000",
                "volume": "2.0",
                "horse_power": "150",
                "time_to_100": "10.0",
                "category": "stub_cat",
                "body_type": "Седан",
            },
            "beta": {
                "id": "stub_m_b",
                "name": "Stub Beta",
                "mark_id": "stub_beta",
                "complectation_id": "stub_c_b",
                "price": "1000000",
                "min_price": "1000000",
                "discount_price": "900000",
                "volume": "2.0",
                "horse_power": "150",
                "time_to_100": "10.0",
                "category": "stub_cat",
                "body_type": "Седан",
            },
        },
        "mixed_model": {
            "id": "stub_m_g",
            "name": "Stub Gamma",
            "mark_id": "stub_gamma",
            "complectation_id": "stub_c_g",
            "price": "1000000",
            "min_price": "1000000",
            "discount_price": "900000",
            "volume": "2.0",
            "horse_power": "150",
            "time_to_100": "10.0",
            "category": "stub_cat",
            "body_type": "Седан",
        },
        "vehicles": {
            "alpha": {"id": str(uuid4()), "model_id": "stub_m_a", "mark_id": "stub_alpha"},
            "beta": {"id": str(uuid4()), "model_id": "stub_m_b", "mark_id": "stub_beta"},
        },
        "mixed_vehicle": {"id": str(uuid4()), "model_id": "stub_m_g", "mark_id": "stub_gamma"},
        "image_filename": "stub-model-showcase.png",
    }


async def upload_model_showcase_e2e_images(*args: Any, **kwargs: Any) -> None:
    pass


@dataclass(frozen=True)
class FixtureContext:
    """Stable namespace-derived identity that never leaks unsafe input into keys."""

    namespace: str
    uuid_namespace: UUID
    prefix: str
    slug_prefix: str
    client_phone: str
    client_non_owner_phone: str
    employee_phone: str

    def entity_id(self, key: str) -> UUID:
        return uuid5(self.uuid_namespace, key)


@dataclass(frozen=True)
class E2EFixtureManifest:
    """Machine-readable hand-off contract consumed by browser workers."""

    schema_version: int
    namespace: str
    prefix: str
    auth: Mapping[str, Any] = field(default_factory=dict)
    companies: Mapping[str, Any] = field(default_factory=dict)
    categories: Mapping[str, Any] = field(default_factory=dict)
    marks: Mapping[str, Any] = field(default_factory=dict)
    models: Mapping[str, Any] = field(default_factory=dict)
    model_showcase: Mapping[str, Any] = field(default_factory=dict)
    modifications: Mapping[str, Any] = field(default_factory=dict)
    trims: Mapping[str, Any] = field(default_factory=dict)
    attribute_groups: Mapping[str, Any] = field(default_factory=dict)
    attributes: Mapping[str, Any] = field(default_factory=dict)
    options: Mapping[str, Any] = field(default_factory=dict)
    products: Mapping[str, Any] = field(default_factory=dict)
    fingerprint_probes: Mapping[str, Any] = field(default_factory=dict)
    relations: Mapping[str, Any] = field(default_factory=dict)
    cart: Mapping[str, Any] = field(default_factory=dict)
    imports: Mapping[str, Any] = field(default_factory=dict)
    counts: Mapping[str, int] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def fixture_context(namespace: str) -> FixtureContext:
    """Return deterministic, mutually isolated fixture identity for a worker."""

    normalized = namespace.strip()
    if not normalized:
        raise ValueError("namespace must not be blank")
    uuid_namespace = uuid5(NAMESPACE_URL, f"carcraft:special-equipment:e2e:{normalized}")
    token = uuid_namespace.hex[:12]
    phone_seed = uuid_namespace.int % 1_000_000
    return FixtureContext(
        namespace=normalized,
        uuid_namespace=uuid_namespace,
        prefix=f"E2E_{token.upper()}",
        slug_prefix=f"e2e-{token}",
        client_phone=f"+76660{phone_seed:06d}",
        client_non_owner_phone=f"+76662{phone_seed:06d}",
        employee_phone=f"+76661{phone_seed:06d}",
    )


async def _merge[OrmRow](session: AsyncSession, row: OrmRow) -> OrmRow:
    return await session.merge(row)


def _owned_ids(context: FixtureContext, family: str, keys: Sequence[str]) -> list[UUID]:
    if family == "product":
        return [_product_id(context, key) for key in keys]
    return [context.entity_id(f"{family}:{key}") for key in keys]


def _product_id(context: FixtureContext, key: str) -> UUID:
    """Keep public representative IDs stable across arbitrary namespaces."""

    pairs = {
        "representative": ("aggregation-base-a", "aggregation-base-b", 0),
        "equivalent": ("aggregation-base-a", "aggregation-base-b", 1),
        "attachmentCompatible": (
            "aggregation-attachment-a",
            "aggregation-attachment-b",
            0,
        ),
        "attachmentStandalone": (
            "aggregation-attachment-a",
            "aggregation-attachment-b",
            1,
        ),
    }
    pair = pairs.get(key)
    if pair is None:
        return context.entity_id(f"product:{key}")
    first, second, position = pair
    return sorted(
        (
            context.entity_id(f"product:{first}"),
            context.entity_id(f"product:{second}"),
        ),
        key=str,
    )[position]


async def _reset_owned_relationships(
    session: AsyncSession,
    context: FixtureContext,
    *,
    category_keys: Sequence[str],
    modification_keys: Sequence[str],
    attribute_keys: Sequence[str],
    product_keys: Sequence[str],
    user_keys: Sequence[str],
) -> None:
    """Delete mutable fixture-owned edges while preserving unrelated rows."""

    import_code_prefix = f"{context.prefix}_IMPORT"

    async def imported_ids(model: Any) -> list[UUID]:
        return list(
            (
                await session.scalars(
                    sa.select(model.id).where(model.code.startswith(import_code_prefix))
                )
            ).all()
        )

    category_ids = _owned_ids(context, "category", category_keys)
    modification_ids = _owned_ids(context, "modification", modification_keys)
    attribute_ids = _owned_ids(context, "attribute", attribute_keys)
    product_ids = _owned_ids(context, "product", product_keys)
    user_ids = _owned_ids(context, "user", user_keys)
    imported_product_ids = await imported_ids(SpecialEquipmentProduct)
    imported_category_ids = await imported_ids(SpecialEquipmentCategory)
    imported_modification_ids = await imported_ids(SpecialEquipmentModification)
    imported_model_ids = await imported_ids(SpecialEquipmentModel)
    imported_mark_ids = await imported_ids(SpecialEquipmentMark)
    imported_attribute_ids = await imported_ids(SpecialEquipmentAttribute)
    imported_group_ids = await imported_ids(SpecialEquipmentAttributeGroup)
    imported_option_ids = await imported_ids(SpecialEquipmentAttributeOption)
    imported_superstructure_ids = await imported_ids(SpecialEquipmentSuperstructure)
    imported_unit_ids = await imported_ids(SpecialEquipmentUnit)
    superstructure_ids = _owned_ids(context, "superstructure", ("crane",))
    superstructure_ids.extend(imported_superstructure_ids)
    product_ids.extend(imported_product_ids)
    category_ids.extend(imported_category_ids)
    modification_ids.extend(imported_modification_ids)
    attribute_ids.extend(imported_attribute_ids)

    await session.execute(
        sa.delete(SpecialEquipmentGuestCartTransfer).where(
            SpecialEquipmentGuestCartTransfer.user_id.in_(user_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentCartItem).where(
            SpecialEquipmentCartItem.user_id.in_(user_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentApplicationItem).where(
            SpecialEquipmentApplicationItem.product_id.in_(product_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentPurchaseOrder).where(
            SpecialEquipmentPurchaseOrder.product_id.in_(product_ids)
        )
    )
    await session.execute(
        sa.delete(UserSession).where(UserSession.user_id.in_(user_ids))
    )
    await session.execute(
        sa.delete(VerificationCode).where(
            VerificationCode.phone.in_(
                [context.client_phone, context.employee_phone]
            )
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentProductAttachment).where(
            sa.or_(
                SpecialEquipmentProductAttachment.product_id.in_(product_ids),
                SpecialEquipmentProductAttachment.attachment_product_id.in_(product_ids),
            )
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentProductChassisValue).where(
            SpecialEquipmentProductChassisValue.product_id.in_(product_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentProductSuperstructureValue).where(
            SpecialEquipmentProductSuperstructureValue.product_id.in_(product_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentSuperstructureAttribute).where(
            SpecialEquipmentSuperstructureAttribute.superstructure_id.in_(superstructure_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentSuperstructureCategory).where(
            SpecialEquipmentSuperstructureCategory.superstructure_id.in_(superstructure_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentSuperstructure).where(
            SpecialEquipmentSuperstructure.id.in_(superstructure_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentUnit).where(
            SpecialEquipmentUnit.id.in_(imported_unit_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentProductCategory).where(
            SpecialEquipmentProductCategory.product_id.in_(product_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentModificationAttributeValue).where(
            SpecialEquipmentModificationAttributeValue.modification_id.in_(
                modification_ids
            )
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentModificationCategory).where(
            SpecialEquipmentModificationCategory.modification_id.in_(
                modification_ids
            )
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentCategoryAttribute).where(
            sa.or_(
                SpecialEquipmentCategoryAttribute.category_id.in_(category_ids),
                SpecialEquipmentCategoryAttribute.attribute_id.in_(attribute_ids),
            )
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentCategoryRelation).where(
            sa.or_(
                SpecialEquipmentCategoryRelation.parent_id.in_(category_ids),
                SpecialEquipmentCategoryRelation.child_id.in_(category_ids),
            )
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentProduct).where(
            SpecialEquipmentProduct.id.in_(imported_product_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentModification).where(
            SpecialEquipmentModification.id.in_(imported_modification_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentModel).where(
            SpecialEquipmentModel.id.in_(imported_model_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentMark).where(
            SpecialEquipmentMark.id.in_(imported_mark_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentAttributeOption).where(
            SpecialEquipmentAttributeOption.id.in_(imported_option_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentAttribute).where(
            SpecialEquipmentAttribute.id.in_(imported_attribute_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentAttributeGroup).where(
            SpecialEquipmentAttributeGroup.id.in_(imported_group_ids)
        )
    )
    await session.execute(
        sa.delete(SpecialEquipmentCategory).where(
            SpecialEquipmentCategory.id.in_(imported_category_ids)
        )
    )


def _category_card(
    row: SpecialEquipmentCategory,
    *,
    parent_ids: Sequence[UUID],
    path: Sequence[str],
) -> dict[str, Any]:
    return {
        "id": str(row.id),
        "code": row.code,
        "name": row.name,
        "slug": row.slug,
        "path": list(path),
        "parent_ids": [str(item) for item in parent_ids],
        "is_attachment_category": row.is_attachment_category,
    }


def _directory_card(row: Any) -> dict[str, Any]:
    return {
        "id": str(row.id),
        "code": row.code,
        "name": row.name,
        "slug": row.slug,
    }


def _product_card(row: SpecialEquipmentProduct, title: str) -> dict[str, Any]:
    return {
        "id": str(row.id),
        "code": row.code,
        "slug": row.slug,
        "title": title,
        "publication_status": row.publication_status,
        "sale_status": row.sale_status,
        "published_at": row.published_at.isoformat() if row.published_at else None,
    }


async def seed_special_equipment_e2e(  # noqa: PLR0915
    session: AsyncSession,
    *,
    namespace: str,
    reset: bool = False,
    artifacts_dir: Path | None = None,
    auth_state_dir: Path | None = None,
) -> E2EFixtureManifest:
    """Upsert one complete fixture graph in the caller's transaction."""

    context = fixture_context(namespace)
    published_at = datetime(2026, 8, 6, 12, 0, tzinfo=UTC)
    manifest_category_keys = (
        "root",
        "emptyRoot",
        "dagParentA",
        "dagParentB",
        "shared",
        "siblingA",
        "siblingB",
        "attachmentRoot",
        "attachmentChild",
        "attachmentDescendant",
        "leaf",
        "count25",
    )
    manifest_product_keys = (
        "representative",
        "equivalent",
        "nonEquivalent",
        "usedMileage",
        "usedHours",
        "onOrder",
        "noVin",
        "attachmentStandalone",
        "attachmentCompatible",
        "kitWithMod",
        "kitWithoutMod",
        "unpublishedAvailable",
        "publishedUnavailable",
        "trimAbsEsp",
        "trimAbsOnly",
        "trimSafetyNone",
    )
    fingerprint_probe_keys = (
        "fingerprintModification",
        "fingerprintCategory",
        "fingerprintSeller",
        "fingerprintCondition",
        "fingerprintYear",
        "fingerprintPrice",
        "fingerprintMileage",
        "fingerprintHours",
        "fingerprintOwners",
        "fingerprintSaleStatus",
        "fingerprintDescription",
        "fingerprintAttributes",
        "fingerprintCompatibility",
    )
    category_specs = {
        "root": ("Корневой каталог", "mileage_km", False, 10, True),
        "emptyRoot": ("Пустой корень", "mileage_km", False, 20, True),
        "dagParentA": ("Родитель DAG A", "mileage_km", False, 30, True),
        "dagParentB": ("Родитель DAG B", "engine_hours", False, 40, True),
        "shared": ("Общая категория DAG", "mileage_km", False, 50, True),
        "siblingA": ("Соседняя категория A", "mileage_km", False, 60, True),
        "siblingB": ("Соседняя категория B", "engine_hours", False, 70, True),
        "attachmentRoot": ("Надстройки", "engine_hours", True, 80, True),
        "attachmentChild": ("Крановые надстройки", "engine_hours", False, 90, True),
        "attachmentDescendant": ("Надстройки-потомки", "engine_hours", False, 95, True),
        "leaf": ("Конечная категория", "mileage_km", False, 100, True),
        "count0": ("Счётчик 0", "mileage_km", False, 200, True),
        "count1": ("Счётчик 1", "mileage_km", False, 210, True),
        "count2": ("Счётчик 2", "mileage_km", False, 220, True),
        "count5": ("Счётчик 5", "mileage_km", False, 230, True),
        "count11": ("Счётчик 11", "mileage_km", False, 240, True),
        "count21": ("Счётчик 21", "mileage_km", False, 250, True),
        "count22": ("Счётчик 22", "mileage_km", False, 260, True),
        "count25": ("Счётчик 25", "mileage_km", False, 270, True),
        "inactiveRoot": ("Неактивный корень", "mileage_km", False, 280, False),
        "inactiveLinked": ("Неактивная дочерняя", "mileage_km", False, 290, False),
        "inactiveLevel3": ("Неактивный уровень 3", "mileage_km", False, 300, False),
        "inactiveLevel4": ("Неактивный уровень 4", "mileage_km", False, 310, False),
        "inactiveLevel5": ("Неактивный уровень 5", "mileage_km", False, 320, False),
    }
    modification_specs = {
        "baseWhite": "Базовая белая",
        "baseWhiteTwin": "Базовая белая — fingerprint twin",
        "baseBlack": "Базовая чёрная",
        "noCategory": "Legacy без категории",
        "hours": "Техника с моточасами",
        "attachment": "Крановая надстройка",
    }
    attribute_specs = {
        "capacity": ("Грузоподъёмность", "number", "тонн", "range"),
        "color": ("Цвет", "select", None, "exact"),
        "description": ("Назначение", "text", None, "search"),
        "stabilizers": ("Опоры", "boolean", None, "exact"),
        "free": ("Свободная характеристика", "text", None, "search"),
        "abs": ("ABS", "boolean", None, "exact"),
        "esp": ("ESP", "boolean", None, "exact"),
    }
    product_specs: dict[str, dict[str, Any]] = {
        "representative": {"mod": "baseWhite", "price": "10000000", "title": "Белый КАМАЗ 1000 т"},
        "equivalent": {"mod": "baseWhite", "price": "10000000", "title": "Белый КАМАЗ 1000 т"},
        "nonEquivalent": {"mod": "baseBlack", "price": "10000000", "title": "Чёрный КАМАЗ 1000 т"},
        "usedMileage": {"mod": "baseWhite", "price": "7000000", "condition": "used", "owners": 0, "mileage": 125000, "title": "КАМАЗ с пробегом"},
        "usedHours": {"mod": "hours", "category": "siblingB", "price": "6500000", "condition": "used", "owners": 2, "hours": 1840, "title": "Экскаватор с моточасами"},
        "onOrder": {"mod": "baseWhite", "price": "12000000", "sale": "on_order", "title": "КАМАЗ под заказ"},
        "noVin": {"mod": "baseWhite", "price": "9000000", "no_vin": True, "title": "КАМАЗ без VIN"},
        "attachmentStandalone": {"mod": "attachment", "category": "attachmentDescendant", "price": "2500000", "sale": "on_order", "title": "Надстройка отдельно", "description": "Самостоятельная крановая надстройка под заказ"},
        "attachmentCompatible": {"mod": "attachment", "category": "attachmentDescendant", "price": "2500000", "title": "Совместимая надстройка", "description": "Совместимая крановая надстройка"},
        "kitWithMod": {
            "kind": "kit",
            "model": "kamaz1000",
            "mod": "baseWhite",
            "superstructure": "crane",
            "superstructure_name": "PK 23500",
            "superstructure_manufacturer": "Palfinger",
            "price": "13000000",
            "title": "КАМАЗ 1000 + Кран Palfinger PK 23500",
        },
        "kitWithoutMod": {
            "kind": "kit",
            "model": "kamaz1000",
            "superstructure": "crane",
            "superstructure_name": "АТЗ-10",
            "superstructure_manufacturer": "НПО Вектор",
            "price": "12500000",
            "title": "КАМАЗ 1000 + Цистерна НПО Вектор АТЗ-10",
        },
        "unpublishedAvailable": {"mod": "attachment", "category": "attachmentDescendant", "price": "2500000", "publication": "draft", "title": "Черновик надстройки в наличии"},
        "publishedUnavailable": {"mod": "attachment", "category": "attachmentDescendant", "price": "2500000", "sale": "unavailable", "title": "Опубликованная недоступная надстройка"},
        "fingerprintModification": {"mod": "baseWhiteTwin", "price": "10000000", "title": "Fingerprint: modification", "description": "Белый КАМАЗ 1000 т"},
        "fingerprintCategory": {"mod": "baseWhite", "category": "root", "price": "10000000", "title": "Fingerprint: categories", "description": "Белый КАМАЗ 1000 т"},
        "fingerprintSeller": {"mod": "baseWhite", "seller": "alternate", "price": "10000000", "title": "Fingerprint: seller", "description": "Белый КАМАЗ 1000 т"},
        "fingerprintCondition": {"mod": "baseWhite", "price": "10000000", "condition": "used", "owners": 0, "mileage": 1, "title": "Fingerprint: condition", "description": "Белый КАМАЗ 1000 т"},
        "fingerprintYear": {"mod": "baseWhite", "price": "10000000", "year": 2024, "title": "Fingerprint: year", "description": "Белый КАМАЗ 1000 т"},
        "fingerprintPrice": {"mod": "baseWhite", "price": "10000001", "title": "Fingerprint: price", "description": "Белый КАМАЗ 1000 т"},
        "fingerprintMileage": {"mod": "baseWhite", "price": "10000000", "condition": "used", "owners": 0, "mileage": 2, "title": "Fingerprint: mileage", "description": "Белый КАМАЗ 1000 т"},
        "fingerprintHours": {"mod": "hours", "category": "siblingB", "price": "10000000", "condition": "used", "owners": 0, "hours": 1, "title": "Fingerprint: engine hours", "description": "Белый КАМАЗ 1000 т"},
        "fingerprintOwners": {"mod": "baseWhite", "price": "10000000", "condition": "used", "owners": 1, "mileage": 1, "title": "Fingerprint: owners", "description": "Белый КАМАЗ 1000 т"},
        "fingerprintSaleStatus": {"mod": "baseWhite", "price": "10000000", "sale": "on_order", "title": "Fingerprint: sale status", "description": "Белый КАМАЗ 1000 т"},
        "fingerprintDescription": {"mod": "baseWhite", "price": "10000000", "title": "Fingerprint: description", "description": "Иное коммерческое описание"},
        "fingerprintAttributes": {"mod": "baseBlack", "price": "10000000", "title": "Fingerprint: attributes", "description": "Белый КАМАЗ 1000 т"},
        "fingerprintCompatibility": {"mod": "baseWhite", "price": "10000000", "title": "Fingerprint: compatibility", "description": "Белый КАМАЗ 1000 т"},
        "trimAbsEsp": {
            "mod": "baseWhite",
            "trim": "safetyAll",
            "price": "11000000",
            "title": "Комплектация с ABS и ESP",
        },
        "trimAbsOnly": {
            "mod": "baseWhite",
            "trim": "safetyAbsOnly",
            "price": "11100000",
            "title": "Комплектация только с ABS",
        },
        "trimSafetyNone": {
            "mod": "baseWhite",
            "trim": "safetyNone",
            "price": "11200000",
            "title": "Комплектация без ABS и ESP",
        },
    }
    count_values = (0, 1, 2, 5, 11, 21, 22, 25)
    for ordinal in range(1, 26):
        product_specs[f"count_{ordinal}"] = {
            "mod": "baseWhite",
            "count_ordinal": ordinal,
            "price": str(20_000_000 + ordinal),
            "title": f"Счётное объявление {ordinal}",
        }
    if reset:
        await reset_model_showcase_e2e(session, context)
        await _reset_owned_relationships(
            session,
            context,
            category_keys=list(category_specs),
            modification_keys=list(modification_specs),
            attribute_keys=list(attribute_specs),
            product_keys=list(product_specs),
            user_keys=["client", "client_non_owner", "employee"],
        )

    company = await _merge(
        session,
        Company(
            id=context.entity_id("company:seller"),
            name=f"{context.prefix} Тестовый дилер",
            inn=f"99{context.uuid_namespace.int % 100_000_000:08d}",
            kpp="770101001",
            ogrn="1027700000001",
            company_type="dealer",
            legal_address="г. Москва, E2E проезд, д. 1",
            actual_address="г. Москва, E2E проезд, д. 1, строение 2",
            phone=f"+7{'9900000001'}",
            email="full-22294@example.test",
            website="https://full-22294.example.test",
            is_active=True,
        ),
    )
    alternate_company = await _merge(
        session,
        Company(
            id=context.entity_id("company:alternate-seller"),
            name=f"{context.prefix} Альтернативный дилер",
            inn=f"98{context.uuid_namespace.int % 100_000_000:08d}",
            company_type="dealer",
            legal_address="г. Москва, E2E проезд, д. 2",
            is_active=True,
        ),
    )
    sparse_optional_company = await _merge(
        session,
        Company(
            id=context.entity_id("company:22294-sparse-optional"),
            name=f"{context.prefix} 22294 Без контактов",
            inn=f"97{context.uuid_namespace.int % 100_000_000:08d}",
            kpp="770101002",
            ogrn="1027700000002",
            company_type="dealer",
            legal_address="г. Москва, E2E проезд, д. 3",
            actual_address=None,
            phone=None,
            email=None,
            website=None,
            is_active=True,
        ),
    )
    legacy_invalid_company = await _merge(
        session,
        Company(
            id=context.entity_id("company:22294-legacy-invalid"),
            name=f"{context.prefix} 22294 Legacy",
            inn=f"96{context.uuid_namespace.int % 100_000_000:08d}",
            kpp="legacy-kpp",
            ogrn="1027700000003",
            company_type="dealer",
            legal_address="г. Москва, E2E проезд, д. 4",
            is_active=True,
        ),
    )
    warehouse = await _merge(
        session,
        Warehouse(
            id=context.entity_id("warehouse:catalog"),
            name=f"{context.prefix} Каталожный склад",
            owner_company_id=company.id,
            owner_company_type=company.company_type,
            address=f"г. Москва, {context.prefix} склад, д. 1",
            is_active=True,
        ),
    )
    users = {
        "client": await _merge(
            session,
            User(
                id=context.entity_id("user:client"),
                phone=context.client_phone,
                name=f"{context.prefix} Клиент",
                role="client",
                company_id=company.id,
                is_active=True,
                phone_verified=True,
                mfa_enabled=False,
            ),
        ),
        "client_non_owner": await _merge(
            session,
            User(
                id=context.entity_id("user:client-non-owner"),
                phone=context.client_non_owner_phone,
                name=f"{context.prefix} 22294 Клиент без компании",
                role="client",
                company_id=None,
                is_active=True,
                phone_verified=True,
                mfa_enabled=False,
            ),
        ),
        "employee": await _merge(
            session,
            User(
                id=context.entity_id("user:employee"),
                phone=context.employee_phone,
                name=f"{context.prefix} Сотрудник",
                role="carcraft_employee",
                is_active=True,
                phone_verified=True,
                mfa_enabled=False,
            ),
        ),
    }
    model_showcase = await seed_model_showcase_e2e(
        session,
        context,
        warehouse=warehouse,
        employee=users["employee"],
    )
    categories: dict[str, SpecialEquipmentCategory] = {}
    for key, (
        name,
        usage_metric,
        is_attachment,
        sort_order,
        is_active,
    ) in category_specs.items():
        categories[key] = await _merge(
            session,
            SpecialEquipmentCategory(
                id=context.entity_id(f"category:{key}"),
                code=f"{context.prefix}_CATEGORY_{key.upper()}",
                name=f"{context.prefix} {name}",
                slug=f"{context.slug_prefix}-category-{key.lower()}",
                usage_metric=usage_metric,
                is_attachment_category=is_attachment,
                sort_order=sort_order,
                is_active=is_active,
            ),
        )
    relation_specs = (
        ("root", "siblingA", 30),
        ("root", "siblingB", 40),
        ("dagParentA", "shared", 10),
        ("dagParentB", "shared", 10),
        ("shared", "leaf", 10),
        ("root", "leaf", 20),
        ("attachmentRoot", "attachmentChild", 10),
        ("root", "attachmentChild", 50),
        ("attachmentChild", "attachmentDescendant", 10),
        ("root", "inactiveLinked", 60),
        ("inactiveRoot", "inactiveLinked", 10),
        ("inactiveLinked", "inactiveLevel3", 10),
        ("inactiveLevel3", "inactiveLevel4", 10),
        ("inactiveLevel4", "inactiveLevel5", 10),
    )
    category_relations = [
        await _merge(
            session,
            SpecialEquipmentCategoryRelation(
                parent_id=categories[parent].id,
                child_id=categories[child].id,
                sort_order=sort_order,
            ),
        )
        for parent, child, sort_order in relation_specs
    ]

    marks = {
        key: await _merge(
            session,
            SpecialEquipmentMark(
                id=context.entity_id(f"mark:{key}"),
                code=f"{context.prefix}_MARK_{key.upper()}",
                name=f"{context.prefix} {name}",
                slug=f"{context.slug_prefix}-mark-{key}",
                is_active=True,
            ),
        )
        for key, name in {"kamaz": "КАМАЗ", "excavator": "Экскаватор", "attachment": "Надстройка"}.items()
    }
    await _merge(session, WarehouseMark(warehouse_id=warehouse.id, mark_id=marks["kamaz"].id))
    model_specs = {
        "kamaz1000": ("kamaz", "1000"),
        "excavator200": ("excavator", "200"),
        "craneBody": ("attachment", "Крановая установка"),
    }
    models = {
        key: await _merge(
            session,
            SpecialEquipmentModel(
                id=context.entity_id(f"model:{key}"),
                mark_id=marks[mark_key].id,
                code=f"{context.prefix}_MODEL_{key.upper()}",
                name=f"{context.prefix} {name}",
                slug=f"{context.slug_prefix}-model-{key.lower()}",
                is_active=True,
            ),
        )
        for key, (mark_key, name) in model_specs.items()
    }
    modification_model = {
        "baseWhite": "kamaz1000",
        "baseWhiteTwin": "kamaz1000",
        "baseBlack": "kamaz1000",
        "noCategory": "kamaz1000",
        "hours": "excavator200",
        "attachment": "craneBody",
    }
    modifications = {
        key: await _merge(
            session,
            SpecialEquipmentModification(
                id=context.entity_id(f"modification:{key}"),
                model_id=models[modification_model[key]].id,
                code=f"{context.prefix}_MODIFICATION_{key.upper()}",
                name=f"{context.prefix} {name}",
                slug=f"{context.slug_prefix}-modification-{key.lower()}",
                year_from=2020,
                year_to=2026,
                is_active=True,
            ),
        )
        for key, name in modification_specs.items()
    }
    trims = {
        key: await _merge(
            session,
            SpecialEquipmentTrim(
                id=context.entity_id(f"trim:{key}"),
                modification_id=modifications["baseWhite"].id,
                code=f"{context.prefix}_TRIM_{key.upper()}",
                name=f"{context.prefix} {name}",
                slug=f"{context.slug_prefix}-trim-{key.lower()}",
                sort_order=sort_order,
                is_active=True,
            ),
        )
        for key, (name, sort_order) in {
            "safetyAll": ("Безопасность: ABS и ESP", 10),
            "safetyAbsOnly": ("Безопасность: только ABS", 20),
            "safetyNone": ("Безопасность: без ABS и ESP", 30),
        }.items()
    }
    modification_category = {
        "baseWhite": "leaf",
        "baseWhiteTwin": "leaf",
        "baseBlack": "leaf",
        "hours": "siblingB",
        "attachment": "attachmentDescendant",
    }
    modification_categories = [
        await _merge(
            session,
            SpecialEquipmentModificationCategory(
                modification_id=modifications[key].id,
                category_id=categories[category_key].id,
                sort_order=0,
                is_primary=True,
            ),
        )
        for key, category_key in modification_category.items()
    ]
    modification_categories.extend(
        [
            await _merge(
                session,
                SpecialEquipmentModificationCategory(
                    modification_id=modifications["baseWhite"].id,
                    category_id=categories[category_key].id,
                    sort_order=sort_order,
                    is_primary=False,
                ),
            )
            for sort_order, category_key in enumerate(
                ("root", *(f"count{value}" for value in count_values)),
                start=1,
            )
        ]
    )

    groups = {
        key: await _merge(
            session,
            SpecialEquipmentAttributeGroup(
                id=context.entity_id(f"attribute-group:{key}"),
                code=f"{context.prefix}_GROUP_{key.upper()}",
                name=f"{context.prefix} {name}",
                slug=f"{context.slug_prefix}-group-{key}",
                sort_order=sort_order,
                is_active=True,
            ),
        )
        for key, (name, sort_order) in {
            "technical": ("Технические параметры", 10),
            "appearance": ("Описание", 20),
            "safety": ("Безопасность", 30),
        }.items()
    }
    attribute_group = {
        "capacity": "technical",
        "color": "appearance",
        "description": "appearance",
        "stabilizers": "technical",
        "free": "technical",
        "abs": "safety",
        "esp": "safety",
    }
    units = {
        "tonn": await _merge(
            session,
            SpecialEquipmentUnit(
                id=context.entity_id("unit:tonn"),
                code=f"{context.prefix}_UNIT_TONN",
                name="тонн",
                slug=f"{context.slug_prefix}-tonn",
                is_active=True,
            ),
        ),
    }
    attributes = {
        key: await _merge(
            session,
            SpecialEquipmentAttribute(
                id=context.entity_id(f"attribute:{key}"),
                code=f"{context.prefix}_ATTRIBUTE_{key.upper()}",
                name=f"{context.prefix} {name}",
                attribute_group_id=groups[attribute_group[key]].id,
                data_type=data_type,
                unit_id=units["tonn"].id if unit == "тонн" else None,
                filter_kind=filter_kind,
                is_active=True,
            ),
        )
        for key, (name, data_type, unit, filter_kind) in attribute_specs.items()
    }
    options = {
        key: await _merge(
            session,
            SpecialEquipmentAttributeOption(
                id=context.entity_id(f"option:{key}"),
                attribute_id=attributes["color"].id,
                code=f"{context.prefix}_OPTION_{key.upper()}",
                name=f"{context.prefix} {name}",
                sort_order=sort_order,
                is_active=True,
            ),
        )
        for key, (name, sort_order) in {
            "white": ("Белый", 10),
            "black": ("Чёрный", 20),
        }.items()
    }
    category_attributes = [
        await _merge(
            session,
            SpecialEquipmentCategoryAttribute(
                category_id=categories[category_key].id,
                attribute_id=attributes[attribute_key].id,
                group_id=groups[attribute_group[attribute_key]].id,
                is_required=attribute_key in {"capacity", "color"},
                is_filterable=True,
                is_visible=True,
                sort_order=sort_order,
            ),
        )
        for category_key, attribute_key, sort_order in (
            ("root", "capacity", 10),
            ("root", "color", 20),
            ("shared", "capacity", 10),
            ("shared", "color", 20),
            ("leaf", "capacity", 10),
            ("leaf", "color", 20),
            ("leaf", "description", 30),
            ("leaf", "abs", 40),
            ("leaf", "esp", 50),
            ("attachmentChild", "capacity", 10),
            ("attachmentChild", "stabilizers", 20),
        )
    ]
    modification_values = [
        await _merge(session, value)
        for value in (
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["baseWhite"].id,
                attribute_id=attributes["capacity"].id,
                value_number=Decimal("1000"),
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["baseWhite"].id,
                attribute_id=attributes["color"].id,
                option_id=options["white"].id,
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["baseWhite"].id,
                attribute_id=attributes["description"].id,
                value_text=f"{context.prefix} магистральное шасси",
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["baseWhiteTwin"].id,
                attribute_id=attributes["capacity"].id,
                value_number=Decimal("1000"),
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["baseWhiteTwin"].id,
                attribute_id=attributes["color"].id,
                option_id=options["white"].id,
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["baseWhiteTwin"].id,
                attribute_id=attributes["description"].id,
                value_text=f"{context.prefix} магистральное шасси",
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["baseBlack"].id,
                attribute_id=attributes["capacity"].id,
                value_number=Decimal("1000"),
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["baseBlack"].id,
                attribute_id=attributes["color"].id,
                option_id=options["black"].id,
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["baseBlack"].id,
                attribute_id=attributes["description"].id,
                value_text=f"{context.prefix} магистральное шасси",
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["hours"].id,
                attribute_id=attributes["capacity"].id,
                value_number=Decimal("200"),
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["hours"].id,
                attribute_id=attributes["color"].id,
                option_id=options["black"].id,
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["hours"].id,
                attribute_id=attributes["description"].id,
                value_text="E2E землеройная техника",
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["attachment"].id,
                attribute_id=attributes["capacity"].id,
                value_number=Decimal("25"),
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["attachment"].id,
                attribute_id=attributes["color"].id,
                option_id=options["black"].id,
            ),
            SpecialEquipmentModificationAttributeValue(
                modification_id=modifications["attachment"].id,
                attribute_id=attributes["stabilizers"].id,
                value_boolean=True,
            ),
        )
    ]
    trim_attributes = [
        await _merge(
            session,
            SpecialEquipmentTrimAttribute(
                trim_id=trims[trim_key].id,
                attribute_id=attributes[attribute_key].id,
                group_id=groups["safety"].id,
                is_filterable=True,
                sort_order=sort_order,
            ),
        )
        for trim_key in trims
        for attribute_key, sort_order in (("abs", 10), ("esp", 20))
    ]
    trim_values = [
        await _merge(
            session,
            SpecialEquipmentTrimAttributeValue(
                trim_id=trims[trim_key].id,
                attribute_id=attributes[attribute_key].id,
                value_boolean=value_boolean,
            ),
        )
        for trim_key, abs_value, esp_value in (
            ("safetyAll", True, True),
            ("safetyAbsOnly", True, False),
            ("safetyNone", False, False),
        )
        for attribute_key, value_boolean in (("abs", abs_value), ("esp", esp_value))
    ]

    superstructures = {
        "crane": await _merge(
            session,
            SpecialEquipmentSuperstructure(
                id=context.entity_id("superstructure:crane"),
                model_id=models["craneBody"].id,
                modification_id=modifications["attachment"].id,
                code=f"{context.prefix}_SUPERSTRUCTURE_CRANE",
                name=f"{context.prefix} Кран-манипулятор",
                slug=f"{context.slug_prefix}-crane",
                is_active=True,
            ),
        ),
    }
    superstructure_attributes = [
        await _merge(
            session,
            SpecialEquipmentSuperstructureAttribute(
                superstructure_id=superstructures["crane"].id,
                attribute_id=attributes["stabilizers"].id,
                group_id=groups["technical"].id,
                is_required=True,
                is_visible=True,
                is_filterable=True,
                sort_order=10,
            ),
        ),
    ]
    superstructure_categories = [
        await _merge(
            session,
            SpecialEquipmentSuperstructureCategory(
                superstructure_id=superstructures["crane"].id,
                category_id=categories[cat_key].id,
            ),
        )
        for cat_key in ("leaf", "root")
    ]

    products: dict[str, SpecialEquipmentProduct] = {}
    for index, (key, spec) in enumerate(product_specs.items(), start=1):
        no_vin = bool(spec.get("no_vin", False))
        publication = str(spec.get("publication", "published"))
        is_kit = spec.get("kind") == "kit"
        if is_kit:
            modification_id = (
                modifications[str(spec["mod"])].id if "mod" in spec else None
            )
            model_id = models[str(spec.get("model", "kamaz1000"))].id
            superstructure_id = superstructures[str(spec["superstructure"])].id
            superstructure_name = str(spec["superstructure_name"])
            superstructure_manufacturer = str(spec["superstructure_manufacturer"])
            superstructure_modification_id = None
            trim_id = None
        else:
            modification_id = modifications[str(spec["mod"])].id
            model_id = None
            superstructure_id = None
            superstructure_name = None
            superstructure_manufacturer = None
            superstructure_modification_id = None
            trim_id = trims[str(spec["trim"])].id if "trim" in spec else None

        products[key] = await _merge(
            session,
            SpecialEquipmentProduct(
                id=_product_id(context, key),
                code=f"{context.prefix}_PRODUCT_{key.upper()}",
                modification_id=modification_id,
                model_id=model_id,
                superstructure_id=superstructure_id,
                superstructure_modification_id=superstructure_modification_id,
                superstructure_name=superstructure_name,
                superstructure_manufacturer=superstructure_manufacturer,
                trim_id=trim_id,
                seller_company_id=(
                    alternate_company.id
                    if spec.get("seller") == "alternate"
                    else company.id
                ),
                warehouse_id=None if no_vin else warehouse.id,
                slug=f"{context.slug_prefix}-product-{key.lower()}",
                description=f"{context.prefix} {spec.get('description', spec['title'])}",
                price=Decimal(str(spec["price"])),
                currency_code="RUB",
                manufacture_year=int(spec.get("year", 2025)),
                vin=None if no_vin else f"{context.uuid_namespace.hex[:10].upper()}{index:07d}",
                no_vin=no_vin,
                condition=str(spec.get("condition", "new")),
                owners_count=spec.get("owners"),
                mileage_km=spec.get("mileage"),
                engine_hours=spec.get("hours"),
                publication_status=publication,
                sale_status=str(spec.get("sale", "available")),
                published_at=(
                    published_at - timedelta(minutes=index // 2)
                    if publication == "published"
                    else None
                ),
            ),
        )
    product_categories: list[SpecialEquipmentProductCategory] = []
    for key, product in products.items():
        spec = product_specs[key]
        count_ordinal_value = spec.get("count_ordinal")
        category_keys = (
            [
                f"count{value}"
                for value in count_values
                if count_ordinal_value <= value
            ]
            if isinstance(count_ordinal_value, int)
            else [str(spec.get("category", "leaf"))]
        )
        product_categories.extend(
            [
                await _merge(
                    session,
                    SpecialEquipmentProductCategory(
                        product_id=product.id,
                        category_id=categories[category_key].id,
                    ),
                )
                for category_key in category_keys
            ]
        )
    compatibility = [
        await _merge(
            session,
            SpecialEquipmentProductAttachment(
                product_id=products[product_key].id,
                attachment_product_id=products[attachment_key].id,
                position=position,
                created_by=users["employee"].id,
            ),
        )
        for product_key in ("representative", "equivalent")
        for position, attachment_key in enumerate(
            (
                "attachmentStandalone",
                "attachmentCompatible",
                "unpublishedAvailable",
                "publishedUnavailable",
            )
        )
    ]
    product_chassis_values: list[SpecialEquipmentProductChassisValue] = [
        await _merge(
            session,
            SpecialEquipmentProductChassisValue(
                product_id=products["kitWithoutMod"].id,
                attribute_id=attributes["capacity"].id,
                value_number=Decimal("15"),
            ),
        ),
    ]
    product_superstructure_values: list[SpecialEquipmentProductSuperstructureValue] = [
        await _merge(
            session,
            SpecialEquipmentProductSuperstructureValue(
                product_id=products[kit_key].id,
                superstructure_id=superstructures["crane"].id,
                attribute_id=attributes["stabilizers"].id,
                value_boolean=True,
            ),
        )
        for kit_key in ("kitWithMod", "kitWithoutMod")
    ]
    root_cart = await _merge(
        session,
        SpecialEquipmentCartItem(
            id=context.entity_id("cart:root"),
            user_id=users["client"].id,
            product_id=products["representative"].id,
            quantity=1,
            is_selected=True,
            equipments=[],
            services=[],
        ),
    )
    await session.flush()
    cart_items = [
        root_cart,
        await _merge(
            session,
            SpecialEquipmentCartItem(
                id=context.entity_id("cart:attachment"),
                user_id=users["client"].id,
                product_id=products["attachmentCompatible"].id,
                quantity=1,
                parent_item_id=root_cart.id,
                is_selected=True,
                equipments=[],
                services=[],
            ),
        ),
        await _merge(
            session,
            SpecialEquipmentCartItem(
                id=context.entity_id("cart:standalone-attachment"),
                user_id=users["client"].id,
                product_id=products["attachmentCompatible"].id,
                quantity=1,
                is_selected=True,
                equipments=[],
                services=[],
            ),
        ),
    ]
    guest_transfer = await _merge(
        session,
        SpecialEquipmentGuestCartTransfer(
            user_id=users["client"].id,
            transfer_id=context.entity_id("cart:guest-transfer"),
            request_hash="0" * 64,
            result={"created": [str(root_cart.id)]},
        ),
    )
    await session.flush()

    guest_parent_id = context.entity_id("guest-cart:parent")
    guest_child_id = context.entity_id("guest-cart:child")
    guest_snapshot = {
        "version": 2,
        "transfer_id": str(context.entity_id("guest-cart:transfer")),
        "items": [
            {
                "local_id": str(guest_parent_id),
                "product_id": str(products["representative"].id),
                "quantity": 1,
                "parent_local_id": None,
                "is_selected": True,
                "comment": None,
                "equipments": [],
                "services": [],
            },
            {
                "local_id": str(guest_child_id),
                "product_id": str(products["attachmentCompatible"].id),
                "quantity": 1,
                "parent_local_id": str(guest_parent_id),
                "is_selected": True,
                "comment": None,
                "equipments": [],
                "services": [],
            },
        ],
    }

    import_paths = _create_import_artifacts(context, artifacts_dir)
    auth_paths = await _create_auth_artifacts(
        session,
        context,
        users,
        auth_state_dir,
    )
    parent_map: dict[str, list[str]] = {key: [] for key in categories}
    for parent, child, _sort_order in relation_specs:
        parent_map[child].append(parent)
    category_paths = {
        "root": [categories["root"].slug],
        "emptyRoot": [categories["emptyRoot"].slug],
        "dagParentA": [categories["dagParentA"].slug],
        "dagParentB": [categories["dagParentB"].slug],
        "shared": [categories["dagParentA"].slug, categories["shared"].slug],
        "siblingA": [categories["root"].slug, categories["siblingA"].slug],
        "siblingB": [categories["root"].slug, categories["siblingB"].slug],
        "attachmentRoot": [categories["attachmentRoot"].slug],
        "attachmentChild": [categories["attachmentRoot"].slug, categories["attachmentChild"].slug],
        "attachmentDescendant": [categories["attachmentRoot"].slug, categories["attachmentChild"].slug, categories["attachmentDescendant"].slug],
        "leaf": [categories["dagParentA"].slug, categories["shared"].slug, categories["leaf"].slug],
        "count25": [categories["count25"].slug],
    }
    auth_manifest = {
        key: {
            "id": str(user.id),
            "phone": user.phone,
            "role": user.role,
            "storage_state_path": auth_paths.get(key),
        }
        for key, user in users.items()
    }
    auth_manifest["client_owner"] = {
        "id": str(users["client"].id),
        "phone": users["client"].phone,
        "role": users["client"].role,
        "storage_state_path": auth_paths.get("client"),
    }
    return E2EFixtureManifest(
        schema_version=1,
        namespace=context.namespace,
        prefix=context.prefix,
        auth=auth_manifest,
        companies={
            "seller": {"id": str(company.id), "inn": company.inn, "name": company.name},
            "alternate": {
                "id": str(alternate_company.id),
                "inn": alternate_company.inn,
                "name": alternate_company.name,
            },
            "22294_full_target": {"id": str(company.id), "inn": company.inn, "name": company.name},
            "22294_sparse_optional": {"id": str(sparse_optional_company.id), "inn": sparse_optional_company.inn, "name": sparse_optional_company.name},
            "22294_duplicate_inn": {"id": str(alternate_company.id), "inn": alternate_company.inn, "name": alternate_company.name},
            "22294_legacy_invalid": {"id": str(legacy_invalid_company.id), "inn": legacy_invalid_company.inn, "name": legacy_invalid_company.name},
        },
        categories={
            key: _category_card(
                row,
                parent_ids=[categories[parent].id for parent in parent_map[key]],
                path=category_paths[key],
            )
            for key, row in categories.items()
            if key in manifest_category_keys
        },
        marks={key: _directory_card(row) for key, row in marks.items()},
        models={key: _directory_card(row) for key, row in models.items()},
        model_showcase=model_showcase,
        modifications={key: _directory_card(row) for key, row in modifications.items()},
        trims={key: _directory_card(row) for key, row in trims.items()},
        attribute_groups={key: _directory_card(row) for key, row in groups.items()},
        attributes={
            key: {"id": str(row.id), "code": row.code, "name": row.name, "data_type": row.data_type}
            for key, row in attributes.items()
        },
        options={
            key: {"id": str(row.id), "code": row.code, "name": row.name}
            for key, row in options.items()
        },
        products={
            key: _product_card(row, str(product_specs[key]["title"]))
            for key, row in products.items()
            if key in manifest_product_keys
        },
        fingerprint_probes={
            key: _product_card(products[key], str(product_specs[key]["title"]))
            for key in fingerprint_probe_keys
        },
        relations={
            "compatibility": [
                {"product_id": str(row.product_id), "attachment_product_id": str(row.attachment_product_id), "position": row.position}
                for row in compatibility
            ],
        },
        cart={
            "server_item_ids": [str(row.id) for row in cart_items],
            "guest_transfer_id": str(guest_transfer.transfer_id),
            "guest_storage_key": "carcraft:special-equipment:cart:v2",
            "guest_snapshot": guest_snapshot,
        },
        imports=import_paths,
        counts={
            "categories": len(categories),
            "category_relations": len(category_relations),
            "marks": len(marks),
            "models": len(models),
            "modifications": len(modifications),
            "modification_categories": len(modification_categories),
            "attribute_groups": len(groups),
            "attributes": len(attributes),
            "category_attributes": len(category_attributes),
            "modification_values": len(modification_values),
            "trim_attributes": len(trim_attributes),
            "trim_values": len(trim_values),
            "products": len(products),
            "product_categories": len(product_categories),
            "compatibility": len(compatibility),
            "superstructures": len(superstructures),
            "superstructure_attributes": len(superstructure_attributes),
            "superstructure_categories": len(superstructure_categories),
            "product_chassis_values": len(product_chassis_values),
            "product_superstructure_values": len(product_superstructure_values),
            "units": len(units),
            "cart_items": len(cart_items),
        },
    )


def _save_v2_workbook(
    path: Path,
    *,
    mode: ImportMode,
    rows: Mapping[str, Sequence[Sequence[Any]]],
) -> None:
    workbook = load_workbook(BytesIO(build_template_v7(mode=mode)))
    try:
        for family, family_rows in rows.items():
            sheet = workbook[DATA_SHEET_NAMES[family]]
            expected_len = len(DATA_SHEET_HEADERS[family])
            for row in family_rows:
                row_tuple = tuple(row)
                if len(row_tuple) == expected_len + 1:
                    row_tuple = row_tuple[1:]
                sheet.append(row_tuple)
        workbook.save(path)
    finally:
        workbook.close()


def _create_import_artifacts(
    context: FixtureContext,
    artifacts_dir: Path | None,
) -> dict[str, str]:
    if artifacts_dir is None:
        return {}
    artifacts_dir.mkdir(parents=True, exist_ok=True)
    seller_inn = f"99{context.uuid_namespace.int % 100_000_000:08d}"
    import_prefix = f"{context.prefix}_IMPORT"
    common_rows: dict[str, Sequence[Sequence[Any]]] = {
        "units": [(None, f"{import_prefix}_UNIT_T", "т", "Да")],
        "marks": [(None, f"{import_prefix}_MARK", "E2E Марка импорта", "Да")],
        "models": [(None, f"{import_prefix}_MODEL", "E2E Модель импорта", f"{import_prefix}_MARK", "Да")],
        "modifications": [
            (None, f"{import_prefix}_MOD", "E2E Модификация импорта", f"{import_prefix}_MODEL", 2024, 2026, "Да"),
            (None, f"{import_prefix}_ATTACHMENT_MOD", "E2E Надстройка импорта", f"{import_prefix}_MODEL", 2024, 2026, "Да"),
        ],
        "trims": [],
        "modification_categories": [
            (None, f"{import_prefix}_MOD", f"{import_prefix}_CATEGORY", 0, "Да"),
            (None, f"{import_prefix}_ATTACHMENT_MOD", f"{import_prefix}_ATTACHMENT_CATEGORY", 0, "Да"),
        ],
        "modification_attribute_values": [(None, f"{import_prefix}_MOD", f"{import_prefix}_ATTRIBUTE", None, 25)],
        "trim_attributes": [],
        "trim_attribute_values": [],
        "categories": [
            (None, f"{import_prefix}_CATEGORY", "E2E Категория импорта", None, "Пробег", "Нет", 10, "Да"),
            (None, f"{import_prefix}_ATTACHMENT_CATEGORY", "E2E Надстройки импорта", None, "Моточасы", "Да", 20, "Да"),
        ],
        "category_relations": [],
        "attribute_groups": [(None, f"{import_prefix}_GROUP", "E2E Группа", 10, "Да")],
        "attributes": [(None, f"{import_prefix}_ATTRIBUTE", "E2E Мощность", f"{import_prefix}_GROUP", "Число", f"{import_prefix}_UNIT_T", "Диапазон", "Да")],
        "attribute_options": [],
        "category_attributes": [(None, f"{import_prefix}_CATEGORY", f"{import_prefix}_ATTRIBUTE", f"{import_prefix}_GROUP", "Да", "Да", "Да", 10)],
        "colors": [],
        "superstructures": [
            (None, f"{import_prefix}_SUPERSTRUCTURE", "E2E Тип надстройки", f"{import_prefix}_MODEL", None, "Да"),
        ],
        "superstructure_attributes": [
            (None, f"{import_prefix}_SUPERSTRUCTURE", f"{import_prefix}_GROUP", f"{import_prefix}_ATTRIBUTE", "Нет", "Да", "Да", 10),
        ],
        "products": [
            (None, f"{import_prefix}_PRODUCT", f"{import_prefix}_MOD", None, None, None, None, None, None, seller_inn, "Новое", 2026, "E2E объявление импорта", None, None, 1000000, None, None, None, "RUB", None, "Нет", f"{context.uuid_namespace.hex[:10].upper()}IMP0001", None, None, "Черновик", "Доступно", None, None),
            (None, f"{import_prefix}_ATTACHMENT", f"{import_prefix}_ATTACHMENT_MOD", None, None, None, None, None, None, seller_inn, "Новое", 2026, "E2E надстройка", None, None, 200000, None, None, None, "RUB", None, "Нет", f"{context.uuid_namespace.hex[:10].upper()}IMP0003", None, None, "Черновик", "Доступно", None, None),
            (None, f"{import_prefix}_KIT", None, None, f"{import_prefix}_MODEL", f"{import_prefix}_SUPERSTRUCTURE", None, "КМУ Название", "КМУ Производитель", seller_inn, "Новое", 2026, "E2E комплект импорта", None, None, 1200000, None, None, None, "RUB", None, "Нет", f"{context.uuid_namespace.hex[:10].upper()}IMP0004", None, None, "Черновик", "Доступно", None, None),
        ],
        "product_categories": [
            (None, f"{import_prefix}_PRODUCT", f"{import_prefix}_CATEGORY"),
            (None, f"{import_prefix}_ATTACHMENT", f"{import_prefix}_ATTACHMENT_CATEGORY"),
            (None, f"{import_prefix}_KIT", f"{import_prefix}_CATEGORY"),
        ],
        "product_chassis_values": [
            (None, f"{import_prefix}_KIT", f"{import_prefix}_ATTRIBUTE", None, 30),
        ],
        "product_superstructure_values": [
            (None, f"{import_prefix}_KIT", f"{import_prefix}_ATTRIBUTE", None, 30),
        ],
        "product_attachments": [
            (None, f"{import_prefix}_PRODUCT", f"{import_prefix}_ATTACHMENT", 0)
        ],
    }
    full_path = (artifacts_dir / "v2-full.xlsx").resolve()
    _save_v2_workbook(full_path, mode=ImportMode.FULL_SNAPSHOT, rows=common_rows)

    patch_path = (artifacts_dir / "v2-patch.xlsx").resolve()
    _save_v2_workbook(
        patch_path,
        mode=ImportMode.PATCH,
        rows={
            "products": [
                (
                    "Задать",
                    f"{context.prefix}_PRODUCT_REPRESENTATIVE",
                    f"{context.prefix}_MODIFICATION_BASEWHITE",
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    seller_inn,
                    None,
                    None,
                    "E2E описание обновлено импортом",
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                )
            ],
            "product_attachments": [
                (
                    "Удалить",
                    f"{context.prefix}_PRODUCT_REPRESENTATIVE",
                    f"{context.prefix}_PRODUCT_ATTACHMENTSTANDALONE",
                    None,
                ),
                (
                    "Задать",
                    f"{context.prefix}_PRODUCT_REPRESENTATIVE",
                    f"{context.prefix}_PRODUCT_ATTACHMENTCOMPATIBLE",
                    7,
                ),
            ],
        },
    )

    empty_relations_path = (artifacts_dir / "v2-empty-relations.xlsx").resolve()
    _save_v2_workbook(
        empty_relations_path,
        mode=ImportMode.FULL_SNAPSHOT,
        rows={
            **common_rows,
            "product_attachments": [],
        },
    )

    mixed_atomic_path = (artifacts_dir / "v2-mixed-atomic.xlsx").resolve()
    atomic_rows = dict(common_rows)
    atomic_rows["categories"] = [
        *common_rows["categories"],
        (None, "INVALID CODE WITH SPACES", "Ошибочная категория", None, "Пробег", "Нет", 30, "Да"),
    ]
    _save_v2_workbook(
        mixed_atomic_path,
        mode=ImportMode.FULL_SNAPSHOT,
        rows=atomic_rows,
    )

    mixed_best_effort_path = (artifacts_dir / "v2-mixed-best-effort.xlsx").resolve()
    _save_v2_workbook(
        mixed_best_effort_path,
        mode=ImportMode.PATCH,
        rows={
            "categories": [
                ("Добавить", f"{import_prefix}_BEST_EFFORT_OK", "Корректная категория", None, "Пробег", "Нет", 40, "Да"),
                ("Добавить", "INVALID CODE WITH SPACES", "Ошибочная категория", None, "Пробег", "Нет", 50, "Да"),
            ]
        },
    )

    old_v2_path = (artifacts_dir / "v2-old-contract.xlsx").resolve()
    old_v2 = Workbook()
    try:
        manifest = old_v2.active
        manifest.title = PARAMETERS_SHEET_NAME
        manifest.append(MANIFEST_HEADERS)
        manifest.append(("Версия шаблона", "2"))
        manifest.append(("Дата формирования", "2026-01-01T00:00:00Z"))
        manifest.append(("Маркер очистки", NULL_TOKEN))
        old_v2.save(old_v2_path)
    finally:
        old_v2.close()
    return {
        "v2Full": str(full_path),
        "v2Patch": str(patch_path),
        "v2EmptyRelations": str(empty_relations_path),
        "v2OldContract": str(old_v2_path),
        "v2MixedAtomic": str(mixed_atomic_path),
        "v2MixedBestEffort": str(mixed_best_effort_path),
    }


def _storage_state(
    *,
    access_token: str,
    refresh_token: str,
    csrf_token: str,
) -> dict[str, Any]:
    base_url = settings.e2e_base_url.rstrip("/")
    parsed = urlsplit(base_url)
    if parsed.scheme not in {"http", "https"} or parsed.hostname is None:
        raise ValueError("E2E_BASE_URL must be an absolute HTTP(S) URL")
    expires = int((datetime.now(UTC) + timedelta(days=settings.refresh_token_expiry_days)).timestamp())
    secure = parsed.scheme == "https"
    same_site = {"lax": "Lax", "strict": "Strict", "none": "None"}[
        settings.cookie_samesite
    ]
    cookies = [
        {
            "name": "accessToken",
            "value": access_token,
            "domain": parsed.hostname,
            "path": "/",
            "expires": expires,
            "httpOnly": True,
            "secure": secure,
            "sameSite": same_site,
        },
        {
            "name": "refreshToken",
            "value": refresh_token,
            "domain": parsed.hostname,
            "path": "/api/v1/auth",
            "expires": expires,
            "httpOnly": True,
            "secure": secure,
            "sameSite": same_site,
        },
    ]
    if settings.csrf_enabled:
        cookies.append(
            {
                "name": settings.csrf_cookie_name,
                "value": csrf_token,
                "domain": parsed.hostname,
                "path": "/",
                "expires": expires,
                "httpOnly": False,
                "secure": secure,
                "sameSite": same_site,
            }
        )
    return {"cookies": cookies, "origins": []}


def _write_private_json(path: Path, payload: Mapping[str, Any]) -> None:
    """Atomically publish authentication material with owner-only permissions."""

    path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    path.parent.chmod(0o700)
    descriptor, temporary_name = tempfile.mkstemp(
        dir=path.parent,
        prefix=f".{path.name}.",
    )
    temporary_path = Path(temporary_name)
    try:
        os.fchmod(descriptor, 0o600)
        with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
            json.dump(payload, stream, ensure_ascii=False)
            stream.write("\n")
            stream.flush()
            os.fsync(stream.fileno())
        temporary_path.replace(path)
        path.chmod(0o600)
    except BaseException:
        with suppress(OSError):
            os.close(descriptor)
        temporary_path.unlink(missing_ok=True)
        raise


async def _create_auth_artifacts(
    session: AsyncSession,
    context: FixtureContext,
    users: Mapping[str, User],
    auth_state_dir: Path | None,
) -> dict[str, str]:
    if auth_state_dir is None:
        return {}
    paths: dict[str, str] = {}
    for key, user in users.items():
        session_id = context.entity_id(f"auth-session:{key}")
        access_token, refresh_token = generate_tokens(
            user.id,
            user.role,
            user.company_id,
            refresh_session_id=session_id,
        )
        await _merge(
            session,
            UserSession(
                id=session_id,
                user_id=user.id,
                refresh_token_hash=hash_refresh_token(refresh_token),
                expires_at=datetime.now(UTC)
                + timedelta(days=settings.refresh_token_expiry_days),
                user_agent="CarCraft E2E fixture",
            ),
        )
        path = (auth_state_dir / f"{key}-storage-state.json").resolve()
        _write_private_json(
            path,
            _storage_state(
                access_token=access_token,
                refresh_token=refresh_token,
                csrf_token=context.entity_id(f"csrf:{key}").hex,
            ),
        )
        paths[key] = str(path)
    return paths


async def _apply_fixture(
    *,
    namespace: str,
    reset: bool,
    artifacts_dir: Path | None,
    auth_state_dir: Path | None = None,
) -> E2EFixtureManifest:
    async with AsyncSessionLocal() as session, session.begin():
        manifest = await seed_special_equipment_e2e(
            session,
            namespace=namespace,
            reset=reset,
            artifacts_dir=artifacts_dir,
            auth_state_dir=auth_state_dir,
        )
        await upload_model_showcase_e2e_images(manifest.model_showcase)
        return manifest


def _arguments(argv: Sequence[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--apply",
        action="store_true",
        required=True,
        help="Explicitly confirm database mutation.",
    )
    parser.add_argument("--namespace", required=True, help="Isolated E2E worker namespace.")
    parser.add_argument(
        "--reset",
        action="store_true",
        help="Reset only rows owned by this namespace before recreating them.",
    )
    parser.add_argument("--artifacts-dir", type=Path)
    parser.add_argument("--auth-state-dir", type=Path)
    output = parser.add_mutually_exclusive_group()
    output.add_argument("--manifest-path", type=Path)
    output.add_argument("--json", action="store_true")
    return parser.parse_args(argv)


def _write_manifest(path: Path, manifest: E2EFixtureManifest) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(manifest.to_dict(), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main(argv: Sequence[str] | None = None) -> int:
    args = _arguments(argv)
    manifest = asyncio.run(
        _apply_fixture(
            namespace=args.namespace,
            reset=args.reset,
            artifacts_dir=args.artifacts_dir,
            auth_state_dir=args.auth_state_dir,
        )
    )
    if args.manifest_path is not None:
        _write_manifest(args.manifest_path, manifest)
    elif args.json:
        sys.stdout.write(json.dumps(manifest.to_dict(), ensure_ascii=False) + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
