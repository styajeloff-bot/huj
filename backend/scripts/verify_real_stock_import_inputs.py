from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

EXPORT = Path("~/projects/carcraft-leadgenerator/leasing_export").expanduser()
FULL = EXPORT / "import" / "full"
CATALOG_DIR = FULL / "catalog_real_stock"


def _read_lca_vehicle_ids() -> set[str]:
    with (FULL / "lca.csv").open("r", encoding="utf-8-sig", newline="") as f:
        return {
            row["vehicle_id"]
            for row in csv.DictReader(f, delimiter=";")
            if row.get("vehicle_id")
        }


def _read_excel_vehicle_ids() -> tuple[set[str], set[str], int]:
    ids: set[str] = set()
    vins: set[str] = set()
    rows = 0
    for path in sorted(CATALOG_DIR.glob("part-*.xlsx")):
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
            index = {header: i for i, header in enumerate(headers)}
            for values in ws.iter_rows(min_row=2, values_only=True):
                rows += 1
                ids.add(str(values[index["ID"]]))
                vins.add(str(values[index["VIN"]]))
                for required in ["Статус", "created_at", "dealer_id", "warehouse_id"]:
                    value = values[index[required]]
                    if value is None or str(value).strip() == "":
                        raise RuntimeError(f"{path}: empty {required} at row {rows + 1}")
        finally:
            wb.close()
    return ids, vins, rows


def main() -> None:
    lca_ids = _read_lca_vehicle_ids()
    excel_ids, vins, row_count = _read_excel_vehicle_ids()
    missing = lca_ids - excel_ids
    if missing:
        raise RuntimeError(f"{len(missing)} LCA vehicle IDs missing from Excel chunks")
    if len(vins) != row_count:
        raise RuntimeError(f"VINs are not unique: rows={row_count} unique_vins={len(vins)}")
    print(f"OK rows={row_count} excel_ids={len(excel_ids)} lca_ids={len(lca_ids)} vins={len(vins)}")  # noqa: T201


if __name__ == "__main__":
    main()
