"""HTTP/PostgreSQL/S3 regressions with a disposable instrumented app process.

Run: compose run --rm --no-deps backend python /e2e/transactions.py.
Only fixture-prefixed records are written. Lost ACK is fault injection after a
real commit, not a claim that a physical network outage was reproduced.
"""
import asyncio
import io
import json
import os
import secrets
import socket
import subprocess
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path
from uuid import UUID, uuid4
from zoneinfo import ZoneInfo

import httpx
from runtime import ROOT, guard, pdf, progress, require

BASE = "/api/v1/document-registry"
PREFIX = "22296-audit-tx"


class InstrumentedAPI:
    def __init__(self, state, port, key):
        self.state, self.port, self.key = state, port, key

    async def request(self, role, method, path, expected=200, raw=False, phase="", target="", **kwargs):
        actor = self.state["users"].get(role)
        cookies = {"accessToken": actor["access_token"], "csrfToken": actor["csrf"]} if actor else {}
        headers = {"Origin": "http://localhost:18296", "X-Audit-Tx-Key": self.key,
                   "X-Audit-Tx-Phase": phase, "X-Audit-Tx-Document": target}
        if actor:
            headers["X-CSRF-Token"] = actor["csrf"]
        headers.update(kwargs.pop("headers", {}))
        async with httpx.AsyncClient(base_url=f"http://127.0.0.1:{self.port}", trust_env=False,
                                     cookies=cookies, headers=headers, timeout=60) as client:
            response = await client.request(method, path, **kwargs)
        allowed = (expected,) if isinstance(expected, int) else expected
        require(response.status_code in allowed, f"{role} {method} {path}: HTTP {response.status_code}; {response.text[:300]}")
        return response if raw else (response.json() if response.content else None)

    async def control(self, action, document_id=None):
        return await self.request("admin", "GET", "/__registry_tx/control", params={"action": action, **({"document_id": document_id} if document_id else {})})


async def database_facts(document):
    import aioboto3
    from sqlalchemy import select, func
    from infrastructure.database import AsyncSessionLocal
    from infrastructure.models.document_registry import reference_documents as docs, reference_document_versions as versions, reference_document_files as files
    settings = guard()
    identifier = UUID(document["id"])
    async with AsyncSessionLocal() as session:
        require(await session.scalar(select(docs.c.contract_number).where(docs.c.id == identifier)) == document["contract_number"], "Wrong fixture row")
        count = await session.scalar(select(func.count()).select_from(versions).where(versions.c.document_id == identifier))
        current = await session.scalar(select(func.count()).select_from(versions).where(versions.c.document_id == identifier, versions.c.is_current.is_(True)))
        rows = (await session.execute(select(files.c.s3_key).join(versions, versions.c.id == files.c.version_id).where(versions.c.document_id == identifier))).scalars().all()
    async with aioboto3.Session().client("s3", endpoint_url=settings.s3_endpoint, region_name=settings.s3_region,
        aws_access_key_id=settings.s3_access_key_id, aws_secret_access_key=settings.s3_secret_access_key) as s3:
        objects = await s3.list_objects_v2(Bucket=settings.s3_bucket, Prefix=f"document-registry/{identifier}/")
    keys = {row["Key"] for row in objects.get("Contents", [])}
    return {"committed_versions": count, "current_version_count": current, "committed_file_rows": len(rows),
            "referenced_object_count_present": len(set(rows) & keys), "referenced_object_count_missing": len(set(rows) - keys), "storage_objects": len(keys)}


class Scenarios:
    def __init__(self, api):
        self.api = api
        self.run = uuid4().hex[:10]
        self.day = datetime.now(ZoneInfo("Europe/Moscow")).date()
        self.owned = []

    def related(self, alias=None):
        return {"platform_ml": False, "leasing_company_ids": [], "dealer_company_ids": [self.api.state["companies"][alias]["company_id"]] if alias else [], "distributor_company_ids": []}

    def metadata(self, label, alias="dealer"):
        return {"document_type": "contract", "contract_number": f"{PREFIX}-{self.run}-{label}",
            "name": f"{PREFIX}-{self.run}-{label}-v1", "valid_from": str(self.day), "valid_to": str(self.day+timedelta(days=100)),
            "participants": {"platform_ml": True, "leasing_company_ids": [self.api.state["companies"]["leasing"]["leasing_company_id"]]}, "related_companies": self.related(alias)}

    def version(self, document, alias="dealer2"):
        return {"expected_current_version_id": document["current_version"]["id"], "name": document["name"] + "-v2",
                "valid_from": str(self.day), "valid_to": str(self.day+timedelta(days=100)),
                "related_companies": self.related(alias), "retained_file_ids": []}

    def upload(self, label):
        return [("files", (f"{PREFIX}-{self.run}-{label}.pdf", pdf(), "application/pdf"))]

    async def create(self, label):
        doc = await self.api.request("admin", "POST", BASE + "/documents", expected=201,
            data={"metadata": json.dumps(self.metadata(label))}, files=self.upload(label))
        self.owned.append(doc["id"])
        return doc

    async def finish_tasks(self, tasks):
        tasks = list(tasks)
        for task in tasks:
            if not task.done():
                task.cancel()
        await asyncio.gather(*tasks, return_exceptions=True)

    async def release_and_finish(self, document_id, tasks):
        try:
            await self.api.control("release", document_id)
        finally:
            await self.finish_tasks(tasks)

    async def wait_reached(self, document_id, count, tasks=()):
        for _ in range(300):
            for task in tasks:
                if task.done():
                    task.result()
                    raise AssertionError("HTTP request completed before reaching its expected barrier")
            if (await self.api.control("status", document_id))["reached"] >= count:
                return
            await asyncio.sleep(.05)
        raise AssertionError("Real HTTP requests did not reach the barrier")

    async def projection_race(self):
        v1 = await self.create("projection")
        path = BASE + "/documents/" + v1["id"]
        await self.api.control("arm", v1["id"])
        tasks = [asyncio.create_task(self.api.request(role, "GET", path, phase="pause_projection", target=v1["id"])) for role in ("admin", "dealer")]
        try:
            await self.wait_reached(v1["id"], 2, tasks)
            v2 = await self.api.request("admin", "POST", path + "/versions", expected=201,
                data={"metadata": json.dumps(self.version(v1))}, files=self.upload("projection-v2"))
            await self.api.control("release", v1["id"])
            responses = await asyncio.gather(*tasks)
        finally:
            await self.release_and_finish(v1["id"], tasks)
        for doc in responses:
            require(doc == v1, "ACL/document/version/related/files did not share the original snapshot")
        await self.api.request("dealer", "GET", path, expected=404)
        await self.api.request("dealer", "GET", v2["current_version"]["files"][0]["download_url"], expected=404)
        require((await self.api.request("outsider", "GET", path)) == v2, "New partner cannot read the committed version")
        progress("registry_projection_snapshot_passed", delayed_reads=2, revoked_partner_fresh_status=404)

    async def route_snapshots(self):
        v1 = await self.create("all-routes")
        path = BASE + "/documents/" + v1["id"]
        company = self.api.state["companies"]
        programs = "/api/v1/admin/monetization/programs"
        program = await self.api.request("admin", "POST", programs, expected=201, json={
            "status": "inactive", "name": v1["name"], "leasing_company_id": company["leasing"]["leasing_company_id"],
            "dealer_company_id": company["dealer"]["company_id"], "period_start": str(self.day),
            "sources": [{"source_type": "platform", "expenses": [{"participant_type": "platform", "base_type": "none", "calc_type": "amount", "value": "100"}], "incomes": []}],
            "reference_document_ids": [v1["id"]]})
        search = {"search": v1["contract_number"]}
        routes = [
            ("history", "GET", path + "/versions", {}),
            ("groups", "GET", BASE + "/groups", {"params": search}),
            ("children", "GET", BASE + f"/groups/{v1['group_id']}/documents", {}),
            ("table", "GET", BASE + "/table", {"params": search}),
            ("export", "GET", BASE + "/table/export", {"params": search, "raw": True}),
            ("download", "GET", v1["current_version"]["files"][0]["download_url"], {"raw": True}),
            ("candidates", "POST", BASE + "/monetization-candidates", {"json": {"program_id": program["id"]}}),
            ("program", "GET", programs + "/" + program["id"], {}),
            ("references", "GET", programs + "/" + program["id"] + "/reference-documents", {}),
        ]
        before = {name: await self.api.request("dealer", method, url, **kwargs) for name, method, url, kwargs in routes}
        await self.api.control("arm", v1["id"])
        tasks = {name: asyncio.create_task(self.api.request("dealer", method, url, phase="pause_actor", target=v1["id"], **kwargs)) for name, method, url, kwargs in routes}
        try:
            await self.wait_reached(v1["id"], len(tasks), tasks.values())
            await self.api.request("admin", "POST", path + "/versions", expected=201,
                data={"metadata": json.dumps(self.version(v1))}, files=self.upload("all-routes-v2"))
            await self.api.control("release", v1["id"])
            responses = dict(zip(tasks, await asyncio.gather(*tasks.values()), strict=True))
        finally:
            await self.release_and_finish(v1["id"], tasks.values())
        from openpyxl import load_workbook
        for name, actual in responses.items():
            expected = before[name]
            if name == "export":
                actual = list(load_workbook(io.BytesIO(actual.content)).active.values)
                expected = list(load_workbook(io.BytesIO(expected.content)).active.values)
            elif name == "download":
                actual, expected = actual.content, expected.content
            elif name == "candidates":
                # Other simultaneous E2E runs share company context. Only our
                # document/group participates in this controlled interleaving.
                def owned_projection(result):
                    return {"items": [item for item in result["items"] if item["id"] == v1["id"]],
                        "groups": [group for group in result["groups"] if group["group_id"] == v1["group_id"]]}
                actual, expected = owned_projection(actual), owned_projection(expected)
                require(len(expected["items"]) == len(expected["groups"]) == 1, "Own candidate missing before snapshot race")
            require(actual == expected, name + " changed after the actor snapshot was acquired")
        await self.api.request("dealer", "GET", path + "/versions", expected=404)
        await self.api.request("dealer", "GET", routes[5][2], expected=404)
        for suffix in ("", "/reference-documents"):
            current = await self.api.request("dealer", "GET", programs + "/" + program["id"] + suffix)
            require(not current["items" if suffix else "reference_documents"], "Fresh monetization response leaked revoked document")
        progress("registry_all_read_snapshots_passed", routes=[row[0] for row in routes])
        await self.write_projection_locks(program["id"], v1["id"])

    async def write_projection_locks(self, program_id, document_id):
        program_path = "/api/v1/admin/monetization/programs/" + program_id
        document_path = BASE + "/documents/" + document_id
        for label, route, payload in (
            ("status", program_path, {"status": "inactive"}),
            ("references", program_path + "/reference-documents", {"document_ids": [document_id]}),
        ):
            before = await self.api.request("admin", "GET", document_path)
            await self.api.control("arm", document_id)
            reader = asyncio.create_task(self.api.request("admin", "PATCH", route, json=payload,
                phase="pause_projection", target=document_id))
            tasks = [reader]
            try:
                await self.wait_reached(document_id, 1, tasks)
                writer = asyncio.create_task(self.api.request("admin", "POST", document_path + "/versions", expected=201,
                    phase="observe_write_lock", target=document_id, data={"metadata": json.dumps(self.version(before))}, files=self.upload("write-projection-" + label)))
                tasks.append(writer)
                for _ in range(300):
                    if writer.done():
                        writer.result()
                        raise AssertionError(label + " writer completed without waiting for the projection lock")
                    if (await self.api.control("status", document_id))["writer_waiting_for_lock"]:
                        break
                    await asyncio.sleep(.05)
                else:
                    raise AssertionError(label + " response did not protect document projection with a PostgreSQL lock")
                await self.api.control("release", document_id)
                response, after = await asyncio.gather(reader, writer)
                require(response["reference_documents" if label == "status" else "items"] == [before], label + " returned inconsistent references")
                require(after["current_version"]["version_number"] == before["current_version"]["version_number"] + 1, "Concurrent version never committed")
            finally:
                await self.release_and_finish(document_id, tasks)
        progress("registry_write_projection_locks_passed", routes=["program-status", "reference-links"], actual_postgresql_lock_waits=2)

    async def commit_failures(self):
        metadata = self.metadata("commit-create")
        await self.api.request("admin", "POST", BASE + "/documents", expected=500, raw=True, phase="commit_ack_loss",
            data={"metadata": json.dumps(metadata)}, files=self.upload("commit-create"))
        groups = await self.api.request("admin", "GET", BASE + "/groups", params={"search": metadata["contract_number"]})
        require(groups["pagination"]["total"] == 1, "Real commit did not persist the new document")
        created = groups["items"][0]["main_document"]
        self.owned.append(created["id"])
        await self.api.request("admin", "GET", created["current_version"]["files"][0]["download_url"], raw=True)
        facts = await database_facts(created)
        require(facts["committed_versions"] == 1 and facts["referenced_object_count_missing"] == 0, "Lost ACK removed committed create objects")
        v1 = await self.create("commit-version")
        path = BASE + "/documents/" + v1["id"]
        body = self.version(v1)
        body["retained_file_ids"] = [v1["current_version"]["files"][0]["id"]]
        await self.api.request("admin", "POST", path + "/versions", expected=500, raw=True, phase="commit_ack_loss",
            data={"metadata": json.dumps(body)}, files=self.upload("commit-version-v2"))
        current = await self.api.request("admin", "GET", path)
        history = await self.api.request("admin", "GET", path + "/versions")
        require(current["version_count"] == 2 and len(history["items"]) == 2, "Lost ACK version did not commit")
        for version in history["items"]:
            for file in version["files"]:
                await self.api.request("admin", "GET", file["download_url"], raw=True)
        before = await database_facts(current)
        require(before["referenced_object_count_missing"] == 0 and before["storage_objects"] == 2, "Lost ACK removed committed version objects")
        await self.api.request("admin", "POST", path + "/versions", expected=500, raw=True, phase="before_commit_failure",
            data={"metadata": json.dumps(self.version(current))}, files=self.upload("rollback"))
        require(await self.api.request("admin", "GET", path) == current, "Precommit failure changed the current document")
        require(await database_facts(current) == before, "Precommit rollback left SQL/S3 partial state")
        faults = (await self.api.control("faults"))["faults"]
        require(sum(row.get("real_commit_returned", False) for row in faults) == 2 and sum(row.get("before_commit", False) for row in faults) == 1, "Fault injection did not execute")
        progress("registry_commit_failures_passed", real_commits_then_injected_ack_loss=2, committed_files_readable=True, precommit_objects_removed=True)


async def run(port, key):
    state = json.loads((ROOT / "state.secret.json").read_text())
    api = InstrumentedAPI(state, port, key)
    for _ in range(200):
        try:
            await api.request("anonymous", "GET", "/api/v1/health")
            break
        except httpx.ConnectError:
            await asyncio.sleep(.1)
    else:
        raise AssertionError("Disposable HTTP process did not start")
    cases = Scenarios(api)
    await cases.projection_race()
    await cases.route_snapshots()
    await cases.commit_failures()
    from upload_boundaries import verify_upload_boundaries, verify_openapi
    await verify_upload_boundaries(api)
    await verify_openapi(api)
    progress("registry_transaction_regressions_passed", own_documents=cases.owned, physical_network_outage=False)


def main():
    guard()
    with tempfile.TemporaryDirectory(prefix="22296-registry-tx-") as directory:
        root = Path(directory)
        key = secrets.token_urlsafe(32)
        control = root / "control-key"
        control.write_text(key)
        control.chmod(0o600)
        with socket.socket() as listener:
            listener.bind(("127.0.0.1", 0))
            port = listener.getsockname()[1]
        env = {**os.environ, "REGISTRY_TX_CONTROL_FILE": str(control), "PYTHONPATH": "/e2e:/source"}
        with (root / "server.log").open("w+") as log:
            process = subprocess.Popen([sys.executable, "-m", "uvicorn", "transaction_server:app", "--host", "127.0.0.1", "--port", str(port), "--lifespan", "off"], env=env, stdout=log, stderr=subprocess.STDOUT)
            try:
                asyncio.run(run(port, key))
            finally:
                process.terminate()
                try:
                    process.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.wait(timeout=5)
                log.seek(0)
                failures = []
                for line in log:
                    try:
                        event = json.loads(line)
                    except ValueError:
                        continue
                    if event.get("event") == "http.request.failed":
                        failures.append({key: event.get(key) for key in ("http_method", "http_route", "http_status_code", "error_type", "error_message")})
                progress("registry_disposable_process_stopped", request_failures=failures)


if __name__ == "__main__":
    main()
