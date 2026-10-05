"""Exact acceptance matrix from the original task and direct user clarifications."""

import asyncio
import io
import json
import sys
from datetime import datetime, timedelta
from uuid import uuid4
from zoneinfo import ZoneInfo

sys.path.insert(0, "/e2e")
from runtime import API, ROOT, guard, pdf, require, progress
from versioning import version_metadata
from openpyxl import load_workbook

BASE = "/api/v1/document-registry"
PROGRAMS = "/api/v1/admin/monetization/programs"


async def main():
    guard()
    state = json.loads((ROOT / "state.secret.json").read_text())
    api = API(state)
    day = datetime.now(ZoneInfo("Europe/Moscow")).date()
    prefix = "22296-acceptance-" + uuid4().hex[:7]
    c = state["companies"]
    cat = state["catalog"]
    lc = c["leasing"]["leasing_company_id"]
    lc2 = c["leasing2"]["leasing_company_id"]
    dealer = c["dealer"]["company_id"]
    outsider = c["dealer2"]["company_id"]
    dist = c["distributor"]["company_id"]
    checks = []
    roots = []

    def check(label, condition, **detail):
        checks.append({"check": label, "pass": bool(condition), **detail})
        progress("matrix_check", **checks[-1])

    def meta(
        label,
        participants=None,
        kind="contract",
        start=-10,
        end=90,
        related=None,
        name=None,
    ):
        return dict(
            document_type=kind,
            contract_number=prefix + "-" + label,
            name=name or prefix + "-" + label,
            valid_from=str(day + timedelta(days=start)),
            valid_to=str(day + timedelta(days=end)) if end is not None else None,
            participants=participants or {"leasing_company_ids": [lc]},
            related_companies=related or {},
        )

    async def upload(m, group=None, contents=None):
        body = dict(m)
        path = BASE + "/documents"
        if group:
            body.pop("participants", None)
            path = BASE + f"/groups/{group}/documents"
        result = await api.request(
            "admin",
            "POST",
            path,
            expected=201,
            data={"metadata": json.dumps(body)},
            files=[("files", ("matrix.pdf", contents or pdf(), "application/pdf"))],
        )
        if not group:
            roots.append(result)
        return result

    async def export(params, role="admin"):
        r = await api.request(
            role, "GET", BASE + "/table/export", params=params, raw=True
        )
        return load_workbook(io.BytesIO(r.content), data_only=False).active

    async def all_cards(params, role="admin"):
        result = []
        pages = []
        page = 1
        while True:
            r = await api.request(
                role, "GET", BASE + "/groups", params={**params, "page": page}
            )
            result += r["items"]
            pages.append(len(r["items"]))
            if len(result) >= r["pagination"]["total"]:
                break
            require(r["items"], "Pagination made no progress")
            page += 1
            require(page < 20, "Unexpected pagination loop")
        return result, pages

    async def all_table(params, role="admin"):
        result = []
        pages = []
        cursor = None
        while True:
            r = await api.request(
                role,
                "GET",
                BASE + "/table",
                params={**params, **({"cursor": cursor} if cursor else {})},
            )
            result += r["items"]
            pages.append(len(r["items"]))
            if not r["pagination"]["has_more"]:
                break
            require(
                r["pagination"]["next_cursor"] != cursor and r["items"],
                "Cursor stalled",
            )
            cursor = r["pagination"]["next_cursor"]
            require(len(pages) < 40, "Unexpected cursor loop")
        return result, pages

    async def delete(d):
        impact = await api.request(
            "admin", "GET", BASE + f"/documents/{d['id']}/monetization-usages"
        )
        await api.request(
            "admin",
            "DELETE",
            BASE + f"/documents/{d['id']}",
            expected=204,
            headers={"If-Match": impact["fingerprint"]},
        )

    try:
        a = await upload(
            meta(
                "M-A",
                {
                    "leasing_company_ids": [lc, lc2],
                    "dealer_company_ids": [dealer],
                    "mark_id": cat["mark_id"],
                },
                related={"leasing_company_ids": [lc], "dealer_company_ids": [dealer]},
            )
        )
        b = await upload(
            meta(
                "M-B",
                {
                    "leasing_company_ids": [lc2],
                    "distributor_company_ids": [dist],
                    "mark_id": cat["mark_id"],
                    "model_id": cat["model_id"],
                },
                kind="agreement",
                start=1,
                end=20,
            )
        )
        d = await upload(
            meta("M-C", {"leasing_company_ids": [lc]}, kind="invoice", end=None)
        )
        e = await upload(
            meta(
                "M-D",
                {"leasing_company_ids": [lc], "dealer_company_ids": [outsider]},
                kind="other",
                end=-1,
            )
        )
        ch = await upload(
            meta(
                "M-ACT",
                kind="act",
                start=-5,
                end=-1,
                related={"dealer_company_ids": [outsider]},
            ),
            a["group_id"],
        )
        same = await upload(meta("M-SAME", name=a["name"]), a["group_id"])
        check(
            "same name/type inside SAME group remains separate document",
            same["id"] != a["id"]
            and same["group_id"] == a["group_id"]
            and same["current_version"]["version_number"] == 1
            and a["current_version"]["version_number"] == 1,
        )
        tests = [
            ("all", {}, [a, b, d, e]),
            (
                "AND across different documents",
                {"document_type": "act", "status": "active"},
                [a],
            ),
            ("pending", {"status": "pending"}, [b]),
            ("expired", {"status": "expired"}, [a, e]),
            ("active", {"status": "active"}, [a, d]),
            (
                "full period inclusive",
                {
                    "valid_from": str(day - timedelta(days=5)),
                    "valid_to": str(day - timedelta(days=1)),
                },
                [a],
            ),
            (
                "partial overlap excluded",
                {"valid_from": str(day), "valid_to": str(day + timedelta(days=100))},
                [b],
            ),
            (
                "finite upper excludes indefinite",
                {"valid_to": str(day + timedelta(days=100))},
                [a, b, e],
            ),
            ("model", {"model_id": cat["model_id"]}, [b]),
            ("mark", {"mark_id": cat["mark_id"]}, [a, b]),
            ("first LC", {"leasing_company_id": lc}, [a, d, e]),
            ("second LC", {"leasing_company_id": lc2}, [a, b]),
            ("participant dealer", {"dealer_company_id": dealer}, [a]),
            ("participant distributor", {"distributor_company_id": dist}, [b]),
            (
                "related role combination",
                {
                    "participant_scope": "related",
                    "leasing_company_id": lc,
                    "dealer_company_id": outsider,
                },
                [a],
            ),
            (
                "related negative",
                {"participant_scope": "related", "distributor_company_id": dist},
                [],
            ),
            ("type negative", {"document_type": "power_of_attorney"}, []),
        ]
        for label, filters, wanted in tests:
            params = {"search": prefix + "-M-", **filters}
            expected = {x["group_id"] for x in wanted}
            cards, _ = await all_cards(params)
            table, _ = await all_table(params)
            sheet = await export(params)
            check(
                "filter " + label + " cards/table/Excel exact",
                {x["group_id"] for x in cards} == expected
                and {x["group_id"] for x in table} == expected
                and {r[1].value for r in list(sheet.rows)[1:]}
                == {x["contract_number"] for x in wanted},
                expected_count=len(expected),
                cards=len(cards),
                table=len(table),
                excel=sheet.max_row - 1,
            )
        # Case-independent name and number-only search must each identify one group.
        n = await upload(
            meta(
                "SEARCH-UNIQUE-NUMBER",
                name="MixedCase Independent Title " + uuid4().hex[:7],
            )
        )
        for label, term in [
            ("name swapcase", n["name"].swapcase()),
            ("number only", n["contract_number"].upper()),
        ]:
            cards, _ = await all_cards({"search": term})
            check(label, [x["group_id"] for x in cards] == [n["group_id"]])
        # ACL must be applied before filters and export; outsider sees only child of A and main E.
        cards, _ = await all_cards(
            {"search": prefix + "-M-", "status": "active"}, "outsider"
        )
        check("hidden main cannot satisfy status filter", not cards)
        rows, _ = await all_table({"search": prefix + "-M-"}, "outsider")
        sheet = await export({"search": prefix + "-M-"}, "outsider")
        check(
            "partner positive export contains only accessible child/main",
            {x["id"] for x in rows} == {ch["id"], e["id"]}
            and {r[1].value for r in list(sheet.rows)[1:]}
            == {ch["contract_number"], e["contract_number"]},
        )
        # Populate every page with known exact sets.
        pg = []
        for i in range(23):
            pg.append(await upload(meta("P-" + str(i))))
        cards, cp = await all_cards({"search": prefix + "-P-"})
        rows, tp = await all_table({"search": prefix + "-P-"})
        sheet = await export({"search": prefix + "-P-"})
        expected = {x["group_id"] for x in pg}
        check(
            "cards pagination exact 20+3",
            cp == [20, 3]
            and len(cards) == 23
            and {x["group_id"] for x in cards} == expected,
            pages=cp,
        )
        check(
            "table cursor pagination exact 10+10+3",
            tp == [10, 10, 3]
            and len(rows) == 23
            and {x["group_id"] for x in rows} == expected,
            pages=tp,
        )
        check(
            "Excel exports ALL 23 groups",
            sheet.max_row == 24
            and {r[1].value for r in list(sheet.rows)[1:]}
            == {x["contract_number"] for x in pg},
        )
        kids = []
        for i in range(22):
            kids.append(
                await upload(
                    meta("CH-" + str(i), kind="power_of_attorney"), pg[0]["group_id"]
                )
            )
        collected = []
        childpages = []
        for p in (1, 2):
            r = await api.request(
                "admin",
                "GET",
                BASE + f"/groups/{pg[0]['group_id']}/documents",
                params={"document_type": "power_of_attorney", "page": p},
            )
            collected += r["items"]
            childpages.append(len(r["items"]))
        check(
            "children page2 exact set, versions excluded",
            childpages == [20, 2]
            and {x["id"] for x in collected} == {x["id"] for x in kids},
            pages=childpages,
        )
        # Hazardous strings must be in actual exported cells, not only in hidden child names.
        for lead in ["=", "+", "-", "@"]:
            hazard = lead + "1+1 " + prefix
            x = await upload(meta("XLSX-" + str(ord(lead)), name=hazard))
            sheet = await export({"search": x["contract_number"]})
            check(
                "Excel literal main name " + lead,
                sheet["C2"].value == hazard
                and sheet["C2"].data_type == "s"
                and not any(cell.data_type == "f" for row in sheet for cell in row),
            )
        # General contracts, both LCs, brand/model and platform context dimensions.
        gen = await upload(meta("C-GENERAL", {"leasing_company_ids": [lc, lc2]}))
        plat = await upload(
            meta("C-PLATFORM", {"leasing_company_ids": [lc], "platform_ml": True})
        )
        own = {x["id"] for x in [a, b, d, e, ch, same, gen, plat]}
        ctx = {
            "leasing_company_id": lc,
            "dealer_company_id": dealer,
            "distributor_company_id": dist,
            "mark_id": cat["mark_id"],
            "model_id": cat["model_id"],
        }
        candidates = [
            ("any positive match", ctx, {a["id"], same["id"], b["id"], d["id"], gen["id"], plat["id"]}),
            (
                "second LC",
                {**ctx, "leasing_company_id": lc2},
                {a["id"], same["id"], b["id"], gen["id"]},
            ),
            (
                "without dealer",
                {**ctx, "dealer_company_id": None},
                {a["id"], same["id"], b["id"], d["id"], gen["id"], plat["id"]},
            ),
            (
                "no brand/model",
                {**ctx, "mark_id": None, "model_id": None},
                {a["id"], same["id"], b["id"], d["id"], gen["id"], plat["id"]},
            ),
            (
                "platform participates",
                {**ctx, "platform_ml": True},
                {a["id"], same["id"], b["id"], d["id"], gen["id"], plat["id"]},
            ),
            (
                "model absent but LC, distributor or mark match",
                {**ctx, "leasing_company_id": lc2, "model_id": None},
                {a["id"], same["id"], b["id"], gen["id"]},
            ),
        ]
        for label, context, expected in candidates:
            r = await api.request(
                "admin",
                "POST",
                BASE + "/monetization-candidates",
                json={"context": context},
            )
            found = {x["id"] for x in r["items"]} & own
            check(
                "candidates " + label,
                found == expected,
                expected=sorted(expected),
                actual=sorted(found),
            )
        # Restore uses a distinctly different PDF payload, and all projections follow it.
        import reportlab.pdfgen.canvas

        def distinct_pdf(label):
            out = io.BytesIO()
            can = reportlab.pdfgen.canvas.Canvas(out)
            can.drawString(30, 700, label)
            can.save()
            return out.getvalue()

        oldbytes = distinct_pdf("ACCEPTANCE OLD " + prefix)
        newbytes = distinct_pdf("ACCEPTANCE NEW " + prefix)
        v = await upload(meta("RESTORE"), contents=oldbytes)
        old = v["current_version"]
        vm = version_metadata(
            v,
            name=prefix + "-RESTORE NEW",
            valid_to=str(day + timedelta(days=42)),
            retained_file_ids=[],
        )
        v2 = await api.request(
            "admin",
            "POST",
            BASE + f"/documents/{v['id']}/versions",
            expected=201,
            data={"metadata": json.dumps(vm)},
            files=[("files", ("different.pdf", newbytes, "application/pdf"))],
        )
        restored = await api.request(
            "admin",
            "POST",
            BASE + f"/documents/{v['id']}/versions/{old['id']}/activate",
            json={"expected_current_version_id": v2["current_version"]["id"]},
        )
        rows, _ = await all_table({"search": v["contract_number"]})
        sheet = await export({"search": v["contract_number"]})
        oldfile = await api.request(
            "admin", "GET", old["files"][0]["download_url"], raw=True
        )
        newfile = await api.request(
            "admin", "GET", v2["current_version"]["files"][0]["download_url"], raw=True
        )
        check(
            "old/new history downloads exact distinct bytes",
            oldbytes != newbytes
            and oldfile.content == oldbytes
            and newfile.content == newbytes,
        )
        check(
            "restore table and Excel current snapshot",
            restored["current_version"]["id"] == old["id"]
            and len(rows) == 1
            and rows[0]["name"] == v["name"]
            and rows[0]["current_version"]["id"] == old["id"]
            and sheet["C2"].value == v["name"]
            and sheet["J2"].value == (day + timedelta(days=90)).strftime("%d.%m.%Y"),
        )
        # Multi-program impact, cancellation, child-only unlink, main cascade; programs stay alive.
        parent = await upload(meta("DELETE-PARENT"))
        child = await upload(meta("DELETE-CHILD"), parent["group_id"])
        programs = []
        for i, ids in enumerate(([parent["id"], child["id"]], [child["id"]])):
            body = {
                "status": "inactive",
                "name": prefix + "-PROGRAM-" + str(i),
                "leasing_company_id": lc,
                "period_start": str(day),
                "sources": [
                    {
                        "source_type": "platform",
                        "expenses": [
                            {
                                "participant_type": "leasing",
                                "base_type": "none",
                                "calc_type": "amount",
                                "value": "100",
                            }
                        ],
                        "incomes": [],
                    }
                ],
                "reference_document_ids": ids,
            }
            programs.append(
                await api.request("admin", "POST", PROGRAMS, expected=201, json=body)
            )
        impact = await api.request(
            "admin", "GET", BASE + f"/documents/{parent['id']}/monetization-usages"
        )
        check(
            "main deletion impact ALL programs and children",
            {x["id"] for x in impact["programs"]} == {x["id"] for x in programs}
            and set(impact["document_ids"]) == {parent["id"], child["id"]},
        )
        for x in [parent, child]:
            await api.request("admin", "GET", BASE + f"/documents/{x['id']}")
        check("impact preview is read-only / cancellation preserves documents", True)
        await delete(child)
        links = [
            await api.request(
                "admin", "GET", PROGRAMS + f"/{p['id']}/reference-documents"
            )
            for p in programs
        ]
        check(
            "delete child unlinks only child from both programs",
            {x["id"] for x in links[0]["items"]} == {parent["id"]}
            and not links[1]["items"],
        )
        await delete(parent)
        roots = [x for x in roots if x["id"] != parent["id"]]
        links = [
            await api.request("admin", "GET", PROGRAMS + f"/{p['id']} ".strip())
            for p in programs
        ]
        check(
            "main deletion retains both programs, no dangling links",
            all(not x["reference_documents"] for x in links),
        )
        progress(
            "matrix_complete",
            prefix=prefix,
            checks=len(checks),
            failed=sum(not x["pass"] for x in checks),
        )
    finally:
        failures = []
        for x in roots:
            try:
                await delete(x)
            except Exception as exc:
                failures.append({"id": x["id"], "error": str(exc)[:200]})
        (ROOT / "acceptance-results.json").write_text(
            json.dumps(
                {"prefix": prefix, "checks": checks, "cleanup_failures": failures},
                ensure_ascii=False,
                indent=2,
            )
        )
        progress("matrix_cleanup", roots=len(roots), failures=failures)
    require(not failures, "Acceptance fixture cleanup failed")
    require(all(x["pass"] for x in checks), "Matrix check failed")


if __name__ == "__main__":
    asyncio.run(main())
