"""Real HTTP upload acceptance/rejection, role boundaries and OpenAPI checks."""
import asyncio
import io
import json
import struct
import zipfile
from uuid import uuid4

from PIL import Image
from runtime import API, ROOT, guard, pdf, progress, require

BASE = "/api/v1/document-registry"


def corrupt_files():
    buffer = io.BytesIO()
    Image.new("RGB", (5, 5), (30, 50, 150)).save(buffer, format="PNG")
    png = buffer.getvalue()
    bad_crc = bytearray(png)
    offset = bad_crc.find(b"IDAT")
    length = struct.unpack(">I", bad_crc[offset-4:offset])[0]
    bad_crc[offset+4+length] ^= 1
    bad_chunk = bytearray(png)
    bad_chunk[offset:offset+4] = b"!#$%"
    buffer = io.BytesIO()
    Image.new("RGB", (5, 5)).save(buffer, format="JPEG")
    yield "crc.png", bytes(bad_crc)
    yield "truncated.png", png[:-13]
    yield "chunk.png", bytes(bad_chunk)
    yield "truncated.jpg", buffer.getvalue()[:-20]
    for extension, member in (("docx", "word/document.xml"), ("xlsx", "xl/workbook.xml")):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr(member, "<document/>")
            archive.writestr("badname", "x")
        malformed = bytearray(buffer.getvalue())
        offset = 0
        while (offset := malformed.find(b"PK\x01\x02", offset)) >= 0:
            size = struct.unpack("<H", malformed[offset+28:offset+30])[0]
            if malformed[offset+46:offset+46+size] == b"badname":
                flags = struct.unpack("<H", malformed[offset+8:offset+10])[0] | 0x800
                malformed[offset+8:offset+10] = struct.pack("<H", flags)
                malformed[offset+46] = 255
            offset += 4
        yield "utf8." + extension, bytes(malformed)


def valid_files(revision):
    from openpyxl import Workbook, load_workbook
    from xml.etree import ElementTree

    for extension, image_format, content_type in (("png", "PNG", "image/png"), ("jpg", "JPEG", "image/jpeg")):
        buffer = io.BytesIO()
        Image.new("RGB", (8, 8), (30 + revision * 60, 50, 150)).save(buffer, format=image_format)
        data = buffer.getvalue()
        with Image.open(io.BytesIO(data)) as decoded:
            decoded.load()
            require(decoded.size == (8, 8) and decoded.format == image_format, "Invalid image fixture")
        yield extension, data, content_type

    # Minimal valid OPC Word package: content types, root relationship, main part.
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr("_rels/.rels", '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        archive.writestr("word/document.xml", f'<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Registry revision {revision}</w:t></w:r></w:p><w:sectPr/></w:body></w:document>')
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as archive:
        require(archive.testzip() is None, "Invalid DOCX fixture CRC")
        for member in archive.namelist():
            ElementTree.fromstring(archive.read(member))
    yield "docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    workbook = Workbook()
    workbook.active["A1"] = f"Registry revision {revision}"
    buffer = io.BytesIO()
    workbook.save(buffer)
    require(load_workbook(io.BytesIO(buffer.getvalue())).active["A1"].value == f"Registry revision {revision}", "Invalid XLSX fixture")
    yield "xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def remove_owned_document(api, document_id):
    path = BASE + "/documents/" + document_id
    preview = await api.request("admin", "GET", path + "/monetization-usages", expected=(200, 404), raw=True)
    if preview.status_code == 200:
        await api.request("admin", "DELETE", path, expected=204, headers={"If-Match": preview.json()["fingerprint"]})
    await api.request("admin", "GET", path, expected=404)


async def verify_valid_uploads(api):
    from versioning import version_metadata

    prefix = "22296-valid-uploads-" + uuid4().hex[:10]
    owned = []
    downloads = 0
    try:
        revisions = [dict((extension, (data, mime)) for extension, data, mime in valid_files(number)) for number in (1, 2)]
        for extension in ("png", "jpg", "docx", "xlsx"):
            first_bytes, content_type = revisions[0][extension]
            second_bytes, _ = revisions[1][extension]
            require(first_bytes != second_bytes, "Revision fixture bytes must differ")
            metadata = {"document_type": "contract", "contract_number": prefix + "-" + extension,
                "name": prefix + "-" + extension, "valid_from": "2026-01-01", "participants": {"platform_ml": True}}
            document = await api.request("admin", "POST", BASE + "/documents", expected=201,
                data={"metadata": json.dumps(metadata)}, files=[("files", ("first." + extension, first_bytes, content_type))])
            owned.append(document["id"])
            path = BASE + "/documents/" + document["id"]
            updated = await api.request("admin", "POST", path + "/versions", expected=201,
                data={"metadata": json.dumps(version_metadata(document, retained_file_ids=[]))},
                files=[("files", ("second." + extension, second_bytes, content_type))])
            require(updated["version_count"] == 2 and updated["current_version"]["id"] != document["current_version"]["id"], "Valid replacement did not create a new version")
            history = await api.request("admin", "GET", path + "/versions")
            require({version["id"] for version in history["items"]} == {document["current_version"]["id"], updated["current_version"]["id"]}, "Valid file history incomplete")
            for version, expected_bytes in ((document["current_version"], first_bytes), (updated["current_version"], second_bytes)):
                require(len(version["files"]) == 1 and version["files"][0]["type"] == content_type, "Wrong file metadata")
                downloaded = await api.request("admin", "GET", version["files"][0]["download_url"], raw=True)
                require(downloaded.content == expected_bytes and downloaded.headers["content-type"] == content_type, "Stored file bytes or download content type changed")
                downloads += 1
        progress("registry_valid_uploads_passed", formats=["PNG", "JPEG", "DOCX", "XLSX"], creates=4, versions=4, exact_downloads=downloads)
    finally:
        for document_id in owned:
            await remove_owned_document(api, document_id)


async def verify_role_boundaries(api):
    from versioning import version_metadata

    prefix = "22296-role-boundaries-" + uuid4().hex[:10]
    companies = api.state["companies"]
    metadata = {"document_type": "contract", "contract_number": prefix, "name": prefix,
        "valid_from": "2026-01-01", "participants": {
            "leasing_company_ids": [companies["leasing"]["leasing_company_id"]],
            "dealer_company_ids": [companies["dealer"]["company_id"]],
            "distributor_company_ids": [companies["distributor"]["company_id"]]}}
    attachments = [("files", ("roles.pdf", pdf(prefix), "application/pdf"))]
    document = await api.request("admin", "POST", BASE + "/documents", expected=201,
        data={"metadata": json.dumps(metadata)}, files=attachments)
    document_path = BASE + "/documents/" + document["id"]
    programs = "/api/v1/admin/monetization/programs"
    rejected = 0
    try:
        program = await api.request("admin", "POST", programs, expected=201, json={
            "status": "inactive", "name": prefix, "leasing_company_id": companies["leasing"]["leasing_company_id"],
            "dealer_company_id": companies["dealer"]["company_id"], "distributor_company_id": companies["distributor"]["company_id"],
            "period_start": "2026-01-01", "sources": [{"source_type": "platform", "expenses": [{
                "participant_type": "leasing", "base_type": "none", "calc_type": "amount", "value": "100"}], "incomes": []}],
            "reference_document_ids": [document["id"]]})
        program_path = programs + "/" + program["id"] + "/reference-documents"
        for role in ("dealer", "leasing", "distributor"):
            await api.request(role, "GET", document_path)
            await api.request(role, "GET", document_path + "/versions")
            await api.request(role, "GET", document["current_version"]["files"][0]["download_url"], raw=True)
            await api.request(role, "GET", program_path, expected=404 if role == "distributor" else 200)
        for path in (document_path, document_path + "/versions", document["current_version"]["files"][0]["download_url"], BASE + "/groups/" + document["group_id"] + "/documents", program_path):
            await api.request("outsider", "GET", path, expected=404)
        impact = await api.request("admin", "GET", document_path + "/monetization-usages")
        missing = str(uuid4())
        child_metadata = {key: value for key, value in metadata.items() if key != "participants"}
        child_metadata["contract_number"] += "-child"
        for role in ("dealer", "leasing", "distributor", "outsider", "client", "anonymous"):
            expected = 401 if role == "anonymous" else 403
            cases = [
                ("POST", BASE + "/documents", {"data": {"metadata": json.dumps({**metadata, "contract_number": prefix + "-forbidden"})}, "files": attachments}),
                ("GET", BASE + "/check-contract-number", {"params": {"number": prefix}}),
            ]
            for existing in (True, False):
                target = document["id"] if existing else missing
                target_path = BASE + "/documents/" + target
                cases.extend([
                    ("POST", BASE + "/groups/" + (document["group_id"] if existing else missing) + "/documents", {"data": {"metadata": json.dumps(child_metadata)}, "files": attachments}),
                    ("POST", target_path + "/versions", {"data": {"metadata": json.dumps(version_metadata(document))}}),
                    ("POST", target_path + "/versions/" + document["current_version"]["id"] + "/activate", {"json": {"expected_current_version_id": document["current_version"]["id"]}}),
                    ("PATCH", target_path + "/activation", {"json": {"active": False}}),
                    ("DELETE", target_path, {"headers": {"If-Match": impact["fingerprint"]}}),
                    ("GET", target_path + "/monetization-usages", {}),
                    ("PATCH", programs + "/" + (program["id"] if existing else missing) + "/reference-documents", {"json": {"document_ids": []}}),
                ])
            for method, path, kwargs in cases:
                await api.request(role, method, path, expected=expected, **kwargs)
                rejected += 1
        require(await api.request("admin", "GET", document_path) == document, "Rejected role request changed the document")
        links = await api.request("admin", "GET", program_path)
        require([item["id"] for item in links["items"]] == [document["id"]], "Rejected role request changed program links")
        require((await api.request("admin", "GET", BASE + "/check-contract-number", params={"number": prefix}))["available"] is False, "Admin number lookup denied")
        child = await api.request("admin", "POST", BASE + "/groups/" + document["group_id"] + "/documents", expected=201,
            data={"metadata": json.dumps(child_metadata)}, files=attachments)
        updated = await api.request("admin", "POST", document_path + "/versions", expected=201, data={"metadata": json.dumps(version_metadata(document, name=prefix + "-v2"))})
        await api.request("admin", "POST", document_path + "/versions/" + document["current_version"]["id"] + "/activate", json={"expected_current_version_id": updated["current_version"]["id"]})
        for active in (False, True):
            await api.request("admin", "PATCH", document_path + "/activation", json={"active": active})
        for selected in ([], [document["id"]]):
            result = await api.request("admin", "PATCH", program_path, json={"document_ids": selected})
            require([item["id"] for item in result["items"]] == selected, "Admin link write failed")
        await remove_owned_document(api, child["id"])
        progress("registry_http_role_boundaries_passed", admin_operations=9, rejected_requests=rejected, existing_and_missing_ids=True, allowed_partner_reads=11, outsider_hidden_reads=5, distributor_without_dealer_link_status=404)
    finally:
        await remove_owned_document(api, document["id"])


async def verify_upload_boundaries(api):
    await verify_valid_uploads(api)
    await verify_role_boundaries(api)
    from sqlalchemy import func, select
    from infrastructure.database import AsyncSessionLocal
    from infrastructure.models.document_registry import reference_documents
    prefix = "22296-upload-boundaries-" + uuid4().hex[:10]
    metadata = {"document_type": "contract", "contract_number": prefix, "name": prefix,
                "valid_from": "2026-01-01", "participants": {"platform_ml": True}}
    original = await api.request("admin", "POST", BASE + "/documents", expected=201,
        data={"metadata": json.dumps(metadata)}, files=[("files", ("valid.pdf", pdf(), "application/pdf"))])
    from versioning import version_metadata
    count = 0
    for filename, content in corrupt_files():
        for new_version in (False, True):
            body = version_metadata(original) if new_version else {**metadata, "contract_number": prefix + "-" + filename}
            path = BASE + "/documents" + ("/" + original["id"] + "/versions" if new_version else "")
            await api.request("admin", "POST", path, expected=415,
                data={"metadata": json.dumps(body)}, files=[("files", (filename, content, "application/octet-stream"))])
            count += 1
    current = await api.request("admin", "GET", BASE + "/documents/" + original["id"])
    require(current == original, "Rejected upload changed the document/version")
    async with AsyncSessionLocal() as session:
        total = await session.scalar(select(func.count()).select_from(reference_documents).where(reference_documents.c.contract_number.startswith(prefix)))
    require(total == 1, "Rejected upload persisted a partial document")
    from transactions import database_facts
    facts = await database_facts(original)
    require(facts["committed_versions"] == 1 and facts["referenced_object_count_present"] == 1 and facts["storage_objects"] == 1, "Rejected upload left partial SQL/S3 files")
    progress("registry_corrupt_uploads_passed", requests=count, partial_documents=0, partial_versions=0, partial_objects=0)


async def verify_openapi(api):
    allowed = {}
    for role in ("admin", "dealer", "leasing", "distributor", "client", "anonymous"):
        schema = await api.request(role, "GET", "/api/v1/openapi.json")
        link_methods = set(schema["paths"].get("/api/v1/admin/monetization/programs/{program_id}/reference-documents", {}))
        require(("patch" in link_methods) == (role == "admin"), role + " has incorrect link-write documentation")
        routes = {path: set(methods) for path, methods in schema["paths"].items() if path.startswith(BASE)}
        allowed[role] = sum(len(methods) for methods in routes.values())
        if role in {"client", "anonymous"}:
            require(not routes, role + " sees internal registry API")
            continue
        require("get" in routes[BASE + "/documents/{document_id}"], role + " lacks readable documents")
        if role != "admin":
            require(all(methods <= {"get"} or (path == BASE + "/monetization-candidates" and methods == {"post"}) for path, methods in routes.items()), role + " sees write routes")
            require(BASE + "/check-contract-number" not in routes and BASE + "/documents/{document_id}/monetization-usages" not in routes, role + " sees admin-only reads")
        else:
            required = (("post", "/documents"), ("post", "/groups/{group_id}/documents"),
                ("post", "/documents/{document_id}/versions"), ("post", "/documents/{document_id}/versions/{version_id}/activate"),
                ("patch", "/documents/{document_id}/activation"), ("delete", "/documents/{document_id}"),
                ("get", "/check-contract-number"), ("get", "/documents/{document_id}/monetization-usages"))
            require(all(method in routes.get(BASE + path, set()) for method, path in required), "Admin operations disappeared")
    progress("registry_openapi_roles_passed", operations=allowed)


async def main():
    guard()
    api = API(json.loads((ROOT / "state.secret.json").read_text()))
    await verify_upload_boundaries(api)
    await verify_openapi(api)


if __name__ == "__main__":
    asyncio.run(main())
