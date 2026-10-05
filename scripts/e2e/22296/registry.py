"""Task 22296 acceptance against real HTTP, PostgreSQL and object storage."""
from __future__ import annotations

import asyncio
import io
import json
from datetime import datetime, timedelta
from uuid import UUID, uuid4
from zoneinfo import ZoneInfo

from runtime import API, ROOT, guard, pdf, private_json, progress, require
from versioning import verify_inactive_relationships, verify_version_lifecycle, version_metadata

BASE = "/api/v1/document-registry"
PROGRAMS = "/api/v1/admin/monetization/programs"


async def verify_program_status_response(api: API, program_id: str, expected_ids: set[str]):
    response = await api.request("admin", "PATCH", PROGRAMS + f"/{program_id}", json={"status": "inactive"})
    require({document["id"] for document in response["reference_documents"]} == expected_ids, "Program status response lost reference documents")
    fresh = await api.request("admin", "GET", PROGRAMS + f"/{program_id}")
    require(response["reference_documents"] == fresh["reference_documents"], "Program status response differs from GET")


async def verify_status_contract():
    guard()
    state = json.loads((ROOT / "state.secret.json").read_text())
    manifest = json.loads((ROOT / "manifest.json").read_text())
    api = API(state)
    program_id = manifest["registry"]["program_id"]
    document_id = manifest["registry"]["main_document_id"]
    before = await api.request("admin", "GET", PROGRAMS + f"/{program_id}/reference-documents")
    try:
        await api.request("admin", "PATCH", PROGRAMS + f"/{program_id}/reference-documents", json={"document_ids": [document_id]})
        await verify_program_status_response(api, program_id, {document_id})
    finally:
        await api.request("admin", "PATCH", PROGRAMS + f"/{program_id}/reference-documents", json={"document_ids": [document["id"] for document in before["items"]]})
    progress("registry_program_status_response_passed")


async def verify():
    guard()
    state = json.loads((ROOT / "state.secret.json").read_text())
    api = API(state)
    day = datetime.now(ZoneInfo("Europe/Moscow")).date()
    prefix = "22296-" + uuid4().hex[:8]
    companies = state["companies"]
    catalog = state["catalog"]
    lc = companies["leasing"]["leasing_company_id"]
    lc2 = companies["leasing2"]["leasing_company_id"]
    dealer = companies["dealer"]["company_id"]
    outsider = companies["dealer2"]["company_id"]
    distributor = companies["distributor"]["company_id"]
    binding = {"leasing_company_ids": [lc, lc2], "dealer_company_ids": [dealer], "distributor_company_ids": [distributor], "mark_id": catalog["mark_id"]}
    general = {"leasing_company_ids": [lc, lc2], "mark_id": catalog["mark_id"]}
    data = pdf()
    created = []

    def metadata(name, *, participants=None, start=-10, end=90, kind="contract", related=None):
        return {"document_type": kind, "contract_number": prefix + "-" + name, "name": prefix + " " + name,
            "valid_from": str(day + timedelta(days=start)), "valid_to": str(day + timedelta(days=end)) if end is not None else None,
            "participants": participants if participants is not None else binding,
            "related_companies": related or {}}

    async def upload(payload, *, group=None, role="admin", expected=201, attachments=None):
        body = dict(payload)
        path = BASE + "/documents"
        if group:
            path = BASE + f"/groups/{group}/documents"
            body.pop("participants", None)
        result = await api.request(role, "POST", path, expected=expected,
            data={"metadata": json.dumps(body)}, files=attachments if attachments is not None else [("files", ("contract.pdf", data, "application/pdf"))])
        if expected == 201:
            created.append(result)
        return result

    async def version(document, *, start=-10, end=90, **changes):
        current = await api.request("admin", "GET", BASE + f"/documents/{document['id']}")
        metadata = version_metadata(current, valid_from=str(day + timedelta(days=start)),
            valid_to=str(day + timedelta(days=end)) if end is not None else None, retained_file_ids=[], **changes)
        return await api.request("admin", "POST", BASE + f"/documents/{document['id']}/versions", expected=201,
            data={"metadata": json.dumps(metadata)}, files=[("files", ("new.pdf", data, "application/pdf"))])

    async def matches(filters):
        response = await api.request("admin", "GET", BASE + "/groups", params={"search": prefix, **filters})
        return {item["group_id"] for item in response["items"]}

    await api.request("guest", "GET", BASE + "/groups", expected=401)
    await api.request("client", "GET", BASE + "/groups", expected=403)
    for invalid_cursor in ("e30=", "W10=", "ImFiYyI=", "!!!badbase64"):
        await api.request("admin", "GET", BASE + "/table", expected=422, params={"cursor": invalid_cursor})
    types = await api.request("admin", "GET", BASE + "/types")
    require(len(types["items"]) == 7, "Missing registry types")
    lookups = await api.request("admin", "GET", BASE + "/lookups/companies", params={"role": "leasing_company", "search": "22296"})
    require(any(item["leasing_company_id"] == lc and item["company_id"] == companies["leasing"]["company_id"] for item in lookups["items"]), "LC identifiers confused")
    await upload(metadata("forbidden"), role="dealer", expected=403)
    await upload(metadata("empty", participants={}), expected=422)
    await upload(metadata("unknown-lc", participants={"leasing_company_ids": [str(uuid4())]}), expected=422)
    await upload(metadata("wrong-role", participants={"dealer_company_ids": [companies["leasing"]["company_id"]]}), expected=422)
    await upload(metadata("wrong-model", participants={"mark_id": str(uuid4()), "model_id": catalog["model_id"]}), expected=422)
    await upload(metadata("wrong-period", start=5, end=4), expected=422)
    expanding_number = metadata("expanding-number")
    expanding_number["contract_number"] = "ß" * 100
    await upload(expanding_number, expected=422)
    await upload(metadata("fake-format"), expected=415, attachments=[("files", ("fake.pdf", b"not-pdf", "application/pdf"))])
    await upload(metadata("empty-file"), expected=422, attachments=[("files", ("empty.pdf", b"", "application/pdf"))])
    await upload(metadata("big-file"), expected=413, attachments=[("files", ("big.pdf", b"x"*(20*1024*1024+1), "application/pdf"))])
    require(not await matches({}), "Validation left partial groups/documents")
    progress("registry_validation_and_role_guards")

    main = await upload(metadata("main", related={"leasing_company_ids": [lc], "dealer_company_ids": [dealer]}), attachments=[("files", ("contract.pdf", data, "application/pdf")), ("files", ("annex.pdf", data, "application/pdf"))])
    require(main["is_main"] and len(main["current_version"]["files"]) == 2, "Wrong main version")
    for literal in ("%", "_"):
        require(not (await api.request("admin", "GET", BASE + "/groups", params={"search": prefix + literal}))["items"], "Search interpreted SQL wildcard")
    repeat = metadata("different")
    repeat["contract_number"] = "  " + main["contract_number"].upper().replace("-", " - ") + " "
    await upload(repeat, expected=409)
    number = await api.request("admin", "GET", BASE + "/check-contract-number", params={"number": repeat["contract_number"]})
    require(number["available"] is False, "Number check differs from save")
    await api.request("admin", "PATCH", BASE + f"/documents/{main['id']}", expected=405, json={"contract_number": "changed"})
    await api.request("admin", "PATCH", BASE + f"/documents/{main['id']}", expected=405, json={"document_type": "act"})
    same_name = metadata("same-name")
    same_name["name"] = main["name"]
    independent = await upload(same_name, group=main["group_id"])
    require(independent["group_id"] == main["group_id"] and independent["id"] != main["id"] and independent["current_version"]["version_number"] == 1, "Separate number treated as version")
    child = await upload(metadata("child", kind="act", end=-1, related={"dealer_company_ids": [outsider]}), group=main["group_id"])
    require(child["related_companies"]["leasing_company_ids"] == [lc] and child["related_companies"]["dealer_company_ids"] == [outsider], "Per-field inheritance incorrect")
    main = await version(main, related_companies={"platform_ml": False, "leasing_company_ids": [], "dealer_company_ids": [], "distributor_company_ids": []})
    unchanged = await api.request("admin", "GET", BASE + f"/documents/{child['id']}")
    require(unchanged["related_companies"]["leasing_company_ids"] == [lc], "Inheritance became live")
    for role in ("dealer", "leasing", "leasing2", "distributor"):
        await api.request(role, "GET", BASE + f"/documents/{main['id']}")
        await api.request(role, "GET", main["current_version"]["files"][0]["download_url"], raw=True)
        await api.request(role, "PATCH", BASE + f"/documents/{main['id']}/activation", expected=403, json={"active": False})
    await api.request("outsider", "GET", BASE + f"/documents/{main['id']}", expected=404)
    await api.request("outsider", "GET", main["current_version"]["files"][0]["download_url"], expected=404)
    await api.request("outsider", "GET", BASE + f"/documents/{main['id']}/versions", expected=404)
    hidden_group = await api.request("outsider", "GET", BASE + "/groups", params={"search": prefix})
    hidden = next(group for group in hidden_group["items"] if group["group_id"] == main["group_id"])
    require(hidden["main_document"] is None and hidden["documents_count"] == 1 and hidden["display_document"]["id"] == child["id"], "Hidden main leaked in group projection")
    require(main["group_id"] in await matches({"document_type": "act", "status": "active"}), "Filters require same document incorrectly")
    outsider_filter = await api.request("outsider", "GET", BASE + "/groups", params={"search": prefix, "status": "active"})
    require(main["group_id"] not in {item["group_id"] for item in outsider_filter["items"]}, "Hidden main satisfied filter")
    require(main["group_id"] in await matches({"valid_from": str(day-timedelta(days=10)), "valid_to": str(day-timedelta(days=1))}), "Inclusive full period failed")
    require(main["group_id"] not in await matches({"valid_from": str(day), "valid_to": str(day+timedelta(days=100))}), "Partial overlap accepted")
    related_match = await matches({"participant_scope": "related", "leasing_company_id": lc, "dealer_company_id": outsider})
    require(main["group_id"] in related_match, "Related filters failed")
    progress("registry_acl_inheritance_group_filters")

    concurrent_payload = metadata("race-number")
    async def race_create():
        return await api.request("admin", "POST", BASE + "/documents", expected=(201,409), raw=True, data={"metadata": json.dumps(concurrent_payload)}, files=[("files", ("contract.pdf", data, "application/pdf"))])
    races = await asyncio.gather(race_create(), race_create())
    require(sorted(response.status_code for response in races) == [201,409], "Concurrent number uniqueness failed")
    concurrent_version = version_metadata(main)
    async def race_version():
        return await api.request("admin", "POST", BASE + f"/documents/{main['id']}/versions", expected=(201,409), raw=True, data={"metadata": json.dumps(concurrent_version)})
    raced_versions = await asyncio.gather(race_version(), race_version())
    require(sorted(response.status_code for response in raced_versions) == [201,409], "Concurrent stale version not rejected")
    history = await api.request("admin", "GET", BASE + f"/documents/{main['id']}/versions")
    require(len(history["items"]) == 3, "Version history lost")
    await api.request("dealer", "GET", main["current_version"]["files"][0]["download_url"], raw=True)
    await api.request("admin", "PATCH", BASE + f"/documents/{child['id']}/activation", json={"active": False})
    updated_child = await version(child, end=90)
    require(updated_child["status"] == "deactivated", "New version activated document")
    child = await api.request("admin", "PATCH", BASE + f"/documents/{child['id']}/activation", json={"active": True})
    require(child["status"] == "active", "Reactivation failed")
    for label, start, end, expected in [("future",1,40,"pending"),("infinite",-1,None,"active"),("today",-1,0,"expiring"),("thirty",-1,30,"expiring"),("thirtyone",-1,31,"active"),("expired",-10,-1,"expired")]:
        document = await upload(metadata(label, participants=general, start=start, end=end))
        require(document["status"] == expected, "Incorrect MSK status: " + label)
        state[label] = document
    require(state["infinite"]["group_id"] not in await matches({"valid_to": str(day+timedelta(days=1000))}), "Infinite period passed upper bound")
    stopped = await upload(metadata("stopped", participants=general))
    stopped = await api.request("admin", "PATCH", BASE + f"/documents/{stopped['id']}/activation", json={"active": False})
    context = {"leasing_company_id": lc, "dealer_company_id": dealer, "distributor_company_id": distributor, "mark_id": catalog["mark_id"], "model_id": catalog["model_id"], "platform_ml": False}
    candidates = await api.request("admin", "POST", BASE + "/monetization-candidates", json={"context": context})
    candidate_ids = {document["id"] for document in candidates["items"]}
    require(main["id"] in candidate_ids and state["future"]["id"] in candidate_ids, "Specific or general document missing")
    require(state["expired"]["id"] not in candidate_ids and stopped["id"] not in candidate_ids, "Inactive candidate offered")
    require(stopped["group_id"] in {group["group_id"] for group in candidates["groups"]}, "Inactive group cannot receive new child")
    no_dealer = await api.request("admin", "POST", BASE + "/monetization-candidates", json={"context": {**context, "dealer_company_id": None}})
    require(main["id"] in {document["id"] for document in no_dealer["items"]} and state["future"]["id"] in {document["id"] for document in no_dealer["items"]}, "Missing dealer vetoed another positive match")
    version_doc = await upload(metadata("version-lifecycle", participants={"leasing_company_ids": [lc]}, related={"dealer_company_ids": [dealer]}), attachments=[("files", ("first.pdf", data, "application/pdf")), ("files", ("second.pdf", data, "application/pdf"))])
    await verify_version_lifecycle(api, version_doc, independent, state)
    await verify_inactive_relationships(api, state)
    from matching_regression import main as verify_matching_regression
    await verify_matching_regression()
    progress("registry_concurrency_statuses_candidates")

    program_input = {"status": "inactive", "name": prefix + " conditions", "leasing_company_id": lc, "dealer_company_id": dealer, "distributor_company_id": distributor, "brand": catalog["brand"], "model": catalog["model"], "period_start": str(day), "sources": [{"source_type": "platform", "expenses": [{"participant_type": "leasing", "base_type": "none", "calc_type": "amount", "value": "100"}], "incomes": []}], "reference_document_ids": [main["id"], child["id"]]}
    program = await api.request("admin", "POST", PROGRAMS, expected=201, json=program_input)
    linked = await api.request("admin", "GET", PROGRAMS + f"/{program['id']}/reference-documents")
    require({document["id"] for document in linked["items"]} == {main["id"], child["id"]}, "Atomic program create links missing")
    await verify_program_status_response(api, program["id"], {main["id"], child["id"]})
    bad_input = {**program_input, "name": prefix + " rejected-program", "reference_document_ids": [state["expired"]["id"]]}
    await api.request("admin", "POST", PROGRAMS, expected=409, json=bad_input)
    program_list = await api.request("admin", "GET", PROGRAMS, params={"page_size": 100})
    require(not any(item["name"] == bad_input["name"] for item in program_list["items"]), "Failed link validation left partial program")
    after_version = await version(main)
    linked = await api.request("admin", "GET", PROGRAMS + f"/{program['id']}/reference-documents")
    require(next(document for document in linked["items"] if document["id"] == main["id"])["current_version"]["id"] == after_version["current_version"]["id"], "Condition did not follow current version")
    restored_main = await api.request("admin", "POST", BASE + f"/documents/{main['id']}/versions/{main['current_version']['id']}/activate", json={"expected_current_version_id": after_version["current_version"]["id"]})
    restored_links = await api.request("admin", "GET", PROGRAMS + f"/{program['id']}/reference-documents")
    require(next(document for document in restored_links["items"] if document["id"] == main["id"])["current_version"]["id"] == restored_main["current_version"]["id"], "Monetization did not follow restored version")
    await version(main, end=-1)
    retained = await api.request("admin", "PATCH", PROGRAMS + f"/{program['id']}/reference-documents", json={"document_ids": [main["id"], child["id"]]})
    require(any(document["id"] == main["id"] and document["status"] == "expired" for document in retained["items"]), "Existing expired relation could not be retained")
    await api.request("admin", "PATCH", PROGRAMS + f"/{program['id']}/reference-documents", expected=409, json={"document_ids": [main["id"], child["id"], stopped["id"]]})
    await api.request("dealer", "PATCH", PROGRAMS + f"/{program['id']}/reference-documents", expected=403, json={"document_ids": []})
    await api.request("dealer", "POST", BASE + "/monetization-candidates", json={"program_id": program["id"]})
    before_delete = await api.request("admin", "GET", BASE + f"/documents/{main['id']}/monetization-usages")
    require(program["id"] in {item["id"] for item in before_delete["programs"]} and child["id"] in before_delete["document_ids"], "Cascade impact omitted child links")
    extra = await upload(metadata("new-child"), group=main["group_id"])
    await api.request("admin", "DELETE", BASE + f"/documents/{main['id']}", expected=409, headers={"If-Match": before_delete["fingerprint"]})
    impact = await api.request("admin", "GET", BASE + f"/documents/{main['id']}/monetization-usages")
    await api.request("admin", "DELETE", BASE + f"/documents/{main['id']}", expected=204, headers={"If-Match": impact["fingerprint"]})
    for document in (main, child, extra, independent):
        await api.request("admin", "GET", BASE + f"/documents/{document['id']}", expected=404)
        await api.request("admin", "GET", document["current_version"]["files"][0]["download_url"], expected=404)
    linked = await api.request("admin", "GET", PROGRAMS + f"/{program['id']}/reference-documents")
    require(linked["items"] == [], "Soft deletion did not remove M2M")
    await api.request("admin", "GET", PROGRAMS + f"/{program['id']}")
    reused = metadata("reuse")
    reused["contract_number"] = main["contract_number"]
    await upload(reused)
    progress("registry_monetization_atomicity_and_deletion")

    export_metadata = metadata("export")
    export_metadata["name"] = "=1+1"
    export_main = await upload(export_metadata)
    from faults import verify_faults
    await verify_faults(api, export_main)
    progress("registry_real_db_failure_s3_cleanup")
    formula = metadata("formula", kind="act")
    formula["name"] = "=1+1"
    export_child = await upload(formula, group=export_main["group_id"])
    table = await api.request("admin", "GET", BASE + "/table", params={"search": prefix, "limit": 10})
    found = next(item for item in table["items"] if item["group_id"] == export_main["group_id"])
    require(found["related_documents"][0]["id"] == export_child["id"] and " + " in found["related_documents"][0]["label"], "Related column not projected")
    from openpyxl import load_workbook
    xlsx = await api.request("admin", "GET", BASE + "/table/export", params={"search": prefix}, raw=True)
    workbook = load_workbook(io.BytesIO(xlsx.content))
    require(all(cell.data_type != "f" for row in workbook.active for cell in row), "Excel contains formulas")
    exported_row = next(row for row in workbook.active if row[1].value == export_main["contract_number"])
    require(exported_row[2].value == "=1+1" and exported_row[2].data_type == "s", "Dangerous main name was not exported as literal text")
    hidden_export = await api.request("outsider", "GET", BASE + "/table/export", params={"search": prefix}, raw=True)
    hidden_book = load_workbook(io.BytesIO(hidden_export.content))
    require(not any(main["contract_number"] == cell.value for row in hidden_book.active for cell in row), "Unauthorized number in Excel")
    pagination_docs = [await upload(metadata(f"pagination-{index}", participants=general)) for index in range(22)]
    page_search = prefix + "-pagination-"
    expected_groups = {document["group_id"] for document in pagination_docs}
    first = await api.request("admin", "GET", BASE + "/groups", params={"search": page_search, "page_size": 20})
    second = await api.request("admin", "GET", BASE + "/groups", params={"search": page_search, "page_size": 20, "page": 2})
    cards = first["items"] + second["items"]
    require([len(first["items"]), len(second["items"])] == [20, 2] and len(cards) == len({item["group_id"] for item in cards}) == 22 and {item["group_id"] for item in cards} == expected_groups, "Page pagination omitted or duplicated groups")
    page_one = await api.request("admin", "GET", BASE + "/table", params={"search": page_search})
    page_two = await api.request("admin", "GET", BASE + "/table", params={"search": page_search, "cursor": page_one["pagination"]["next_cursor"]})
    page_three = await api.request("admin", "GET", BASE + "/table", params={"search": page_search, "cursor": page_two["pagination"]["next_cursor"]})
    table_items = page_one["items"] + page_two["items"] + page_three["items"]
    require([len(page["items"]) for page in (page_one, page_two, page_three)] == [10, 10, 2] and not page_three["pagination"]["has_more"] and len(table_items) == len({item["group_id"] for item in table_items}) == 22 and {item["group_id"] for item in table_items} == expected_groups, "Cursor pagination omitted or duplicated groups")
    all_export = await api.request("admin", "GET", BASE + "/table/export", params={"search": page_search}, raw=True)
    all_sheet = load_workbook(io.BytesIO(all_export.content)).active
    require(all_sheet.max_row == 23 and {row[1].value for row in list(all_sheet.rows)[1:]} == {document["contract_number"] for document in pagination_docs}, "Export omitted or duplicated groups")
    progress("registry_table_xlsx_pagination")

    notification_doc = await upload(metadata("notification", participants={"leasing_company_ids": [lc], "dealer_company_ids": [dealer]}, end=10, related={"dealer_company_ids": [dealer], "distributor_company_ids": [distributor]}))
    from notifications import verify_cleanup, verify_notifications, verify_reactivation, verify_version_switches
    expected_users = {UUID(state["users"][role]["id"]) for role in ("leasing", "dealer", "distributor")}
    await verify_notifications(UUID(notification_doc["id"]), expected_users)
    lifecycle_doc = await upload(metadata("reactivation", participants={"leasing_company_ids": [lc], "dealer_company_ids": [dealer]}, end=90, related={"distributor_company_ids": [distributor]}))
    await api.request("admin", "PATCH", BASE + f"/documents/{lifecycle_doc['id']}/activation", json={"active": False})
    await version(lifecycle_doc, end=10)
    await verify_reactivation(api, UUID(lifecycle_doc["id"]), expected_users)
    await verify_version_switches(api, state)
    from sqlalchemy import select
    from infrastructure.database import AsyncSessionLocal
    from infrastructure.models.document_registry import reference_document_files
    async with AsyncSessionLocal() as session:
        key = await session.scalar(select(reference_document_files.c.s3_key).where(reference_document_files.c.id == UUID(notification_doc["current_version"]["files"][0]["id"])))
    await verify_cleanup(key)
    manifest = json.loads((ROOT / "manifest.json").read_text())
    manifest["registry"] = {"main_document_id": export_main["id"], "group_id": export_main["group_id"], "child_document_id": export_child["id"], "program_id": program["id"], "prefix": prefix}
    private_json(ROOT / "manifest.json", manifest)
    progress("registry_notifications_cleanup", recipients=len(expected_users))
    progress("registry_e2e_passed", prefix=prefix, created_documents=len(created), program_id=program["id"])


async def verify_proxy_upload():
    """A valid multi-file version must pass nginx even above 20 MiB in total."""
    import httpx

    guard()
    state = json.loads((ROOT / "state.secret.json").read_text())
    actor = state["users"]["admin"]
    api = API(state)
    # A PDF comment before EOF leaves all existing object offsets intact.
    data = pdf().replace(b"%%EOF", b"%" + b"x" * (11 * 1024 * 1024) + b"\n%%EOF")
    metadata = {"document_type": "contract", "contract_number": "proxy-" + uuid4().hex,
                "name": "Large multipart proxy check", "valid_from": "2026-01-01",
                "participants": {"platform_ml": True}}
    async with httpx.AsyncClient(base_url="http://nginx", trust_env=False, timeout=60,
            cookies={"accessToken": actor["access_token"], "csrfToken": actor["csrf"]},
            headers={"Origin": "http://localhost:18296", "Host": "localhost:18296",
                     "X-CSRF-Token": actor["csrf"]}) as client:
        response = await client.post(BASE + "/documents", data={"metadata": json.dumps(metadata)},
            files=[("files", (name, data, "application/pdf")) for name in ("one.pdf", "two.pdf")])
    require(response.status_code == 201, f"Proxy rejected valid multipart: {response.status_code}")
    document = response.json()
    files = document["current_version"]["files"]
    require(len(files) == 2, "Proxy upload lost a file")
    downloaded = await api.request("admin", "GET", files[0]["download_url"], raw=True)
    require(downloaded.content == data, "Proxy upload changed file bytes")
    impact = await api.request("admin", "GET", BASE + f"/documents/{document['id']}/monetization-usages")
    await api.request("admin", "DELETE", BASE + f"/documents/{document['id']}", expected=204,
                      headers={"If-Match": impact["fingerprint"]})
    progress("registry_proxy_large_multipart_passed", total_file_bytes=len(data) * 2)
